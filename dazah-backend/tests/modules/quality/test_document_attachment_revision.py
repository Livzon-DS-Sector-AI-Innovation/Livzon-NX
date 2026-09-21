"""附件正文驱动的版本替换回归，不连接数据库或外部服务。"""

import io
from datetime import date
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from docx import Document
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AppException
from app.modules.quality.models.document_catalog import DocumentEntry
from app.modules.quality.service import document_catalog_attachment as storage
from app.modules.quality.service import document_catalog_docx_md as converter


def _docx(code: str = "SMP-XF2-001/04", effective: str = "2026年09月01日") -> bytes:
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows,
        [("文件编号", code), ("批准日期", "2025-01-01"), ("生效日期", effective)],
        strict=True,
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_conversion_uses_document_code_not_filename():
    md, _ = converter.convert_docx_content_to_md(_docx(), "SMP-XF2-001-99.docx")
    assert "**文件编号**: SMP-XF2-001/04" in md
    assert "**生效日期**: 2026年09月01日" in md
    assert "**文件编号**: SMP-XF2-001-99" not in md


def _docm() -> bytes:
    from zipfile import ZipFile

    output = io.BytesIO()
    with ZipFile(io.BytesIO(_docx())) as source, ZipFile(output, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                    b"application/vnd.ms-word.document.macroEnabled.main+xml",
                )
            target.writestr(item, data)
    return output.getvalue()


@pytest.mark.anyio
@pytest.mark.parametrize("extension", [".docm", ".docx", ".doc", ".wps"])
async def test_macro_document_extracts_metadata_without_running_office(
    monkeypatch, extension
):
    from app.modules.quality.service import document_catalog_md as conversion
    from app.modules.quality.service.document_catalog_revision import prepare_revision

    def forbidden(*args, **kwargs):
        pytest.fail("DOCM parsing must not execute Office or macros")

    monkeypatch.setattr(conversion, "convert_legacy_to_docx", forbidden)
    revision = await prepare_revision("wrong-99" + extension, _docm())
    assert revision.code == "SMP-XF2-001/04"
    assert revision.effective_date == date(2026, 9, 1)
    regular, _ = conversion.convert_word_attachment("same.docx", _docx())
    assert revision.prepared[0].split("\n---", 1)[1] == regular.split("\n---", 1)[1]


@pytest.mark.anyio
async def test_text_fallback_reports_conversion_failure_not_missing_number(monkeypatch):
    from app.modules.quality.service import document_catalog_revision as revision

    monkeypatch.setattr(
        revision,
        "convert_word_attachment",
        lambda *a: ("# 文本提取模式（无表格，来自 catdoc）\n正文", []),
    )
    with pytest.raises(AppException, match="转换.*页眉.*表格"):
        await revision.prepare_revision("test.doc", b"old-format")


def test_bilingual_header_vertical_code_and_history():
    doc = Document()
    table = doc.sections[0].header.add_table(rows=2, cols=3, width=6000000)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "文件名称\nFile Name"
    table.cell(0, 1).merge(table.cell(1, 1)).text = "虫害控制管理程序\nPest control"
    table.cell(0, 2).text = "文件编码\nFile Number"
    table.cell(1, 2).text = "SMP-XF2-001/04"
    date_table = doc.add_table(rows=1, cols=2)
    date_table.cell(0, 0).text = "生效日期\nEffective Date"
    date_table.cell(0, 1).text = "2026年10月01日"
    doc.add_paragraph("修订简历")
    history = doc.add_table(rows=1, cols=2)
    history.cell(0, 0).text = "文件编码"
    history.cell(0, 1).text = "SMP-XF2-001/03"
    out = io.BytesIO()
    doc.save(out)
    md, _ = converter.convert_docx_content_to_md(out.getvalue(), "wrong-99.docx")
    head = md.split("\n---", 1)[0]
    assert head.startswith("# 虫害控制管理程序\n")
    assert "**文件编号**: SMP-XF2-001/04" in head
    assert "**生效日期**: 2026年10月01日" in head
    assert "/03" not in head


def test_header_textbox_and_truncated_english_label():
    from docx.oxml import OxmlElement

    from app.modules.quality.service.document_catalog_metadata import (
        docx_metadata_text,
        extract_field,
    )

    doc = Document()
    header = doc.sections[0].header
    table = header.add_table(rows=2, cols=1, width=6000000)
    table.cell(0, 0).text = "文件编码\nFile Numbe"
    table.cell(1, 0).text = "STP-QS-BO-301/05"
    box = OxmlElement("w:txbxContent")
    box.append(table._tbl)
    header.paragraphs[0]._p.append(box)
    before = header._element.xml
    assert extract_field(docx_metadata_text(doc), "code") == "STP-QS-BO-301/05"
    assert header._element.xml == before


@pytest.mark.parametrize("translation", ["File Nμmber", "ile Number", "File Numbe"])
def test_chinese_label_survives_damaged_english_translation(translation):
    from app.modules.quality.service.document_catalog_metadata import (
        docx_metadata_text,
        extract_field,
    )

    doc = Document()
    table = doc.sections[0].header.add_table(rows=2, cols=1, width=6000000)
    table.cell(0, 0).text = "文件编码\n" + translation
    table.cell(1, 0).text = "SOP-QC(LN)-201/04"
    dates = doc.add_table(rows=1, cols=2)
    dates.cell(0, 0).text = "生效日期\nEective Date"
    dates.cell(0, 1).text = "2023年12月01日"
    metadata = docx_metadata_text(doc)
    assert extract_field(metadata, "code") == "SOP-QC(LN)-201/04"
    assert extract_field(metadata, "date") == "2023年12月01日"


def test_later_section_cannot_supply_first_page_identity():
    from docx.enum.section import WD_SECTION_START

    from app.modules.quality.service.document_catalog_metadata import (
        docx_metadata_text,
        extract_field,
    )

    doc = Document()
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.header.is_linked_to_previous = False
    section.header.paragraphs[
        0
    ].text = "文件编码: SMP-QA-018/07\n生效日期: 2024年04月01日"
    metadata = docx_metadata_text(doc)
    assert extract_field(metadata, "code") == ""
    assert extract_field(metadata, "date") == ""


def test_process_procedure_cover_metadata():
    doc = Document()
    doc.add_paragraph("霉酚酸菌种工艺规程")
    doc.add_paragraph("文件编码：KP-SC-MC-001/09")
    doc.add_paragraph("生效日期：2026年03月04日")
    doc.add_paragraph("目 录")
    doc.add_paragraph("文件编码：KP-SC-MC-001/08")
    out = io.BytesIO()
    doc.save(out)
    md, _ = converter.convert_docx_content_to_md(out.getvalue(), "wrong.docx")
    head = md.split("\n---", 1)[0]
    assert head.startswith("# 霉酚酸菌种工艺规程\n")
    assert "KP-SC-MC-001/09" in head
    assert "2026年03月04日" in head


def test_page_break_stops_metadata_and_generic_xml_does_not_crash():
    from lxml import etree

    from app.modules.quality.service.document_catalog_metadata import (
        docx_metadata_text,
        extract_field,
    )

    doc = Document()
    doc.element.body.insert(0, etree.Element("{urn:custom-document}metadata"))
    doc.add_paragraph("文件编码: KP-SC-MC-001/09")
    doc.add_paragraph("生效日期: 2026年03月04日")
    doc.add_page_break()
    doc.add_paragraph("文件编码: KP-SC-MC-001/08")
    doc.add_paragraph("生效日期: 2025年01月01日")
    text = docx_metadata_text(doc)
    assert extract_field(text, "code") == "KP-SC-MC-001/09"
    assert extract_field(text, "date") == "2026年03月04日"


def test_metadata_changes_do_not_change_body_tables_or_images(monkeypatch):
    from app.modules.quality.service import document_catalog_metadata as metadata
    from tests.modules.quality.test_document_catalog_md_rules import (
        _build_template_docx,
    )

    content = _build_template_docx()
    actual, images = converter.convert_docx_content_to_md(content, "sample.docx")
    monkeypatch.setattr(metadata, "docx_metadata_text", lambda _: "")
    without_metadata, other_images = converter.convert_docx_content_to_md(
        content, "sample.docx"
    )
    assert actual.split("\n---", 1)[1] == without_metadata.split("\n---", 1)[1]
    assert [(image.name, image.data, image.content_type) for image in images] == [
        (image.name, image.data, image.content_type) for image in other_images
    ]


def test_conflicting_cover_and_header_codes_are_not_selected():
    doc = Document()
    doc.add_paragraph("文件编号: STP-QS-MC-105/03")
    doc.add_paragraph("生效日期: 2026-01-01")
    doc.sections[0].header.paragraphs[0].text = "文件编号: STP-QS-MC-105/04"
    out = io.BytesIO()
    doc.save(out)
    md, _ = converter.convert_docx_content_to_md(out.getvalue(), "wrong-99.docx")
    assert "**文件编号**" not in md.split("\n---", 1)[0]


def test_numbered_body_table_does_not_supply_historical_metadata():
    from app.modules.quality.service.document_catalog_metadata import (
        docx_metadata_text,
        extract_field,
    )

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "文件编码: SOP-XF2-001/06"
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "1."
    table.cell(0, 1).text = "1."
    table.cell(0, 2).text = "范围"
    table.cell(1, 0).text = "文件编码"
    table.cell(1, 1).text = "SOP-XF2-001/01"
    assert extract_field(docx_metadata_text(doc), "code") == "SOP-XF2-001/06"


def test_code_with_typesetting_spaces():
    from app.modules.quality.service.document_catalog_metadata import (
        docx_metadata_text,
        extract_field,
    )

    doc = Document()
    table = doc.sections[0].header.add_table(rows=2, cols=1, width=6000000)
    table.cell(0, 0).text = "文件编码\nFile Number"
    table.cell(1, 0).text = "SOP- XT2-104 / 01"
    assert extract_field(docx_metadata_text(doc), "code") == "SOP-XT2-104/01"
    dates = doc.add_table(rows=1, cols=2)
    dates.cell(0, 0).text = "生效日期\nEffective Date"
    dates.cell(0, 1).text = "2024 年 0 3月 01 日"
    assert extract_field(docx_metadata_text(doc), "date") == "2024 年 03月 01 日"


def test_cover_and_body_in_one_table_excludes_history():
    from app.modules.quality.service.document_catalog_metadata import (
        docx_metadata_text,
        extract_field,
    )

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "文件编码: SMP-PM-013/04"
    table = doc.add_table(rows=4, cols=2)
    for row, values in zip(
        table.rows,
        [
            ("生效日期", "2023年05月01日"),
            ("分发-1", "质量保证部"),
            ("1.", "范围"),
            ("文件编码", "SMP-PM-013/01"),
        ],
        strict=True,
    ):
        row.cells[0].text, row.cells[1].text = values
    text = docx_metadata_text(doc)
    assert extract_field(text, "code") == "SMP-PM-013/04"
    assert extract_field(text, "date") == "2023年05月01日"


def test_first_page_header_is_not_confused_with_later_page_header():
    from app.modules.quality.service.document_catalog_metadata import (
        docx_metadata_text,
        extract_field,
    )

    doc = Document()
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.header.paragraphs[0].text = "文件编码: SOP-PM-104/01"
    section.first_page_header.paragraphs[0].text = "文件编码: SOP-PM-104/05"
    assert extract_field(docx_metadata_text(doc), "code") == "SOP-PM-104/05"


def test_inserted_revision_digits_are_read_without_changing_source():
    from docx.oxml import OxmlElement

    from app.modules.quality.service.document_catalog_metadata import (
        docx_metadata_text,
        extract_field,
    )

    doc = Document()
    paragraph = doc.add_paragraph("文件编号: QP-SC-001/0")
    for tag, text_tag, value in [("del", "delText", "2"), ("ins", "t", "3")]:
        change = OxmlElement(f"w:{tag}")
        run = OxmlElement("w:r")
        text = OxmlElement(f"w:{text_tag}")
        text.text = value
        run.append(text)
        change.append(run)
        paragraph._p.append(change)
    before = doc.element.xml
    assert extract_field(docx_metadata_text(doc), "code") == "QP-SC-001/03"
    assert doc.element.xml == before


@pytest.fixture
def state(monkeypatch, tmp_path):
    monkeypatch.setattr(storage, "_local_upload_dir", lambda: tmp_path)
    monkeypatch.setattr(storage, "minio_enabled", lambda: False)
    entry = DocumentEntry(
        id=uuid4(),
        department_id=uuid4(),
        name="虫害控制管理程序",
        code="SMP-XF2-001/03",
        effective_date=date(2024, 9, 26),
        effective_date_text="旧日期",
        attachments=[
            {
                "storage_key": "old.docx",
                "converted_md_key": "old.md",
                "asset_keys": ["old.png"],
                "file_name": "old.docx",
            }
        ],
    )
    for key in ["old.docx", "old.md", "old.png"]:
        storage._store_file(key, b"old", "application/octet-stream")
    db = AsyncSession()
    db.execute = AsyncMock(
        return_value=SimpleNamespace(scalar_one_or_none=lambda: entry)
    )
    db.flush = AsyncMock()
    return db, entry, tmp_path


@pytest.mark.anyio
async def test_upgrade_replaces_metadata_and_cleans_old_objects_after_commit(state):
    from app.modules.quality.service.document_catalog_revision import replace_attachment

    db, entry, root = state
    attachment, info = await replace_attachment(
        db, entry, "SMP-XF2-001-99.docx", _docx(), "application/octet-stream"
    )
    assert info.new_code == entry.code == "SMP-XF2-001/04"
    assert entry.effective_date == date(2026, 9, 1)
    assert entry.effective_date_text is None
    assert entry.attachments == [attachment]
    assert (root / "old.docx").exists()
    db.sync_session.dispatch.after_commit(db.sync_session)
    assert not any((root / key).exists() for key in ["old.docx", "old.md", "old.png"])
    assert storage.read_attachment_preview(entry, attachment["storage_key"])[0]
    await db.close()


@pytest.mark.anyio
@pytest.mark.parametrize(
    "code,effective,reason",
    [
        ("SMP-XF2-001/03", "2026-09-01", "高于"),
        ("SMP-XF2-001/02", "2026-09-01", "高于"),
        ("SMP-XF2-002/04", "2026-09-01", "编号"),
        ("SMP-XF2-001/04", "", "生效日期"),
        ("SMP-XF2-001/04", "2026-02-30", "生效日期"),
        ("", "2026-09-01", "编号"),
    ],
)
async def test_invalid_or_not_newer_leaves_everything_unchanged(
    state, code, effective, reason
):
    from app.modules.quality.service.document_catalog_revision import replace_attachment

    db, entry, root = state
    before = list(entry.attachments)
    with pytest.raises(AppException, match=reason):
        await replace_attachment(
            db,
            entry,
            "SMP-XF2-001-99.docx",
            _docx(code, effective),
            "application/octet-stream",
        )
    assert entry.code == "SMP-XF2-001/03"
    assert entry.effective_date == date(2024, 9, 26)
    assert entry.attachments == before
    assert len(list(root.iterdir())) == 3
    await db.close()


@pytest.mark.anyio
async def test_rollback_preserves_old_objects_and_removes_new_objects(state):
    from app.modules.quality.service.document_catalog_revision import replace_attachment

    db, entry, root = state
    await replace_attachment(db, entry, "new.docx", _docx(), "application/octet-stream")
    db.sync_session.dispatch.after_rollback(db.sync_session)
    assert sorted(path.name for path in root.rglob("*") if path.is_file()) == [
        "old.docx",
        "old.md",
        "old.png",
    ]
    db.sync_session.dispatch.after_commit(db.sync_session)
    assert (root / "old.docx").exists()
    await db.close()


@pytest.mark.anyio
async def test_flush_failure_preserves_original_entry(state):
    from app.modules.quality.service.document_catalog_revision import replace_attachment

    db, entry, root = state
    before = list(entry.attachments)
    db.flush.side_effect = RuntimeError("write failed")
    with pytest.raises(RuntimeError, match="write failed"):
        await replace_attachment(
            db, entry, "new.docx", _docx(), "application/octet-stream"
        )
    assert entry.code == "SMP-XF2-001/03"
    assert entry.effective_date == date(2024, 9, 26)
    assert entry.attachments == before
    assert len([p for p in root.rglob("*") if p.is_file()]) == 3
    await db.close()


@pytest.mark.anyio
@pytest.mark.parametrize("route", ["batch", "auto", "entry"])
@pytest.mark.parametrize("extension", [".docx", ".docm"])
async def test_http_upload_paths_use_body_and_reject_repeat(
    state, monkeypatch, route, extension
):
    from fastapi import FastAPI
    from httpx import ASGITransport, AsyncClient

    from app.core.database import get_db
    from app.modules.quality.api import document_catalog as api
    from app.platform.identity.deps import get_current_user

    db, entry, _ = state
    result = SimpleNamespace(
        scalar_one_or_none=lambda: entry,
        scalar_one=lambda: entry,
        scalars=lambda: SimpleNamespace(all=lambda: [entry]),
    )
    db.execute.return_value = result
    monkeypatch.setattr(
        api, "_resolve_quality_list_scope", AsyncMock(return_value=None)
    )
    monkeypatch.setattr(api, "_visible_entry", AsyncMock(return_value=entry))
    app = FastAPI()
    app.include_router(api.router, prefix="/quality")
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=uuid4())
    url = {
        "batch": "/quality/document-catalog/attachments/import",
        "auto": "/quality/document-entries/attachments/auto-bind",
        "entry": f"/quality/document-entries/{entry.id}/attachments",
    }[route]
    field = "files" if route == "batch" else "file"
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        uploads = [
            (
                field,
                (
                    "SMP-OTHER-999-99" + extension,
                    _docm() if extension == ".docm" else _docx(),
                ),
            )
        ]
        if route == "batch":
            uploads.append((field, ("missing.md", b"# no metadata")))
        response = await client.post(url, files=uploads)
        assert response.status_code == 200, response.text
        if route == "batch":
            data = response.json()["data"]
            assert data["bound"] == 1 and data["failed"] == 1
            assert "编号" in data["results"][1]["reason"]
        assert entry.code == "SMP-XF2-001/04"
        assert entry.effective_date == date(2026, 9, 1)
        assert len(entry.attachments) == 1
        original = list(entry.attachments)
        repeat = await client.post(url, files={field: ("SMP-XF2-001-99.docx", _docx())})
        assert entry.attachments == original
        if route == "batch":
            assert repeat.status_code == 200
            data = repeat.json()["data"]
            assert data["bound"] == 0 and data["failed"] == 1
            assert "未高于" in data["results"][0]["reason"]
        else:
            assert repeat.status_code == 409
    await db.close()


@pytest.mark.anyio
async def test_metadata_missing_ambiguous_and_pdf(monkeypatch):
    from app.modules.quality.service import document_catalog_revision as revision

    for text in [
        "**文件编号**: SMP-XF2-001/04\n批准日期: 2026-09-01",
        "**文件编号**: SMP-XF2-001/04\n生效日期: 2026-09-01\n生效日期: 2025-01-01",
        "文件编号: SMP-XF2-001/04\n文件编号: SMP-XF2-002/04\n生效日期: 2026-09-01",
    ]:
        with pytest.raises(AppException):
            await revision.prepare_revision("SMP-XF2-001-99.md", text.encode())
    with pytest.raises(AppException, match="无法可靠"):
        await revision.prepare_revision("SMP-XF2-001-99.png", b"image")
    monkeypatch.setattr(
        revision,
        "_pdf_text",
        lambda _: "文件编号: SMP-XF2-001/04\n生效日期: 2026-09-01",
    )
    info = await revision.prepare_revision("wrong.pdf", b"pdf")
    assert info.code == "SMP-XF2-001/04"
    assert info.effective_date == date(2026, 9, 1)


@pytest.mark.anyio
async def test_matching_requires_unique_subject_number(state):
    from app.modules.quality.service import document_catalog_revision as revision

    db, entry, _ = state
    info = await revision.prepare_revision("new.docx", _docx())
    other = DocumentEntry(code="SMP-XF2-002/04", name=entry.name)
    rows = [entry, other]
    db.execute.return_value = SimpleNamespace(
        scalars=lambda: SimpleNamespace(all=lambda: rows)
    )
    assert await revision.find_revision_entry(db, info, None) is entry
    rows.append(DocumentEntry(code=entry.code, name=entry.name))
    assert await revision.find_revision_entry(db, info, None) is None
    rows[:] = [other]
    assert await revision.find_revision_entry(db, info, None) is None
    await db.close()


@pytest.mark.anyio
async def test_comparison_uses_locked_latest_revision(state):
    from app.modules.quality.service.document_catalog_revision import replace_attachment

    db, entry, root = state
    # 模拟等待锁期间另一次上传已把目录升级至 10，不能按调用前的 03 覆盖。
    latest = DocumentEntry(
        id=entry.id,
        code="SMP-XF2-001/10",
        attachments=entry.attachments,
        effective_date=date(2026, 9, 20),
    )
    db.execute.return_value = SimpleNamespace(scalar_one_or_none=lambda: latest)
    with pytest.raises(AppException, match="未高于"):
        await replace_attachment(
            db, entry, "new.docx", _docx(), "application/octet-stream"
        )
    assert "FOR UPDATE" in str(db.execute.call_args.args[0])
    assert latest.code == "SMP-XF2-001/10"
    assert len(list(root.iterdir())) == 3
    await db.close()


@pytest.mark.anyio
async def test_numeric_revision_nine_to_ten(state):
    from app.modules.quality.service.document_catalog_revision import replace_attachment

    db, entry, _ = state
    entry.code = "SMP-XF2-001/09"
    await replace_attachment(
        db, entry, "wrong.docx", _docx("SMP-XF2-001/10"), "application/octet-stream"
    )
    assert entry.code == "SMP-XF2-001/10"
    await db.close()


@pytest.mark.anyio
async def test_pdf_reads_labeled_text_and_rejects_unreadable_content():
    import pymupdf

    from app.modules.quality.service.document_catalog_revision import prepare_revision

    with pymupdf.open() as doc:
        page = doc.new_page()
        page.insert_text(
            (40, 40), "Document No.: SMP-XF2-001/04\nEffective Date: 2026-09-01"
        )
        content = doc.tobytes()
    revision = await prepare_revision("SMP-XF2-001-99.pdf", content)
    assert revision.code == "SMP-XF2-001/04"
    assert revision.effective_date == date(2026, 9, 1)
    with pytest.raises(AppException, match="解析失败"):
        await prepare_revision("broken.pdf", b"%PDF-1.7 broken")
