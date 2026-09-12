"""OOT 限度告知单 docx 导入/导出测试。"""

from __future__ import annotations

import json
import zipfile
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
from uuid import uuid4

import pytest
from docx import Document

from app.core.exceptions import AppException
from app.modules.quality.api import oot_limit as api
from app.modules.quality.models.oot_limit import OotLimitProduct
from app.modules.quality.service import oot_limit_docx as docx_service


class _Result:
    def __init__(
        self,
        row: object | None = None,
        rows: list[object] | None = None,
        total: int = 0,
    ) -> None:
        self.row = row
        self.rows = rows or []
        self.total = total

    def scalar(self) -> int:
        return self.total

    def scalar_one_or_none(self) -> object | None:
        return self.row

    def scalar_one(self) -> object:
        assert self.row is not None
        return self.row

    def scalars(self) -> _Result:
        return self

    def all(self) -> list[object]:
        return self.rows


class _FakeDb:
    def __init__(self, results: list[_Result] | None = None) -> None:
        self._results = list(results or [])
        self.added: list[object] = []
        self.commits = 0
        self.executed_stmts: list[object] = []

    def add(self, obj: object) -> None:
        self.added.append(obj)

    async def flush(self) -> None:
        for obj in self.added:
            if getattr(obj, "id", None) is None:
                obj.id = uuid4()

    async def commit(self) -> None:
        self.commits += 1

    async def rollback(self) -> None:
        pass

    async def execute(self, stmt: object) -> _Result:
        self.executed_stmts.append(stmt)
        return self._results.pop(0)


def _build_simple_notice(
    title: str = "2026年 测试产品A 产品OOT限度通知单",
    rows: tuple[tuple[str, str, str], ...] = (
        ("比旋度", "+1°~+2°", "+1°~+1.5°"),
        ("炽灼残渣", "≤0.1%", "≤0.1%"),
    ),
    include_title: bool = True,
) -> bytes:
    doc = Document()
    if include_title:
        doc.add_paragraph(title)
    doc.add_table(rows=2, cols=6)
    table = doc.add_table(rows=1 + len(rows), cols=3)
    header = table.rows[0].cells
    header[0].text = "项目"
    header[1].text = "标准"
    header[2].text = "OOT限度"
    for index, (name, standard, oot) in enumerate(rows, start=1):
        cells = table.rows[index].cells
        cells[0].text = name
        cells[1].text = standard
        cells[2].text = oot
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_grouped_notice() -> bytes:
    doc = Document()
    doc.add_paragraph("2026年 测试产品B 产品OOT限度通知单")
    doc.add_table(rows=2, cols=6)
    table = doc.add_table(rows=5, cols=5)
    header = table.rows[0]
    merged_name = header.cells[0].merge(header.cells[1])
    merged_name.text = "项目"
    merged_standard = header.cells[2].merge(header.cells[3])
    merged_standard.text = "标准"
    header.cells[4].text = "OOT限度"

    data = [
        ("", "比旋度（按无水计）", "+55°~+65°", "+58°~+62°"),
        ("有关物质", "杂质A", "≤1.0%", "0.2%~0.8%"),
        ("有关物质", "杂质B", "≤0.5%", "≤0.3%"),
        ("", "含量（按无水计）", "95.0%~102.0%", "95.0%~99.8%"),
    ]
    for row_index, (group, name, standard, oot) in enumerate(data, start=1):
        row = table.rows[row_index]
        if group:
            row.cells[0].text = group
            row.cells[1].text = name
            merged = row.cells[2].merge(row.cells[3])
            merged.text = standard
        else:
            name_cell = row.cells[0].merge(row.cells[1])
            name_cell.text = name
            standard_cell = row.cells[2].merge(row.cells[3])
            standard_cell.text = standard
        row.cells[4].text = oot
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _product(product_id=None, **overrides) -> OotLimitProduct:
    return OotLimitProduct(
        id=product_id or uuid4(),
        product_code=overrides.get("product_code", "OOT-2026-01"),
        product_name=overrides.get("product_name", "测试产品A"),
        document_title=overrides.get(
            "document_title", "2026年 测试产品A 产品OOT限度通知单"
        ),
        document_year=overrides.get("document_year", 2026),
        source_file_name=overrides.get("source_file_name", "告知单.docx"),
        is_active=True,
        is_deleted=False,
    )


# ============ 解析 ============


def test_parse_simple_notice_docx() -> None:
    content = _build_simple_notice()
    parsed = docx_service.parse_notice_docx(content, "2026年 LV OOT限度告知单.docx")
    assert parsed.product_name == "测试产品A"
    assert parsed.document_year == 2026
    assert parsed.document_title == "2026年 测试产品A 产品OOT限度通知单"
    assert parsed.warnings == []
    assert [item.item_name for item in parsed.items] == ["比旋度", "炽灼残渣"]
    assert parsed.items[0].standard_value == "+1°~+2°"
    assert parsed.items[0].oot_limit_value == "+1°~+1.5°"
    # 标准与 OOT 限度相同的行两边都保留真实值
    assert parsed.items[1].standard_value == "≤0.1%"
    assert parsed.items[1].oot_limit_value == "≤0.1%"
    assert [item.display_order for item in parsed.items] == [1, 2]
    assert all(item.item_group is None for item in parsed.items)


def test_parse_grouped_notice_docx() -> None:
    content = _build_grouped_notice()
    parsed = docx_service.parse_notice_docx(content, "多拉菌素 OOT限度告知单.docx")
    assert parsed.product_name == "测试产品B"
    assert [(item.item_group, item.item_name) for item in parsed.items] == [
        (None, "比旋度（按无水计）"),
        ("有关物质", "杂质A"),
        ("有关物质", "杂质B"),
        (None, "含量（按无水计）"),
    ]
    assert parsed.items[1].standard_value == "≤1.0%"
    assert parsed.items[1].oot_limit_value == "0.2%~0.8%"


def test_parse_notice_missing_standard_reports_row_error() -> None:
    doc = Document()
    doc.add_paragraph("2026年 测试产品A 产品OOT限度通知单")
    table = doc.add_table(rows=3, cols=3)
    header = table.rows[0].cells
    header[0].text = "项目"
    header[1].text = "标准"
    header[2].text = "OOT限度"
    table.rows[1].cells[0].text = "干燥失重"
    table.rows[1].cells[2].text = "≤0.1%"
    row = table.rows[2]
    row.cells[0].text = "含量"
    row.cells[1].text = "98%"
    row.cells[2].text = "99%"
    buffer = BytesIO()
    doc.save(buffer)

    parsed = docx_service.parse_notice_docx(buffer.getvalue(), "告知单.docx")
    assert parsed.row_errors == ["第1行：缺少标准值"]
    assert [item.item_name for item in parsed.items] == ["含量"]


def test_parse_notice_without_title_falls_back_to_filename() -> None:
    content = _build_simple_notice(include_title=False)
    parsed = docx_service.parse_notice_docx(content, "2027年 XYZ OOT限度告知单.docx")
    assert parsed.product_name == "XYZ"
    assert parsed.document_year == 2027
    assert any("文件名" in warning for warning in parsed.warnings)


def test_parse_notice_without_limits_table_reports_error() -> None:
    doc = Document()
    doc.add_paragraph("2026年 测试产品A 产品OOT限度通知单")
    only_table = doc.add_table(rows=1, cols=2)
    only_table.rows[0].cells[0].text = "其他"
    only_table.rows[0].cells[1].text = "表"
    buffer = BytesIO()
    doc.save(buffer)

    parsed = docx_service.parse_notice_docx(buffer.getvalue(), "告知单.docx")
    assert parsed.items == []
    assert any("限度表" in error for error in parsed.row_errors)


def test_parse_notice_invalid_docx_raises() -> None:
    with pytest.raises(AppException) as exc_info:
        docx_service.parse_notice_docx(b"not a docx", "坏文件.docx")
    assert exc_info.value.status_code == 400


# ============ 导入（preview / confirm） ============


async def test_preview_marks_new_product_when_no_match() -> None:
    db = _FakeDb([_Result(None), _Result(None)])
    result = await docx_service.preview_oot_limit_notice_import(
        db, [("2026年 测试产品A OOT限度告知单.docx", _build_simple_notice())]
    )
    entry = result["files"][0]
    assert entry["status"] == "ok"
    assert entry["mode"] == "create"
    assert entry["matched_product_code"] is None
    assert entry["item_count"] == 2
    assert entry["items"][0]["item_name"] == "比旋度"


async def test_preview_marks_update_when_product_matched() -> None:
    product = _product()
    db = _FakeDb([_Result(product)])
    result = await docx_service.preview_oot_limit_notice_import(
        db, [("2026年 测试产品A OOT限度告知单.docx", _build_simple_notice())]
    )
    entry = result["files"][0]
    assert entry["mode"] == "update"
    assert entry["matched_product_code"] == "OOT-2026-01"


async def test_confirm_creates_product_with_generated_code() -> None:
    db = _FakeDb([_Result(None), _Result(None), _Result(rows=[])])
    result = await docx_service.confirm_oot_limit_notice_import(
        db, [("2026年 测试产品A OOT限度告知单.docx", _build_simple_notice())]
    )
    assert result["created_count"] == 1
    assert result["update_count"] == 0
    assert result["error_count"] == 0
    products = [obj for obj in db.added if isinstance(obj, OotLimitProduct)]
    assert len(products) == 1
    assert products[0].product_code == "OOT-2026-01"
    assert products[0].source_file_name == "2026年 测试产品A OOT限度告知单.docx"
    items = [obj for obj in db.added if not isinstance(obj, OotLimitProduct)]
    assert len(items) == 2
    assert items[0].specification == "+1°~+2°"
    assert items[0].oot_limit == "+1°~+1.5°"
    assert db.commits == 1


async def test_confirm_updates_product_and_replaces_items() -> None:
    product = _product()
    db = _FakeDb([_Result(product), _Result()])
    result = await docx_service.confirm_oot_limit_notice_import(
        db, [("2026年 测试产品A OOT限度告知单.docx", _build_simple_notice())]
    )
    assert result["created_count"] == 0
    assert result["update_count"] == 1
    assert product.source_file_name == "2026年 测试产品A OOT限度告知单.docx"
    # 旧明细通过 DELETE 语句物理删除（display_order 唯一约束覆盖软删行）
    assert any("DELETE" in str(stmt).upper() for stmt in db.executed_stmts)
    items = [obj for obj in db.added if not isinstance(obj, OotLimitProduct)]
    assert len(items) == 2
    assert db.commits == 1


async def test_confirm_reports_file_without_items() -> None:
    doc = Document()
    doc.add_paragraph("2026年 空文件 产品OOT限度通知单")
    table = doc.add_table(rows=1, cols=3)
    header = table.rows[0].cells
    header[0].text = "项目"
    header[1].text = "标准"
    header[2].text = "OOT限度"
    buffer = BytesIO()
    doc.save(buffer)
    db = _FakeDb()
    result = await docx_service.confirm_oot_limit_notice_import(
        db, [("空文件.docx", buffer.getvalue())]
    )
    assert result["created_count"] == 0
    assert result["error_count"] == 1
    assert "未解析到限度数据" in result["error_details"][0]["error"]


# ============ 导出 ============


def _export_product(**overrides) -> SimpleNamespace:
    return SimpleNamespace(
        id=uuid4(),
        product_code=overrides.get("product_code", "OOT-2026-01"),
        product_name=overrides.get("product_name", "测试产品A"),
        document_title=overrides.get(
            "document_title", "2026年 测试产品A 产品OOT限度通知单"
        ),
        document_year=overrides.get("document_year", 2026),
        source_file_name=overrides.get(
            "source_file_name", "2026年 测试产品A OOT限度告知单.docx"
        ),
    )


def _export_items() -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            item_group=None,
            item_name="比旋度",
            standard_value="+1°~+2°",
            oot_limit_value="+1°~+1.5°",
        ),
        SimpleNamespace(
            item_group=None,
            item_name="炽灼残渣",
            standard_value="≤0.1%",
            oot_limit_value="≤0.1%",
        ),
    ]


def test_export_simple_notice_keeps_original_layout() -> None:
    product = _export_product()
    data = docx_service.export_notice_docx(product, _export_items())
    document = Document(BytesIO(data))

    title = next(p.text for p in document.paragraphs if "OOT限度通知单" in p.text)
    assert "测试产品A" in title
    assert "2026" in title
    # 签批表保留（第一个表格）
    assert len(document.tables) == 2
    limits = document.tables[1]
    assert [cell.text for cell in limits.rows[0].cells[:3]] == [
        "项目",
        "标准",
        "OOT限度",
    ]
    assert [cell.text for cell in limits.rows[1].cells] == [
        "比旋度",
        "+1°~+2°",
        "+1°~+1.5°",
    ]
    assert limits.rows[2].cells[1].text == "≤0.1%"

    # 导出文件可被导入逻辑原样读回
    reparsed = docx_service.parse_notice_docx(data, product.source_file_name)
    assert [
        (i.item_name, i.standard_value, i.oot_limit_value) for i in reparsed.items
    ] == [
        ("比旋度", "+1°~+2°", "+1°~+1.5°"),
        ("炽灼残渣", "≤0.1%", "≤0.1%"),
    ]


def test_export_grouped_notice_roundtrip() -> None:
    product = _export_product(product_name="测试产品B")
    items = [
        SimpleNamespace(
            item_group=None,
            item_name="比旋度（按无水计）",
            standard_value="+55°~+65°",
            oot_limit_value="+58°~+62°",
        ),
        SimpleNamespace(
            item_group="有关物质",
            item_name="杂质A",
            standard_value="≤1.0%",
            oot_limit_value="0.2%~0.8%",
        ),
        SimpleNamespace(
            item_group="有关物质",
            item_name="杂质B",
            standard_value="≤0.5%",
            oot_limit_value="≤0.3%",
        ),
    ]
    data = docx_service.export_notice_docx(product, items)
    document = Document(BytesIO(data))
    title = next(p.text for p in document.paragraphs if "OOT限度通知单" in p.text)
    assert "测试产品B" in title

    reparsed = docx_service.parse_notice_docx(data, "导出.docx")
    assert [
        (i.item_group, i.item_name, i.standard_value, i.oot_limit_value)
        for i in reparsed.items
    ] == [
        (None, "比旋度（按无水计）", "+55°~+65°", "+58°~+62°"),
        ("有关物质", "杂质A", "≤1.0%", "0.2%~0.8%"),
        ("有关物质", "杂质B", "≤0.5%", "≤0.3%"),
    ]


def test_export_filename_prefers_source_file_name() -> None:
    assert (
        docx_service.notice_export_filename(_export_product())
        == "2026年 测试产品A OOT限度告知单.docx"
    )
    fallback = _export_product(source_file_name=None, document_title="2026年 X 通知单")
    assert docx_service.notice_export_filename(fallback) == "2026年 X 通知单.docx"


def test_build_notice_zip_entries() -> None:
    entry_a = ("a.docx", b"aaaa")
    entry_b = ("b.docx", b"bbbb")
    data = docx_service.build_notice_zip([entry_a, entry_b, entry_a])
    with zipfile.ZipFile(BytesIO(data)) as archive:
        names = archive.namelist()
    assert names == ["a.docx", "b.docx", "a(1).docx"]


# ============ API 端点 ============


def _upload(filename: str) -> SimpleNamespace:
    return SimpleNamespace(filename=filename)


async def test_import_endpoints_require_login() -> None:
    with pytest.raises(AppException) as exc_info:
        await api.preview_oot_limit_notice_import(
            files=[_upload("a.docx")], db=_FakeDb(), current_user=None
        )
    assert exc_info.value.status_code == 401
    with pytest.raises(AppException):
        await api.export_oot_limit_product(uuid4(), db=_FakeDb(), current_user=None)


async def test_import_endpoint_rejects_non_docx(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(api, "_require_user", Mock())
    monkeypatch.setattr(api, "_assert_quality_edit_scope", AsyncMock())
    with pytest.raises(AppException) as exc_info:
        await api.preview_oot_limit_notice_import(
            files=[_upload("说明.txt")],
            db=_FakeDb(),
            current_user=SimpleNamespace(id=uuid4()),
        )
    assert exc_info.value.status_code == 400


async def test_preview_endpoint_returns_files(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(api, "_require_user", Mock())
    monkeypatch.setattr(api, "_assert_quality_edit_scope", AsyncMock())

    async def _fake_read(file: object, max_bytes: int, what: str) -> bytes:
        return _build_simple_notice()

    monkeypatch.setattr(api, "read_upload_with_limit", _fake_read)
    db = _FakeDb([_Result(None), _Result(None)])
    response = await api.preview_oot_limit_notice_import(
        files=[_upload("2026年 测试产品A OOT限度告知单.docx")],
        db=db,
        current_user=SimpleNamespace(id=uuid4()),
    )
    body = json.loads(response.body)  # type: ignore[union-attr]
    entry = body["data"]["files"][0]
    assert entry["mode"] == "create"
    assert entry["item_count"] == 2


async def test_confirm_endpoint_reports_counts(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(api, "_require_user", Mock())
    monkeypatch.setattr(api, "_assert_quality_edit_scope", AsyncMock())

    async def _fake_read(file: object, max_bytes: int, what: str) -> bytes:
        return _build_simple_notice()

    monkeypatch.setattr(api, "read_upload_with_limit", _fake_read)
    db = _FakeDb([_Result(None), _Result(None), _Result(rows=[])])
    response = await api.confirm_oot_limit_notice_import(
        files=[_upload("2026年 测试产品A OOT限度告知单.docx")],
        db=db,
        current_user=SimpleNamespace(id=uuid4()),
    )
    body = json.loads(response.body)  # type: ignore[union-attr]
    assert body["data"]["created_count"] == 1
    assert db.commits == 1


async def test_export_product_endpoint_streams_docx(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(api, "_require_user", Mock())
    product = _product()
    items = _export_items()
    db = _FakeDb([_Result(product), _Result(rows=items)])
    response = await api.export_oot_limit_product(
        product.id, db=db, current_user=SimpleNamespace(id=uuid4())
    )
    assert response.media_type == docx_service.DOCX_MEDIA_TYPE
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    content = b"".join([chunk async for chunk in response.body_iterator])
    document = Document(BytesIO(content))
    assert any("测试产品A" in p.text for p in document.paragraphs)


async def test_export_product_endpoint_404(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(api, "_require_user", Mock())
    db = _FakeDb([_Result(None)])
    with pytest.raises(AppException) as exc_info:
        await api.export_oot_limit_product(
            uuid4(), db=db, current_user=SimpleNamespace(id=uuid4())
        )
    assert exc_info.value.status_code == 404


async def test_export_all_endpoint_returns_zip(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(api, "_require_user", Mock())
    product_a = _export_product()
    product_b = _export_product(product_name="测试产品B")
    db = _FakeDb(
        [
            _Result(rows=[product_a, product_b]),
            _Result(rows=_export_items()),
            _Result(rows=_export_items()),
        ]
    )
    response = await api.export_all_oot_limit_products(
        db=db, current_user=SimpleNamespace(id=uuid4())
    )
    assert response.media_type == docx_service.ZIP_MEDIA_TYPE
    content = b"".join([chunk async for chunk in response.body_iterator])
    with zipfile.ZipFile(BytesIO(content)) as archive:
        names = archive.namelist()
    assert len(names) == 2
    assert all(name.endswith(".docx") for name in names)


async def test_export_all_endpoint_404_when_empty(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(api, "_require_user", Mock())
    db = _FakeDb([_Result(rows=[])])
    with pytest.raises(AppException) as exc_info:
        await api.export_all_oot_limit_products(
            db=db, current_user=SimpleNamespace(id=uuid4())
        )
    assert exc_info.value.status_code == 404
