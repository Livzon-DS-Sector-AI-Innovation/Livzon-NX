from __future__ import annotations

import io
import re
from datetime import UTC, date, datetime
from zoneinfo import ZoneInfo

from docx import Document
from docx.oxml.ns import qn

from app.modules.quality.service import capa_ledger_export

_LOCAL_TIMEZONE = ZoneInfo("Asia/Shanghai")

_HEADERS = [
    "CAPA编号",
    "启动日期",
    "事件部门",
    "涉及产品",
    "来源编号",
    "CAPA简述",
    "CAPA效果评估",
    "关闭日期",
    "QA质量员/日期",
]


def _body_paragraph_texts(doc: Document) -> list[str]:
    return [
        "".join(t.text or "" for t in paragraph.iter(qn("w:t")))
        for paragraph in doc.element.body.findall(qn("w:p"))
    ]


def _perm_marker_count(doc: Document) -> tuple[int, int]:
    body = doc.element.body
    return (
        len(list(body.iter(qn("w:permStart")))),
        len(list(body.iter(qn("w:permEnd")))),
    )


def _export(items: list[dict], end_date: date | None = None) -> Document:
    docx_bytes = capa_ledger_export.generate_capa_ledger_export_docx(items, end_date)
    return Document(io.BytesIO(docx_bytes))


def test_export_uses_capa_ledger_template_layout() -> None:
    doc = _export(
        [
            {
                "capa_code": "CAPA-PC2607001",
                "created_at": datetime(2026, 9, 20, 1, 21, tzinfo=UTC),
                # 同步自飞书的启动日期：国内零点，存储为 UTC 16:00
                "expected_completion_date": datetime(2026, 7, 3, 16, 0, tzinfo=UTC),
                "department": "QC",
                "affected_product": "原料A",
                "source_code": "PC-2607001",
                "title": "1、复核空调系统参数\n2、补充培训",
                "evaluation_result": "有效",
                "closure_date": datetime(2026, 7, 20, tzinfo=UTC),
                "qa_confirmer": "杨小芹",
                "qa_confirm_date": datetime(2026, 7, 21, tzinfo=UTC),
            },
            {
                "capa_code": "CAPA-PC2607002",
                "created_at": "2026-07-05T08:00:00+00:00",
                "department": "QA",
                "title": "称量记录缺页",
                "evaluation_result": "进行中",
                "closure_date": None,
                "qa_confirmer": "进行中",
                "qa_confirm_date": None,
            },
        ]
    )

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3
    assert [table.cell(0, index).text for index in range(9)] == _HEADERS

    assert table.cell(1, 0).text == "CAPA-PC2607001"
    assert table.cell(1, 1).text == "2026.07.04"
    assert table.cell(1, 2).text == "QC"
    assert table.cell(1, 3).text == "原料A"
    assert table.cell(1, 4).text == "PC-2607001"
    assert table.cell(1, 5).text == "1、复核空调系统参数\n2、补充培训"
    assert table.cell(1, 6).text == "有效"
    assert table.cell(1, 7).text == "2026.07.20"
    assert table.cell(1, 8).text == "杨小芹2026.07.21"

    # 编号是台账编号（不再带序号列），第二行日期与编号列按存储值导出
    assert table.cell(2, 0).text == "CAPA-PC2607002"
    assert table.cell(2, 1).text == "2026.07.05"
    assert table.cell(2, 4).text == ""


def test_start_date_uses_synced_start_date_then_created_at() -> None:
    doc = _export(
        [
            {
                "capa_code": "SYNCED",
                # 飞书启动日期 2026-05-12（国内零点 -> UTC 16:00）
                "expected_completion_date": datetime(2026, 5, 11, 16, 0, tzinfo=UTC),
                "created_at": datetime(2026, 9, 20, 1, 21, tzinfo=UTC),
            },
            {
                "capa_code": "LOCAL-DRAFT",
                "expected_completion_date": None,
                "created_at": datetime(2026, 9, 20, 1, 21, tzinfo=UTC),
            },
        ]
    )

    table = doc.tables[0]
    assert table.cell(1, 1).text == "2026.05.12"
    assert table.cell(2, 1).text == "2026.09.20"


def test_multiline_summary_reuses_template_paragraphs() -> None:
    doc = _export(
        [{"capa_code": "CAPA-PC2607001", "title": "多行一\n多行二\n多行三\n多行四"}]
    )
    cell = doc.tables[0].cell(1, 5)
    assert [paragraph.text for paragraph in cell.paragraphs] == [
        "多行一",
        "多行二",
        "多行三",
        "多行四",
    ]
    # 自动编号已清除，避免复制的行跨行连续编号
    for paragraph in cell.paragraphs:
        ppr = paragraph._p.find(qn("w:pPr"))
        assert ppr is None or ppr.find(qn("w:numPr")) is None


def test_closure_column_keeps_ledger_in_progress_wording() -> None:
    doc = _export(
        [
            {"capa_code": "A", "evaluation_result": "进行中", "closure_date": None},
            {"capa_code": "B", "evaluation_result": None, "closure_date": None},
            {
                "capa_code": "C",
                "evaluation_result": "进行中",
                "closure_date": datetime(2026, 8, 1, tzinfo=UTC),
            },
        ]
    )
    table = doc.tables[0]
    assert table.cell(1, 7).text == "进行中"
    assert table.cell(2, 7).text == ""
    assert table.cell(3, 7).text == "2026.08.01"


def test_cover_date_ends_on_export_date_and_keeps_template_start() -> None:
    doc = _export([{"capa_code": "A"}], date(2026, 9, 23))

    date_paragraphs = [
        text for text in _body_paragraph_texts(doc) if "－" in text and "年" in text
    ]
    assert date_paragraphs == ["   2026年 01 月 01 日－ 2026 年 09 月 23 日"]


def test_cover_date_defaults_to_local_export_day() -> None:
    before = datetime.now(_LOCAL_TIMEZONE).date()
    doc = _export([{"capa_code": "A"}])
    after = datetime.now(_LOCAL_TIMEZONE).date()

    text = next(
        text for text in _body_paragraph_texts(doc) if "－" in text and "年" in text
    )
    match = re.search(r"－\s*(\d{4})\s*年\s*(\d{2})\s*月\s*(\d{2})\s*日", text)
    assert match is not None
    exported = date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    assert before <= exported <= after


def test_export_keeps_ledger_title_and_single_editable_range() -> None:
    doc = _export([{"capa_code": "A"}, {"capa_code": "B"}])

    paragraphs = _body_paragraph_texts(doc)
    assert paragraphs.count("CAPA登记汇总表") == 2
    assert all("偏差" not in text for text in paragraphs)
    # 模板表头的可编辑区域不能被逐行复制放大
    assert _perm_marker_count(doc) == (1, 1)
