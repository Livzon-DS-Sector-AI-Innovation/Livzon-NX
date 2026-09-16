"""维保汇总导出：模板复刻填充、YYYY.MM.DD 文本日期、行数调整与文件名。"""

from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO
from typing import Any

import pytest
from openpyxl import load_workbook

from app.core.exceptions import AppException
from app.modules.quality.service import maintenance_export as exporter

pytestmark = pytest.mark.anyio

_TZ_HOURS = 8


def _ms(target: date) -> int:
    from datetime import datetime, timezone

    tz = timezone(timedelta(hours=_TZ_HOURS))
    stamp = datetime(target.year, target.month, target.day, tzinfo=tz).timestamp()
    return int(stamp * 1000)


def _install_mirror(
    monkeypatch: pytest.MonkeyPatch, items: list[dict[str, Any]]
) -> None:
    from app.modules.quality.service import maintenance_export as mod

    async def fake_list(db, entity_code, *, page=1, page_size=500, **kwargs):
        return {
            "items": items,
            "total": len(items),
            "page": 1,
            "page_size": page_size,
            "configured": True,
            "fields": [],
            "last_sync_time": None,
        }

    monkeypatch.setattr(mod, "list_instrument_mirror", fake_list)


def _row(day: date, code: str, name: str, **extra: Any) -> dict[str, Any]:
    row = {
        "record_id": f"rec-{code}-{day.isoformat()}",
        "设备名称": name,
        "设备编号": code,
        "安装位置": "灭菌室2",
        "维护内容": "清洗内桶、内壁、硅胶密封圈。",
        "维护日期": str(_ms(day)),
        "维护人": [{"name": "贾慧"}],
        "复核人": [{"name": "马晓慧"}],
        "维保类型": "1",
        "维备注": extra.get("维备注"),
    }
    row.update({k: v for k, v in extra.items() if k != "维备注"})
    return row


async def test_export_renders_template_with_dot_dates(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _install_mirror(
        monkeypatch,
        [
            _row(date(2026, 7, 29), "QC-2-2-086", "手提式高压蒸汽灭菌器"),
            _row(date(2026, 7, 12), "QC-2-2-086", "手提式高压蒸汽灭菌器"),
            _row(date(2026, 8, 1), "QC-2-2-086", "手提式高压蒸汽灭菌器"),  # 非当月
        ],
    )

    content, filename = await exporter.export_maintenance_summary(None, "2026-07")

    assert filename == "2026年7月  设备年度预防性维护保养汇总表.xlsx"
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    # 标题与表头保持模板
    assert sheet.cell(1, 1).value == "设备年度预防性维护保养汇总表"
    assert sheet.cell(2, 1).value == "序号"
    assert sheet.cell(2, 2).value == "设备编号"  # 模板表头原样（与数据列对调）
    assert sheet.cell(2, 3).value == "设备名称"
    # 数据按维护日期升序，仅含当月
    assert sheet.cell(3, 1).value == 1
    assert sheet.cell(3, 6).value == "2026.07.12"
    assert sheet.cell(4, 1).value == 2
    assert sheet.cell(4, 6).value == "2026.07.29"
    # 列序复刻手工表：名称在编号前
    assert sheet.cell(3, 2).value == "手提式高压蒸汽灭菌器"
    assert sheet.cell(3, 3).value == "QC-2-2-086"
    assert sheet.cell(3, 7).value == "贾慧"
    assert sheet.cell(3, 10).value == "-"  # 备注空值占位
    # 末尾备注行紧随数据区（第 5 行），合并至 J 列且保留说明文字
    footer_text = str(sheet.cell(5, 1).value or "")
    assert "预防性维护" in footer_text
    merged = {str(m) for m in sheet.merged_cells.ranges}
    assert "A5:J5" in merged
    assert "A1:J1" in merged
    # 数据行样式来自模板示例行（有边框）
    assert sheet.cell(3, 2).border.left.style is not None
    # 日期为文本（非数字/日期类型）
    assert isinstance(sheet.cell(3, 6).value, str)


async def test_export_remark_prefers_beizhu_and_falls_back(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """备注列兼容两种列名：优先「备注」，缺省回退「维备注」。"""
    _install_mirror(
        monkeypatch,
        [
            _row(date(2026, 7, 1), "QC-2-2-001", "甲", 备注="新列名备注内容"),
            _row(date(2026, 7, 2), "QC-2-2-002", "乙", 维备注="旧列名备注内容"),
            _row(date(2026, 7, 3), "QC-2-2-003", "丙"),
        ],
    )

    content, _ = await exporter.export_maintenance_summary(None, "2026-07")
    sheet = load_workbook(BytesIO(content)).active
    assert sheet.cell(3, 10).value == "新列名备注内容"
    assert sheet.cell(4, 10).value == "旧列名备注内容"
    assert sheet.cell(5, 10).value == "-"


async def test_export_more_rows_than_template(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    days = [date(2026, 7, d) for d in (2, 6, 10, 14, 18)]
    _install_mirror(
        monkeypatch,
        [
            _row(day, f"QC-2-2-0{index}", f"设备{index}")
            for index, day in enumerate(days, 1)
        ],
    )

    content, _ = await exporter.export_maintenance_summary(None, "2026-07")

    sheet = load_workbook(BytesIO(content)).active
    for offset, day in enumerate(days):
        excel_row = 3 + offset
        assert sheet.cell(excel_row, 6).value == day.strftime("%Y.%m").replace(
            "-", "."
        ) + f".{day.day:02d}"
    # 备注行在第 8 行（3+5）
    merged = {str(m) for m in sheet.merged_cells.ranges}
    assert "A8:J8" in merged
    assert "预防性维护" in str(sheet.cell(8, 1).value or "")
    # 行高沿用模板数据行（34）
    assert sheet.row_dimensions[3].height == 34
    assert sheet.row_dimensions[7].height == 34
    assert sheet.row_dimensions[8].height == 47


async def test_export_empty_month_keeps_header_and_footer(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _install_mirror(monkeypatch, [_row(date(2026, 8, 1), "QC-2-2-086", "设备")])

    content, filename = await exporter.export_maintenance_summary(None, "2026-07")

    assert filename.startswith("2026年7月")
    sheet = load_workbook(BytesIO(content)).active
    # 无数据：表头（第2行）后直接备注行（第3行）
    assert sheet.cell(2, 1).value == "序号"
    assert "预防性维护" in str(sheet.cell(3, 1).value or "")
    assert "A3:J3" in {str(m) for m in sheet.merged_cells.ranges}
    assert sheet.cell(4, 1).value is None


async def test_export_rejects_bad_month(monkeypatch: pytest.MonkeyPatch) -> None:
    _install_mirror(monkeypatch, [])
    for bad in ("2026-13", "202607", "abc"):
        with pytest.raises(AppException) as exc_info:
            await exporter.export_maintenance_summary(None, bad)
        assert exc_info.value.status_code == 400
