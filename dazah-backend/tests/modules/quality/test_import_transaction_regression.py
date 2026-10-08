"""Real PostgreSQL savepoint and HTTP regressions for production import failures."""

from collections.abc import AsyncIterator
from io import BytesIO
from types import SimpleNamespace
from typing import Any
from uuid import uuid4

import pytest
from docx import Document
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from sqlalchemy import delete, select
from sqlalchemy.exc import DBAPIError
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.deps import get_current_user
from app.modules.quality.api import quality_change
from app.modules.quality.models import CAPA, ChangeControl
from app.modules.quality.service import quality_import_export as service


def _document(headers: list[str], rows: list[list[str]]) -> bytes:
    doc = Document()
    table = doc.add_table(rows=1, cols=len(headers))
    for index, label in enumerate(headers):
        table.cell(0, index).text = label
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            row.cells[index].text = value
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


@pytest.mark.anyio
async def test_change_import_http_rejects_long_fields_without_poisoning_other_rows(
    db_session: AsyncSession,
) -> None:
    prefix = f"import-{uuid4().hex}"
    document = _document(
        ["变更控制号", "变更对象"],
        [
            [f"{prefix}-1", "合法对象"],
            [f"{prefix}-2", "长" * 256],
            [f"{prefix}-3", "后续对象"],
        ],
    )
    test_app = FastAPI()
    test_app.include_router(quality_change.router, prefix="/api/v1/quality")

    async def session_override() -> AsyncIterator[AsyncSession]:
        yield db_session

    test_app.dependency_overrides[get_db] = session_override
    test_app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(
        id=uuid4()
    )
    try:
        async with AsyncClient(
            transport=ASGITransport(app=test_app), base_url="http://test"
        ) as client:
            malformed = await client.post(
                "/api/v1/quality/changes/import/preview",
                files={"file": ("import.docx", b"invalid document")},
            )
            assert malformed.status_code == 422
            assert "有效的 .docx" in malformed.json()["detail"]
            preview = await client.post(
                "/api/v1/quality/changes/import/preview",
                files={"file": ("import.docx", document)},
            )
            assert preview.status_code == 200
            assert preview.json()["data"]["valid_rows"] == 2
            assert "255" in preview.json()["data"]["error_rows"][0]["error_message"]
            response = await client.post(
                "/api/v1/quality/changes/import/confirm",
                files={"file": ("import.docx", document)},
            )
        assert response.status_code == 200
        result = response.json()["data"]
        assert result["success_count"] == 2
        assert result["error_count"] == 1
        assert result["error_details"] == [
            {"row": 3, "error": "变更对象不能超过255个字符（当前256个）"}
        ]
        records = (
            await db_session.scalars(
                select(ChangeControl).where(
                    ChangeControl.change_code.startswith(prefix)
                )
            )
        ).all()
        assert {record.change_code for record in records} == {
            f"{prefix}-1",
            f"{prefix}-3",
        }
    finally:
        await db_session.rollback()
        await db_session.execute(
            delete(ChangeControl).where(ChangeControl.change_code.startswith(prefix))
        )
        await db_session.commit()


@pytest.mark.anyio
@pytest.mark.parametrize("failure", ["duplicate", "overflow"])
async def test_change_import_real_flush_conflict_rolls_back_only_failed_row(
    db_session: AsyncSession,
    monkeypatch: pytest.MonkeyPatch,
    failure: str,
) -> None:
    prefix = f"import-{uuid4().hex}"
    db_session.add(ChangeControl(change_code=f"{prefix}-duplicate"))
    await db_session.commit()
    original = service.repo.create_change

    async def conflicting_create(
        db: AsyncSession, data: dict[str, Any]
    ) -> ChangeControl:
        if data["change_code"].endswith("-bad"):
            if failure == "duplicate":
                return await original(
                    db, {**data, "change_code": f"{prefix}-duplicate"}
                )
            return await original(db, {**data, "change_object": "长" * 256})
        return await original(db, data)

    monkeypatch.setattr(service.repo, "create_change", conflicting_create)
    try:
        result = await service.confirm_change_import(
            db_session,
            _document(
                ["变更控制号"],
                [[f"{prefix}-before"], [f"{prefix}-bad"], [f"{prefix}-after"]],
            ),
        )
        assert result["success_count"] == 2
        assert result["error_count"] == 1
        assert "SQL" not in str(result["error_details"])
        codes = set(
            await db_session.scalars(
                select(ChangeControl.change_code).where(
                    ChangeControl.change_code.startswith(prefix)
                )
            )
        )
        assert codes == {f"{prefix}-duplicate", f"{prefix}-before", f"{prefix}-after"}
    finally:
        await db_session.rollback()
        await db_session.execute(
            delete(ChangeControl).where(ChangeControl.change_code.startswith(prefix))
        )
        await db_session.commit()


@pytest.mark.parametrize(
    "sqlstate, expected",
    [
        ("22001", "字段格式"),
        ("23505", "编号已存在"),
        ("08006", None),
    ],
)
def test_import_database_error_classification_does_not_swallow_connection_failures(
    sqlstate: str,
    expected: str | None,
) -> None:
    class DatabaseFailureError(Exception):
        def __init__(self, state: str) -> None:
            super().__init__("private parameters")
            self.sqlstate = state

    error = DBAPIError("private SQL", {}, DatabaseFailureError(sqlstate))
    message = service._import_database_error_message(error)
    if expected is None:
        assert message is None
    else:
        assert message is not None and expected in message
        assert "private" not in message


@pytest.mark.anyio
async def test_capa_import_does_not_recreate_or_restore_deleted_code(
    db_session: AsyncSession,
) -> None:
    prefix = f"import-{uuid4().hex}"
    db_session.add(CAPA(capa_code=f"{prefix}-deleted", is_deleted=True))
    await db_session.commit()
    try:
        result = await service.confirm_capa_import(
            db_session,
            _document(
                ["CAPA编号"],
                [[f"{prefix}-deleted"], [f"{prefix}-valid"]],
            ),
            update_existing=True,
        )
        assert result["error_count"] == 1
        assert result["success_count"] == 1
        assert result["update_count"] == 0
        deleted = await service.repo.get_capa_by_code(
            db_session, f"{prefix}-deleted", include_deleted=True
        )
        assert deleted is not None and deleted.is_deleted
    finally:
        await db_session.rollback()
        await db_session.execute(delete(CAPA).where(CAPA.capa_code.startswith(prefix)))
        await db_session.commit()
