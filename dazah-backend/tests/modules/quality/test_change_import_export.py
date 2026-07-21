from __future__ import annotations

import io
import uuid
from datetime import date

import pytest
from docx import Document
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.quality.models.change_control import ChangeControl
from app.modules.quality.service import quality_import_export as ie_service


def build_change_docx() -> bytes:
    doc = Document()
    table = doc.add_table(rows=2, cols=10)
    headers = [
        "序号",
        "变更控制号",
        "变更申请部门",
        "变更对象",
        "变更内容",
        "变更等级",
        "变更申请日期",
        "变更计划批准日期",
        "变更正式执行日期",
        "变更关闭日期",
    ]
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    values = [
        "1",
        "BG-2026-001",
        "质量部",
        "反应釜",
        "更换搅拌电机",
        "二级",
        "2026-06-30",
        "2026-07-01",
        "2026-07-05",
        "2026-07-10",
    ]
    for index, value in enumerate(values):
        table.cell(1, index).text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.mark.anyio
async def test_confirm_change_import_updates_existing_row(
    db_session: AsyncSession,
) -> None:
    db_session.add(
        ChangeControl(
            id=uuid.uuid4(),
            serial_number="1",
            change_code="BG-2026-001",
            applicant_department="旧部门",
            change_object="旧对象",
            change_content="旧内容",
            change_level="三级",
            application_date=date(2026, 6, 1),
        )
    )
    await db_session.commit()

    result = await ie_service.confirm_change_import(
        db_session,
        build_change_docx(),
        skip_duplicates=False,
        update_existing=True,
    )

    updated = await db_session.scalar(
        select(ChangeControl).where(ChangeControl.change_code == "BG-2026-001")
    )
    assert result["update_count"] == 1
    assert updated is not None
    assert updated.applicant_department == "质量部"
    assert updated.change_level == "二级"


@pytest.mark.anyio
async def test_preview_change_import_accepts_desktop_template(
    db_session: AsyncSession,
) -> None:
    result = await ie_service.preview_change_import(db_session, build_change_docx())

    assert result["headers"] == [
        "序号",
        "变更控制号",
        "变更申请部门",
        "变更对象",
        "变更内容",
        "变更等级",
        "变更申请日期",
        "变更计划批准日期",
        "变更正式执行日期",
        "变更关闭日期",
    ]
    assert result["total_rows"] >= 0
