"""Excel imports must reject chart sheets before treating them as data tables."""

from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from app.core.exceptions import AppException
from app.modules.hr import esg_api
from app.modules.hr.trainer_api import _parse_excel_trainers
from app.platform.identity.deps import get_current_user


def test_trainer_import_rejects_active_chart_sheet() -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.append([1, 2])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=1, max_col=2, min_row=1))
    workbook.create_chartsheet("图表").add_chart(chart)
    workbook.active = 1
    output = BytesIO()
    workbook.save(output)
    with pytest.raises(AppException) as error:
        _parse_excel_trainers(output.getvalue())
    assert error.value.status_code == 422
    assert "工作表" in error.value.message


@pytest.mark.asyncio
async def test_esg_excel_import_normalizes_text_cells_through_real_route() -> None:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.append(["培训名称", "姓名", "培训方式", "培训日期"])
    sheet.append(["GMP", "张三", 1, "2026-07-01"])
    output = BytesIO()
    workbook.save(output)
    service = SimpleNamespace(create_record=AsyncMock())
    app = FastAPI()
    app.include_router(esg_api.router, prefix="/api/v1/hr")
    app.dependency_overrides[esg_api.get_esg_service] = lambda: service
    app.dependency_overrides[get_current_user] = lambda: SimpleNamespace(id=uuid4())
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/api/v1/hr/esg-training-records/import",
            params={"department": "质量部"},
            files={
                "file": (
                    "培训.xlsx",
                    output.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert response.status_code == 200, response.text
    service.create_record.assert_awaited_once()
    record = service.create_record.await_args.args[0]
    assert record.training_method == "1"
    assert record.employee_name == "张三"
    assert record.department == "质量部"
