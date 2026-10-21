"""危险源辨识台账 Excel 导入解析器单元测试。

覆盖四类车间表格格式、多子表识别、合并单元格展开、勾选列、
列错位子表的锚点解析、跳过规则与建库字段映射。
"""

import io
from typing import Any

import openpyxl
import pytest

from app.modules.safety.service.hazard_ledger_import import (
    FORMAT_ACTIVITY,
    FORMAT_FACILITY,
    FORMAT_KEY_RISK,
    FORMAT_LEGACY,
    SheetMatrix,
    build_hazard_records,
    detect_format,
    guess_department,
    load_workbook_matrices,
    parse_sheet,
    parse_workbook,
)


def _xlsx_bytes(
    sheets: dict[str, list[list[Any]]],
    merges: dict[str, list[str]] | None = None,
) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        ws = workbook.create_sheet(title=name)
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, value in enumerate(row, start=1):
                if value is not None:
                    ws.cell(row=r_idx, column=c_idx, value=value)
        for rng in (merges or {}).get(name, []):
            ws.merge_cells(rng)
    buf = io.BytesIO()
    workbook.save(buf)
    return buf.getvalue()


# ── 作业活动分级管控清单（含错位行）──

ACTIVITY_SHEET: list[list[Any]] = [
    ["作业活动安全风险分级管控措施清单"],
    ["单位名称："],
    ["风险点", None, None, "危险源", None, "可能发生的事故类型及后果",
     "时态（过去、现在、将来）", "状态（正常、异常、紧急）", "安全风险等级评定",
     None, None, None, None, "风险管控措施", None, None, None, "管控主体",
     None, None, "检查频次"],
    ["序号", "类别", "名称", "作业步骤", "危险源辨识", None, None, None,
     "可能性L", "暴露频次E", "严重性C", "危险度D", "风险等级",
     "工程措施", "管理措施", "个体防护", "应急措施", "管控层级",
     "责任单位", "责任人员", None],
    [1, "常规作业", "罐内发酵", "打开进料阀门", "管道憋压", "其他伤害",
     "现在", "正常", 1, 6, 1, 6, "低风险",
     "加防护罩", "定期检查", "手套", "应急卡", "班组级", "消毒班", "李涛", "每月一次"],
    # 错位行：整行左移两列（省略名称/作业步骤），由锚点解析兜底
    [None, "特殊作业：动火作业", "动火作业产生火花", "火灾、触电",
     "现在", "异常", 1, 6, 6, 36, "一般风险",
     "隔离", "管理措施B", "面罩", "演练", "车间级", "101车间",
     "朱健", "每天一次", None, None],
    ["编制/日期：xxx"],
]


def test_parse_activity_sheet_aligned_and_shifted_rows():
    matrix = SheetMatrix(name="消毒", rows=[list(r) for r in ACTIVITY_SHEET])
    fmt, rows, errors = parse_sheet(matrix)
    assert fmt == FORMAT_ACTIVITY
    assert not errors
    assert len(rows) == 2

    aligned = rows[0]
    assert aligned.category == "常规作业"
    assert aligned.name == "罐内发酵"
    assert aligned.step == "打开进料阀门"
    assert aligned.hazard_desc == "管道憋压"
    assert aligned.accident == "其他伤害"
    assert aligned.tense == "现在"
    assert aligned.state == "正常"
    assert (aligned.l_value, aligned.e_value, aligned.c_value, aligned.d_value) == (
        1.0, 6.0, 1.0, 6.0,
    )
    assert aligned.level_key == "level_4"
    assert aligned.engineering == "加防护罩"
    assert aligned.management == "定期检查"
    assert aligned.ppe == "手套"
    assert aligned.emergency == "应急卡"
    assert aligned.control_level == "班组级"
    assert aligned.responsible_unit == "消毒班"
    assert aligned.responsible_person == "李涛"
    assert aligned.check_frequency == "每月一次"

    shifted = rows[1]
    assert shifted.name == "特殊作业：动火作业"
    assert shifted.hazard_desc == "动火作业产生火花"
    assert shifted.accident == "火灾、触电"
    assert (shifted.l_value, shifted.e_value, shifted.c_value, shifted.d_value) == (
        1.0, 6.0, 6.0, 36.0,
    )
    assert shifted.level_key == "level_3"
    assert shifted.engineering == "隔离"


# ── 设备设施分级管控清单（合并单元格展开）──

FACILITY_SHEET: list[list[Any]] = [
    ["设备设施安全风险分级管控措施清单"],
    ["单位名称："],
    ["风险点", None, None, None, "危险源", None, None,
     "可能发生的事故类型及后果", "安全风险等级评定", None, None, None,
     "风险管控措施", None, None, None, "管控主体", None, None, "检查频次"],
    ["序号", "类别", "名称", "位置/区域", "辨识对象", "危险有害因素", "检查标准",
     None, "可能性L", "严重性S", "风险度R", "风险等级",
     "工程措施", "管理措施", "个体防护", "应急措施", "管控层级",
     "责任单位", "责任人员", None],
    [1, "储罐及容器类", "调浆罐", "一楼配料间", "电机搅拌", "转轴伤害", "查防护罩",
     "机械伤害", 1, 3, 3, "低风险", "防护罩", "每班检查", "工服", "应急物资",
     "岗位级", "配料班", "岗位员工", "每班一次"],
    [None, None, None, None, "电源开关", "漏电", "查密封", "触电",
     1, 3, 3, "低风险", "低压用电", "每班检查", "工服", "应急物资",
     "岗位级", "配料班", "岗位员工", "每班一次"],
]


def test_parse_facility_sheet_with_merged_cells():
    data = _xlsx_bytes(
        {"101配料": FACILITY_SHEET},
        merges={"101配料": ["A5:A6", "B5:B6", "C5:C6", "D5:D6"]},
    )
    matrices = load_workbook_matrices(data, "设备设施分级管控总清单.xlsx")
    assert len(matrices) == 1
    fmt, rows, errors = parse_sheet(matrices[0])
    assert fmt == FORMAT_FACILITY
    assert not errors
    assert len(rows) == 2

    second = rows[1]
    # 合并单元格值应已向下展开
    assert second.category == "储罐及容器类"
    assert second.name == "调浆罐"
    assert second.location == "一楼配料间"
    assert second.target == "电源开关"
    assert second.hazard_desc == "漏电"
    assert second.check_standard == "查密封"
    assert second.accident == "触电"
    assert (second.l_value, second.c_value, second.d_value) == (1.0, 3.0, 3.0)
    assert second.e_value is None  # LSR 口径无 E 值
    assert second.level_key == "level_4"


# ── 重点风险管控措施清单 ──

KEY_RISK_SHEET: list[list[Any]] = [
    ["作业活动重点风险管控措施清单（常规作业）"],
    ["      单位名称：201一车间           "],
    ["风险点", None, None, "危险源", None, "可能发生的事故类型及后果",
     "风险管控措施", None, None, None, "AI搜索", None, None],
    ["序号", "类别", "名称", "作业步骤", "危险源辨识", None,
     "工程措施", "管理措施", "个体防护", "应急措施",
     "有效措施", "技术设备更新", "案例"],
    [1, "常规作业", "硫酸卸车", "取样检测", "取样滑落", "高处坠落",
     "围堰", "培训", "面罩", "预案", "许可制度", "智能系统", "恒瑞案例"],
]


def test_parse_key_risk_sheet_maps_ai_columns():
    matrix = SheetMatrix(name="Sheet1", rows=[list(r) for r in KEY_RISK_SHEET])
    fmt, rows, errors = parse_sheet(matrix)
    assert fmt == FORMAT_KEY_RISK
    assert not errors
    assert len(rows) == 1
    row = rows[0]
    assert row.name == "硫酸卸车"
    assert row.step == "取样检测"
    assert row.hazard_desc == "取样滑落"
    assert row.accident == "高处坠落"
    assert row.engineering == "围堰"
    assert row.level_key is None  # 无 LEC 数值
    assert "有效措施：许可制度" in (row.extra_recommendation or "")
    assert "技术设备更新：智能系统" in (row.extra_recommendation or "")
    assert "同业案例：恒瑞案例" in (row.extra_recommendation or "")


# ── 危险源识别评价表（legacy 勾选表头）──

LEGACY_SHEET: list[list[Any]] = [
    ["危险源识别评价表"],
    ["部门：103车间                    "],
    ["序号", "工序", "作业活动", "危险源", "可能导致事故", "时态", None, None,
     "状态", None, None, "评价", None, None, None, "危\n险\n等\n级",
     "是否重要危险源", "防范措施", "红色", "D值大于320分", "一级"],
    [None, None, None, None, None, "过去", "现在", "将来", "正常", "异常", "紧急",
     "L", "E", "C", "D", None, None, None, "橙色", "D值160－320分", "二级"],
    [1.0, "准备工作", "取样", "长发卷入运行设备", "机械致害", None, "√", None,
     "√", None, None, 1.0, 6.0, 1.0, 6.0, "四级", "否", "戴工帽",
     "黄色", "D值70－160分", "三级"],
    ["编制/日期：                "],
]


def test_parse_legacy_sheet_with_check_columns():
    matrix = SheetMatrix(name="化验室", rows=[list(r) for r in LEGACY_SHEET])
    fmt, rows, errors = parse_sheet(matrix)
    assert fmt == FORMAT_LEGACY
    assert not errors
    assert len(rows) == 1
    row = rows[0]
    assert row.name == "准备工作"
    assert row.step == "取样"
    assert row.hazard_desc == "长发卷入运行设备"
    assert row.accident == "机械致害"
    assert row.tense == "现在"
    assert row.state == "正常"
    assert (row.l_value, row.e_value, row.c_value, row.d_value) == (1.0, 6.0, 1.0, 6.0)
    # 表头「危\n险\n等\n级」带换行也应命中等级列
    assert row.level_key == "level_4"
    assert row.is_major == "否"
    assert row.management == "戴工帽"


def test_legacy_d_value_computed_when_missing():
    sheet = [list(r) for r in LEGACY_SHEET]
    sheet[4][14] = None  # 去掉 D 列
    matrix = SheetMatrix(name="化验室", rows=sheet)
    _, rows, _ = parse_sheet(matrix)
    assert rows[0].d_value == 6.0  # 1*6*1


# ── 工作簿级：多子表与跳过规则 ──


def test_parse_workbook_multi_sheet_and_skips():
    summary_sheet = [
        ["101一车间危险源辨识风险点统计"],
        ["分布点", "风险等级", "一级（红色）", "二级（橙色）",
         "三级（黄色）", "四级（蓝色）", "备注"],
        ["液氨", None, 0, 0, 3, 46, None],
    ]
    data = _xlsx_bytes(
        {
            "消毒": ACTIVITY_SHEET,
            "风险点总计": summary_sheet,
            "Sheet1": [[None]],
        }
    )
    result = parse_workbook(data, "作业活动分级管控总清单101-102.xlsx")

    assert result["department_guess"] == "101-102车间"
    assert result["total_valid_rows"] == 2
    by_name = {sheet["name"]: sheet for sheet in result["sheets"]}
    assert by_name["消毒"]["format"] == FORMAT_ACTIVITY
    assert by_name["消毒"]["row_count"] == 2
    assert len(by_name["消毒"]["sample_rows"]) == 2
    assert by_name["风险点总计"]["skipped"] is True
    assert by_name["风险点总计"]["skip_reason"] == "无法识别的表格格式"
    assert by_name["Sheet1"]["skipped"] is True
    assert by_name["Sheet1"]["skip_reason"] == "空子表"


def test_template_instruction_sheet_produces_no_rows():
    instruction_sheet = [
        ["设备设施安全风险分级管控措施清单"],
        ["单位名称："],
        ["风险点", None, None, None, "危险源", None, None,
         "可能发生的事故类型及后果", "安全风险等级评定", None, None, None,
         "风险管控措施", None, None, None, "管控主体", None, None, "检查频次"],
        ["序号", "类别", "名称", "位置/区域", "辨识对象", "危险有害因素", "检查标准",
         None, "可能性L", "严重性S", "风险度R", "风险等级",
         "工程措施", "管理措施", "个体防护", "应急措施", "管控层级",
         "责任单位", "责任人员", None],
        [1, "设备分类。如：炉类、塔类", "《具体风险点名称》", "部位", "具体指向",
         "具体内容", " 提问式检查表", "国标6441事故分类类型",
         "按制度评分", "按制度评分", "按制度评分", "按制度分级", None, None],
    ]
    result = parse_workbook(
        _xlsx_bytes({"各岗位分开填写": instruction_sheet}),
        "201一车间设备设施分级管控总清单.xlsx",
    )
    sheet = result["sheets"][0]
    assert sheet["format"] == FORMAT_FACILITY
    assert sheet["skipped"] is True
    assert sheet["skip_reason"] == "未解析出有效数据行"


def test_load_rejects_unsupported_extension():
    with pytest.raises(ValueError, match="仅支持"):
        load_workbook_matrices(b"x", "清单.csv")


# ── 等级归一与部门推断 ──


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("低风险", "level_4"),
        ("四级", "level_4"),
        ("一般风险", "level_3"),
        ("三级", "level_3"),
        ("较大风险", "level_2"),
        ("二级", "level_2"),
        ("重大风险", "level_1"),
        ("一级", "level_1"),
        ("一级（红色）", "level_1"),
        ("未知等级", None),
        (None, None),
    ],
)
def test_normalize_level_aliases(raw, expected):
    matrix = SheetMatrix(name="t", rows=[["作业活动安全风险分级管控措施清单"], [None]])
    assert detect_format(matrix) == FORMAT_ACTIVITY
    # 直接经由解析行验证等级归一
    from app.modules.safety.service.hazard_ledger_import import _normalize_level

    assert _normalize_level(raw) == expected


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("作业活动分级管控总清单101-102.xlsx", "101-102车间"),
        ("103车间危险源辨识评价表2026.04.06.xls", "103车间"),
        ("201一车间作业活动分级管控总清单.xlsx", "201一车间"),
        ("作业活动重点风险管控措施清单.xlsx", ""),
    ],
)
def test_guess_department(filename, expected):
    assert guess_department(filename) == expected


# ── 建库字段映射 ──


def _sample_parsed_rows() -> list[Any]:
    _, rows, _ = parse_sheet(
        SheetMatrix(name="消毒", rows=[list(r) for r in ACTIVITY_SHEET])
    )
    return rows


def test_build_hazard_records_fields_and_status():
    records = build_hazard_records(
        _sample_parsed_rows(),
        department="101-102车间",
        position="消毒",
        source_label="作业活动清单.xlsx/消毒",
    )
    assert len(records) == 2
    first = records[0]
    assert first["hazard_id_no"].startswith("IMP-")
    assert len(first["hazard_id_no"]) <= 64
    assert first["department"] == "101-102车间"
    assert first["position"] == "消毒"
    assert first["production_step"] == "罐内发酵"
    assert first["specific_activity"] == "打开进料阀门"
    assert first["hazard_type"] == "常规作业"
    assert first["unsafe_behavior"] == "管道憋压"
    assert first["possible_accident"] == "其他伤害"
    assert first["inherent_risk_level"] == "level_4"
    assert first["inherent_risk_label"] == "四级/低风险"
    assert first["existing_engineering_controls"] == "加防护罩"
    assert first["responsible_person"] == "李涛"
    assert first["responsible_unit"] == "消毒班"
    assert first["control_level"] == "班组级"
    assert first["check_frequency"] == "每月一次"
    assert first["tense"] == "现在"
    assert first["state"] == "正常"
    assert first["overall_status"] == "completed"
    assert first["ai_node_progress"] == "completed"
    assert first["script3_review_status"] == "approved"
    assert first["notes"] == "来源：导入《作业活动清单.xlsx/消毒》"


def test_build_hazard_records_dedup_identical_rows():
    rows = _sample_parsed_rows()
    records = build_hazard_records(
        rows + [rows[0]],
        department="101-102车间",
        position="消毒",
        source_label="清单.xlsx/消毒",
    )
    # 追加的重复行按内容指纹去重
    assert len(records) == 2


def test_build_hazard_records_facility_equipment():
    data = _xlsx_bytes(
        {"101配料": FACILITY_SHEET},
        merges={"101配料": ["A5:A6", "B5:B6", "C5:C6", "D5:D6"]},
    )
    matrices = load_workbook_matrices(data, "设备设施分级管控总清单.xlsx")
    _, rows, _ = parse_sheet(matrices[0])
    records = build_hazard_records(
        rows,
        department="101-102车间",
        position="101配料",
        source_label="设备设施清单.xlsx/101配料",
        is_facility=True,
    )
    assert records[0]["equipment_facilities"] == "调浆罐；一楼配料间"
    assert records[0]["specific_activity"] == "电机搅拌"
    assert records[0]["e_inherent"] is None


def test_build_hazard_records_level_fallback_by_d_value():
    rows = _sample_parsed_rows()
    rows[0].level_key = None
    rows[0].level_raw = None
    records = build_hazard_records(
        rows[:1],
        department="车间",
        position="消毒",
        source_label="清单.xlsx/消毒",
    )
    # D=6 且存在 E 值 → 按 LEC 阈值兜底为 level_4
    assert records[0]["inherent_risk_level"] == "level_4"


# ── 201 车间尾部多列变体（检查标准/建议措施列，无管控主体） ──

ACTIVITY_201_SHEET: list[list[Any]] = [
    ["作业活动安全风险分级管控措施清单"],
    ["单位名称："],
    ["风险点", None, None, "危险源", None, "可能发生的事故类型及后果",
     "时态（过去、现在、将来）", "状态（正常、异常、紧急）", "安全风险等级评定",
     None, None, None, None, "风险管控措施", None, None, None, None, None],
    ["序号", "类别", "名称", "作业步骤", "危险源辨识", None, None, None,
     "可能性L", "暴露频次E", "严重性C", "危险度D", "风险等级",
     "工程措施", "管理措施", "个体防护", "应急措施", "检查标准", "建议措施"],
    [1, "常规作业", "接收发酵液", "打开进料阀门", "管道憋压", "其他伤害",
     "现在", "正常", 1, 3, 7, 21, "四级",
     "管道保温", "定期检查", "面罩", "应急卡",
     "管道压力稳定在0.2-0.4MPa范围内，无憋压现象", "建议3年内逐步更新老旧设备"],
]


def test_parse_activity_201_variant_tail_columns():
    """201 表尾部是检查标准/建议措施列：不得把建议文本误当管控层级/责任单位。"""
    matrix = SheetMatrix(name="压滤", rows=[list(r) for r in ACTIVITY_201_SHEET])
    fmt, rows, errors = parse_sheet(matrix)
    assert fmt == FORMAT_ACTIVITY
    assert not errors
    row = rows[0]
    assert row.engineering == "管道保温"
    assert row.management == "定期检查"
    assert row.ppe == "面罩"
    assert row.emergency == "应急卡"
    assert row.check_standard == "管道压力稳定在0.2-0.4MPa范围内，无憋压现象"
    # 建议措施列无表头映射，不得串位
    assert row.control_level is None
    assert row.responsible_unit is None
    assert row.responsible_person is None
    assert row.check_frequency is None


def test_fit_truncates_bounded_fields():
    from app.modules.safety.service.hazard_ledger_import import _fit

    long_text = "长" * 80
    assert len(_fit("responsible_unit", long_text) or "") == 80  # 100 上限未到不截断
    assert len(_fit("control_level", long_text) or "") == 50
    assert (_fit("control_level", long_text) or "").endswith("…")
    assert _fit("engineering", long_text) == long_text  # Text 列不截断
    assert _fit("control_level", None) is None
