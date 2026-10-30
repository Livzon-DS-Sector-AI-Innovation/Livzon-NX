"""MSDS AI 识别服务测试：mock llm_client，覆盖正常/校验失败/各类 LLM 异常。"""

from __future__ import annotations

from types import SimpleNamespace
from typing import Any
from unittest.mock import AsyncMock

import pytest

from app.core.llm.exceptions import (
    LLMConfigError,
    LLMOutputError,
    LLMProviderError,
    LLMRateLimitError,
)
from app.modules.safety.service import msds_extract


def _install_chat_json(
    monkeypatch: pytest.MonkeyPatch, result: dict[str, Any]
) -> AsyncMock:
    fake_client = SimpleNamespace(
        chat_json=AsyncMock(return_value=result),
        chat_vision_json=AsyncMock(return_value=result),
    )
    monkeypatch.setattr(msds_extract, "llm_client", fake_client)
    return fake_client.chat_json


def _docx_bytes() -> bytes:
    """构造一个最小可解析 docx（document_parser 走 python-docx）。"""
    from io import BytesIO

    from docx import Document

    doc = Document()
    doc.add_paragraph("化学品名称：丙酮")
    doc.add_paragraph("CAS 号：67-64-1")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


_GOOD_OUTPUT = {
    "name": "丙酮",
    "alias": "阿西通",
    "cas_no": "67-64-1",
    "category": "易燃液体",
    "hazard_statements": "高度易燃液体和蒸气 H225",
    "storage_requirements": "阴凉通风，远离火源",
    "remark": "灭火剂：抗醇泡沫",
}


async def test_extract_docx_success(monkeypatch: pytest.MonkeyPatch) -> None:
    chat_json = _install_chat_json(monkeypatch, dict(_GOOD_OUTPUT))
    suggestion, error = await msds_extract.extract_msds_suggestion(
        file_name="丙酮MSDS.docx", content=_docx_bytes()
    )
    assert error is None
    assert suggestion is not None
    assert suggestion.name == "丙酮"
    assert suggestion.cas_no == "67-64-1"
    assert suggestion.category == "flammable_liquid"
    assert suggestion.category_label == "易燃液体"
    chat_json.assert_awaited_once()
    prompt = chat_json.await_args.args[0][0]["content"]
    assert "丙酮" in prompt and "67-64-1" in prompt


async def test_extract_invalid_cas_and_unknown_category(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    _install_chat_json(
        monkeypatch,
        {"name": "X", "cas_no": "67-64", "category": "神秘物质", "alias": None},
    )
    suggestion, error = await msds_extract.extract_msds_suggestion(
        file_name="丙酮MSDS.docx", content=_docx_bytes()
    )
    assert error is None
    assert suggestion is not None
    assert suggestion.cas_no is None
    assert suggestion.category is None
    assert suggestion.category_label == ""


@pytest.mark.parametrize(
    ("exc", "expected"),
    [
        (LLMConfigError("no config"), "平台未配置 AI 模型，无法自动识别"),
        (LLMRateLimitError("rate"), "AI 模型限流，请稍后重试识别"),
        (LLMOutputError("bad json"), "AI 识别结果格式无效，请重试或手动填写"),
        (LLMProviderError("down"), "AI 模型服务暂不可用，请稍后重试识别"),
    ],
)
async def test_extract_llm_errors(
    monkeypatch: pytest.MonkeyPatch, exc: Exception, expected: str
) -> None:
    fake_client = SimpleNamespace(chat_json=AsyncMock(side_effect=exc))
    monkeypatch.setattr(msds_extract, "llm_client", fake_client)
    suggestion, error = await msds_extract.extract_msds_suggestion(
        file_name="丙酮MSDS.docx", content=_docx_bytes()
    )
    assert suggestion is None
    assert error == expected


async def test_extract_unsupported_extension(monkeypatch: pytest.MonkeyPatch) -> None:
    _install_chat_json(monkeypatch, dict(_GOOD_OUTPUT))
    suggestion, error = await msds_extract.extract_msds_suggestion(
        file_name="file.zip", content=b"PK"
    )
    assert suggestion is None
    assert "不支持" in (error or "")


def _patch_doc_extraction(
    monkeypatch: pytest.MonkeyPatch,
    *,
    converted_pdf: bytes,
    pdf_text: str,
) -> AsyncMock:
    """模拟 .doc 链路：直接解析失败 → soffice 转 PDF（结果可控）→ PDF 文本提取。

    返回 chat_json mock 以便断言 prompt 内容。
    """
    chat_json = _install_chat_json(monkeypatch, dict(_GOOD_OUTPUT))

    def _extract(file_path: str, max_chars: int = 50000) -> str:
        from pathlib import Path

        if Path(file_path).suffix.lower() == ".doc":
            # 模拟 python-docx 无法解析旧二进制 .doc
            raise RuntimeError("PackageNotFoundError")
        assert Path(file_path).suffix.lower() == ".pdf"
        return pdf_text

    monkeypatch.setattr(msds_extract, "extract_to_markdown", _extract)
    monkeypatch.setattr(
        "app.platform.storage.preview.convert_office_to_pdf",
        lambda content, file_name: converted_pdf,
    )
    return chat_json


async def test_extract_doc_falls_back_to_soffice_conversion(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """真 .doc（OLE2）python-docx 解析失败时，走 LibreOffice 转 PDF 后再提取。"""
    chat_json = _patch_doc_extraction(
        monkeypatch,
        converted_pdf=b"%PDF-1.4 fake",
        pdf_text="化学品名称：丙酮（经转换提取）",
    )
    suggestion, error = await msds_extract.extract_msds_suggestion(
        file_name="旧版丙酮MSDS.doc", content=b"\xd0\xcf\x11\xe0ole2"
    )
    assert error is None
    assert suggestion is not None
    assert suggestion.name == "丙酮"
    prompt = chat_json.await_args.args[0][0]["content"]
    assert "经转换提取" in prompt


async def test_extract_doc_conversion_unavailable_returns_hint(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """无 LibreOffice（转换返回空字节）时识别降级为提示，不抛异常。"""
    _patch_doc_extraction(monkeypatch, converted_pdf=b"", pdf_text="")
    suggestion, error = await msds_extract.extract_msds_suggestion(
        file_name="旧版丙酮MSDS.doc", content=b"\xd0\xcf\x11\xe0ole2"
    )
    assert suggestion is None
    assert error == "未能从附件中提取到文本内容"


async def test_extract_image_uses_vision(monkeypatch: pytest.MonkeyPatch) -> None:
    vision = AsyncMock(return_value=dict(_GOOD_OUTPUT))
    fake_client = SimpleNamespace(chat_json=AsyncMock(), chat_vision_json=vision)
    monkeypatch.setattr(msds_extract, "llm_client", fake_client)
    suggestion, error = await msds_extract.extract_msds_suggestion(
        file_name="scan.png", content=b"\x89PNG\r\n\x1a\n"
    )
    assert error is None
    assert suggestion is not None
    vision.assert_awaited_once()
    data_uri = vision.await_args.args[1][0]
    assert data_uri.startswith("data:image/png;base64,")


async def test_extract_all_fields_empty(monkeypatch: pytest.MonkeyPatch) -> None:
    _install_chat_json(monkeypatch, {key: None for key in _GOOD_OUTPUT})
    suggestion, error = await msds_extract.extract_msds_suggestion(
        file_name="丙酮MSDS.docx", content=_docx_bytes()
    )
    assert suggestion is None
    assert error == "AI 未能从该附件识别出有效档案信息"
