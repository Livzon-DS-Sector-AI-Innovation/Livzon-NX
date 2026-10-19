"""Import/calculation regressions from the July energy workbook review."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import openpyxl
import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient

from app.core.exceptions import AppException
from app.modules.energy.ledger_calculations import (
    Calculation,
    Observation,
    calculate_point,
    calculate_rule,
)
from app.modules.energy.ledger_imports import confirm_import, inspect_xlsx
from app.modules.energy.ledger_schemas import (
    EntrySubmission,
    ImportConfirm,
    MeterDefinition,
    MeterWrite,
    RuleDefinition,
)
from app.modules.energy.ledger_template import extract_template, stable_code


def content(book: openpyxl.Workbook) -> bytes:
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def electricity_book() -> openpyxl.Workbook:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "电量"
    sheet.append(
        [
            "日期",
            "总计",
            "设备甲",
            "设备乙",
            "小计",
            "其他",
            "分表总和",
            "西动力总",
            "东动力总",
        ]
    )
    sheet.append(["7/1", 100, 30, None, 30, 70, 100, 40, 60])
    sheet.append(["7/2", 120, 40, 20, 60, 60, 120, 50, 70])
    return book


@pytest.mark.parametrize("quoted", [False, True])
def test_counter_formula_accepts_both_sheet_reference_forms(quoted: bool) -> None:
    book = openpyxl.Workbook()
    source = book.active
    assert source is not None
    source.title = "饮用水码"
    source.append(["日期", "水表"])
    source.append(["7/1", 100])
    source.append(["7/2", 110])
    usage = book.create_sheet("饮用水量")
    usage.append(["日期", "水量"])
    name = "'饮用水码'" if quoted else "饮用水码"
    usage.append(["7/1", f"={name}!B3-{name}!B2"])
    draft = extract_template(inspect_xlsx(content(book), 2026))
    assert len(draft.rules) == 1
    assert draft.rules[0].terms[0].metric_key == "point:" + stable_code("饮用水码", 2)


def test_value_only_electricity_has_four_rules_and_no_calculated_meters() -> None:
    preview = inspect_xlsx(content(electricity_book()), 2026)
    sheet = preview["sheets"][0]
    assert sheet["calculated_columns"] == [2, 5, 6, 7]
    assert sheet["input_columns"] == [3, 4, 8, 9]
    draft = extract_template(preview)
    assert {point.source_column for point in draft.meters} == {3, 4, 8, 9}
    assert {rule.source_column for rule in draft.rules} == {2, 5, 6, 7}
    subtotal = next(rule for rule in draft.rules if rule.source_column == 5)
    assert (
        calculate_rule(
            subtotal,
            {
                "point:" + stable_code("电量", 3): Calculation(
                    Decimal(30), "complete", {}
                )
            },
        ).value
        is None
    )
    difference = next(rule for rule in draft.rules if rule.source_column == 6)
    assert difference.report_scope == "check"


def test_preview_limits_readings_and_formulas_to_the_selected_period() -> None:
    book = electricity_book()
    sheet = book.create_sheet("日总量")
    sheet.append(["日期", "电量"])
    sheet.append(["7/1", "=电量!B2"])
    sheet.append(["7/2", "=电量!B3"])
    preview = inspect_xlsx(content(book), 2026, date(2026, 7, 1), date(2026, 7, 1))
    assert preview["start_date"] == "2026-07-01"
    assert all(
        cell["date"] == "2026-07-01"
        for sheet in preview["sheets"]
        for cell in sheet["cells"]
    )
    assert [item["date"] for item in preview["formulas"]] == ["2026-07-01"]


def test_ambiguous_formula_composition_does_not_create_a_misleading_rule() -> None:
    book = electricity_book()
    sheet = book.create_sheet("201-2车间")
    sheet.append(["日期", "电"])
    sheet.append(["7/1", "=电量!C2"])
    sheet.append(["7/2", "=电量!C3+电量!D3"])
    draft = extract_template(inspect_xlsx(content(book), 2026))
    assert not any(rule.source_sheet == "201-2车间" for rule in draft.rules)
    assert any("计算组成" in issue and "待确认" in issue for issue in draft.issues)


def test_shifted_electricity_reference_uses_the_same_business_date() -> None:
    book = electricity_book()
    sheet = book.create_sheet("101车间")
    sheet.append(["日期", "电"])
    sheet.append(["7/1", "=电量!C2"])
    sheet.append(["7/2", "=电量!C2"])
    draft = extract_template(inspect_xlsx(content(book), 2026))
    rule = next(rule for rule in draft.rules if rule.source_sheet == "101车间")
    result = calculate_rule(
        rule,
        {"point:" + stable_code("电量", 3): Calculation(Decimal(40), "complete", {})},
    )
    assert result.value == 40
    assert any("同日" in issue for issue in draft.issues)


def test_mixed_temperature_reference_keeps_the_confirmed_seven_degree_source() -> None:
    book = openpyxl.Workbook()
    source = book.active
    assert source is not None
    source.title = "冰水码"
    source.append(["日期", "201-2车间7°", "201二车间-5°"])
    source.append(["7/1", 100, 200])
    source.append(["7/2", 110, 220])
    source.append(["7/3", 120, 240])
    usage = book.create_sheet("冰水量")
    usage.append(["日期", "201-2车间7°", "201二车间-5°"])
    for row in [2, 3]:
        usage.append(
            [
                f"7/{row - 1}",
                f"=冰水码!B{row + 1}-冰水码!B{row}",
                f"=冰水码!C{row + 1}-冰水码!C{row}",
            ]
        )
    workshop = book.create_sheet("201-2车间")
    workshop.append(["日期", "7°冷水"])
    workshop.append(["7/1", "=冰水量!B2"])
    workshop.append(["7/2", "=冰水量!C3"])
    draft = extract_template(inspect_xlsx(content(book), 2026))
    rule = next(rule for rule in draft.rules if rule.source_sheet == "201-2车间")
    assert rule.temperature == "7℃"
    assert rule.terms[0].metric_key == "rule:" + stable_code("冰水量", 2)
    assert any("温度" in issue for issue in draft.issues)


@pytest.mark.asyncio
async def test_confirm_rejects_flattened_total_even_for_an_old_preview() -> None:
    batch = SimpleNamespace(
        payload={
            "sheets": [
                {
                    "name": "电量",
                    "headers": ["日期", "总计"],
                    "cells": [{"date": "2026-07-01", "column": 2, "value": "100"}],
                }
            ]
        }
    )
    meter = SimpleNamespace(
        definition=MeterDefinition(
            code="bad",
            name="总计",
            category="electricity",
            area="测试",
            department="测试",
            unit="kWh",
            reading_type="direct",
            active_from=date(2026, 7, 1),
            enabled=True,
        ).model_dump(mode="json")
    )
    service = SimpleNamespace(
        require_manager=AsyncMock(),
        repo=SimpleNamespace(
            serialize_write=AsyncMock(), get=AsyncMock(side_effect=[batch, meter])
        ),
    )
    request = ImportConfirm(
        mappings=[{"sheet": "电量", "column": 2, "point_id": uuid4()}],
        reason="回归测试",
    )
    with pytest.raises(AppException, match="自动计算"):
        await confirm_import(service, uuid4(), request)


@pytest.mark.asyncio
async def test_preview_route_accepts_a_period_and_rejects_an_inverted_period() -> None:
    from app.modules.energy.ledger_api import ledger_service, router

    app = FastAPI()
    app.include_router(router, prefix="/api/v1/energy")
    service = SimpleNamespace(
        require_manager=AsyncMock(),
        repo=SimpleNamespace(serialize_write=AsyncMock()),
        db=SimpleNamespace(
            scalar=AsyncMock(return_value=None),
            add=lambda item: None,
            flush=AsyncMock(),
        ),
        actor_id=None,
    )
    app.dependency_overrides[ledger_service] = lambda: service
    files = {
        "file": (
            "energy.xlsx",
            content(electricity_book()),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/api/v1/energy/imports/preview",
            params={"year": 2026, "start_date": "2026-07-01", "end_date": "2026-07-01"},
            files=files,
        )
        assert response.status_code == 200, response.text
        assert response.json()["data"]["payload"]["end_date"] == "2026-07-01"
        response = await client.post(
            "/api/v1/energy/imports/preview",
            params={"year": 2026, "start_date": "2026-07-02", "end_date": "2026-07-01"},
            files=files,
        )
        assert response.status_code == 422


def legacy_total() -> MeterDefinition:
    return MeterDefinition(
        code="legacy_total",
        name="总计",
        source_sheet="电量",
        source_column=2,
        category="electricity",
        reading_type="direct",
        area="测试",
        department="测试",
        unit="kWh",
        active_from=date(2026, 7, 1),
        enabled=True,
    )


def test_legacy_total_is_readable_but_cannot_feed_calculations_or_be_enabled() -> None:
    meter = legacy_total()
    with pytest.raises(ValueError, match="自动计算"):
        MeterWrite(definition=meter)
    result = calculate_point(
        meter,
        date(2026, 7, 1),
        {
            date(2026, 7, 1): Observation(Decimal(999), 1),
        },
    )
    assert result is not None and result.value is None
    assert result.status == "configuration_conflict"
    assert MeterWrite(definition=meter.model_copy(update={"enabled": False}))
    with pytest.raises(ValueError, match="自动计算"):
        MeterWrite(definition=meter.model_copy(update={"source_sheet": ""}))


def test_legacy_electricity_difference_keeps_negative_values_for_checks() -> None:
    rule = RuleDefinition(
        code="difference",
        name="电量·其他",
        source_sheet="电量",
        source_column=42,
        category="electricity",
        unit="kWh",
        department="测试",
        active_from=date(2026, 7, 1),
        report_scope="source",
        terms=[
            {"metric_key": "point:total"},
            {"metric_key": "point:subtotal", "coefficient": -1},
        ],
    )
    assert rule.report_scope == "check"
    assert (
        calculate_rule(
            rule,
            {
                "point:total": Calculation(Decimal(90), "complete", {}),
                "point:subtotal": Calculation(Decimal(100), "complete", {}),
            },
        ).value
        == -10
    )


@pytest.mark.asyncio
async def test_entry_does_not_generate_or_accept_manual_totals() -> None:
    from app.modules.energy.ledger_service import EnergyLedgerService

    service = EnergyLedgerService(AsyncMock())
    point_id, owner_id, reading_id = uuid4(), uuid4(), uuid4()
    definition = legacy_total()
    service.require_manager = AsyncMock()
    service.definitions_on = AsyncMock(
        return_value=({"point:legacy_total": (point_id, 1, definition)}, {})
    )

    plan = SimpleNamespace(
        id=uuid4(),
        definition={
            "name": "旧填报组",
            "point_ids": [str(point_id)],
            "owner_ids": [str(owner_id)],
            "publish_time": "08:00",
            "deadline_time": "10:00",
            "reminder_time": "10:30",
        },
    )
    service.repo = SimpleNamespace(
        serialize_write=AsyncMock(),
        all=AsyncMock(return_value=[plan]),
        task_for=AsyncMock(return_value=None),
        get=AsyncMock(
            return_value=SimpleNamespace(definition=definition.model_dump(mode="json"))
        ),
        task_readings=AsyncMock(
            return_value=[SimpleNamespace(id=reading_id, point_id=point_id, version=0)]
        ),
    )
    assert await service.generate_tasks(date(2026, 7, 2)) == []
    service.task = AsyncMock()
    with pytest.raises(AppException, match="自动计算"):
        await service.submit(
            uuid4(),
            EntrySubmission(
                values=[{"reading_id": reading_id, "version": 0, "value": "100"}]
            ),
        )


@pytest.mark.asyncio
async def test_workbook_marks_legacy_manual_total_as_a_configuration_conflict() -> None:
    from app.modules.energy.ledger_workbook import EnergyWorkbookService

    workbook = EnergyWorkbookService(SimpleNamespace())
    workbook.loaded = (2026, 7)
    workbook.headers["电量"] = ["日期", "总计"]
    workbook.definitions["point:legacy_total"] = legacy_total()
    workbook.readings["point:legacy_total", date(2026, 7, 1)] = SimpleNamespace(
        value=Decimal(999)
    )
    report = await workbook.report(2026, 7, "电量")
    assert report.rows[0].cells["2"].value is None
    assert report.rows[0].cells["2"].status == "configuration_conflict"


@pytest.mark.asyncio
async def test_http_rejects_manual_total_configuration_and_confirmation() -> None:
    from app.modules.energy.ledger_api import ledger_service, router

    app = FastAPI()
    app.include_router(router, prefix="/api/v1/energy")
    batch = SimpleNamespace(
        payload={"sheets": [{"name": "电量", "headers": ["日期", "总计"]}]}
    )
    meter = SimpleNamespace(definition=legacy_total().model_dump(mode="json"))
    service = SimpleNamespace(
        require_manager=AsyncMock(),
        save_meter=AsyncMock(),
        repo=SimpleNamespace(
            serialize_write=AsyncMock(), get=AsyncMock(side_effect=[batch, meter])
        ),
    )
    app.dependency_overrides[ledger_service] = lambda: service
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/api/v1/energy/meter-points", json={"definition": meter.definition}
        )
        assert response.status_code == 422
        service.save_meter.assert_not_called()
        response = await client.post(
            f"/api/v1/energy/imports/{uuid4()}/confirm",
            json={
                "mappings": [{"sheet": "电量", "column": 2, "point_id": str(uuid4())}],
                "reason": "回归测试",
            },
        )
        assert response.status_code == 422
        assert "自动计算" in response.json()["detail"]


def test_counter_missing_next_day_does_not_become_zero_or_negative() -> None:
    meter = MeterDefinition(
        code="water",
        name="水表",
        category="drinking_water",
        source_sheet="饮用水码",
        reading_type="counter",
        area="测试",
        department="测试",
        unit="m3",
        active_from=date(2026, 7, 1),
        enabled=True,
    )
    result = calculate_point(
        meter,
        date(2026, 7, 29),
        {
            date(2026, 7, 29): Observation(Decimal(10000), 1),
        },
    )
    assert (
        result is not None
        and result.value is None
        and result.status == "missing_reading"
    )


def test_detached_circulation_fields_are_left_for_review() -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "循环水码"
    sheet.append(["日期", "循环水表", None, None, "大地蒸汽内用表"])
    sheet.append(["7/1", 100, None, None, 999])
    preview = inspect_xlsx(content(book), 2026)
    assert preview["sheets"][0]["input_columns"] == [2]
    draft = extract_template(preview)
    assert [point.source_column for point in draft.meters] == [2]
    assert any("待确认" in issue for issue in draft.issues)


@pytest.mark.asyncio
async def test_confirm_rechecks_stored_date_window_before_writing_readings() -> None:
    from unittest.mock import Mock

    point_id = uuid4()
    definition = MeterDefinition(
        code="water",
        name="水表",
        category="drinking_water",
        area="测试",
        department="测试",
        unit="m3",
        enabled=True,
        active_from=date(2026, 7, 1),
        source_sheet="饮用水码",
        source_column=2,
    )
    batch = SimpleNamespace(
        id=uuid4(),
        status="preview",
        payload={
            "start_date": "2026-07-01",
            "end_date": "2026-07-01",
            "sheets": [
                {
                    "name": "饮用水码",
                    "headers": ["日期", "水表"],
                    "cells": [
                        {"date": "2026-07-01", "column": 2, "value": "100"},
                        {"date": "2026-07-02", "column": 2, "value": "200"},
                    ],
                }
            ],
        },
    )
    service = SimpleNamespace(
        actor_id=None,
        require_manager=AsyncMock(),
        revision=Mock(),
        repo=SimpleNamespace(
            serialize_write=AsyncMock(),
            get=AsyncMock(
                side_effect=[
                    batch,
                    SimpleNamespace(
                        id=point_id, definition=definition.model_dump(mode="json")
                    ),
                ]
            ),
        ),
        db=SimpleNamespace(
            scalar=AsyncMock(return_value=SimpleNamespace(value=Decimal(100))),
            flush=AsyncMock(),
            add=Mock(),
        ),
    )
    result = await confirm_import(
        service,
        batch.id,
        ImportConfirm(
            mappings=[{"sheet": "饮用水码", "column": 2, "point_id": point_id}],
            reason="范围测试",
        ),
    )
    assert result.status == "imported"
    assert service.db.scalar.await_count == 1
    assert service.db.scalar.call_args.args[0].compile().params[
        "reading_date_1"
    ] == date(2026, 7, 1)
    service.db.add.assert_not_called()


@pytest.mark.asyncio
async def test_preview_period_changes_its_idempotency_key() -> None:
    from unittest.mock import Mock

    from app.modules.energy.ledger_imports import preview_import

    batches = []
    service = SimpleNamespace(
        actor_id=None,
        require_manager=AsyncMock(),
        repo=SimpleNamespace(serialize_write=AsyncMock()),
        db=SimpleNamespace(
            scalar=AsyncMock(return_value=None),
            flush=AsyncMock(),
            add=Mock(side_effect=batches.append),
        ),
    )
    data = content(electricity_book())
    await preview_import(service, data, 2026, date(2026, 7, 1), date(2026, 7, 1))
    await preview_import(service, data, 2026, date(2026, 7, 1), date(2026, 7, 2))
    assert batches[0].content_hash != batches[1].content_hash
    assert {cell["date"] for cell in batches[0].payload["sheets"][0]["cells"]} == {
        "2026-07-01"
    }


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "sheet,category",
    [
        ("饮用水码", "drinking_water"),
        ("冰水码", "chilled_water"),
        ("蒸汽码", "steam"),
        ("循环水码", "circulating_water"),
        ("空气码", "air"),
        ("电量", "electricity"),
    ],
)
async def test_http_accepts_the_six_confirmed_manual_input_types(
    sheet: str, category: str
) -> None:
    from app.modules.energy.ledger_api import ledger_service, router
    from app.modules.energy.ledger_schemas import MeterResponse

    definition = MeterDefinition.model_validate(
        {
            "code": "allowed",
            "name": "东动力总" if category == "electricity" else "表码",
            "source_sheet": sheet,
            "source_column": 2,
            "category": category,
            "reading_type": "direct" if category == "electricity" else "counter",
            "area": "测试",
            "department": "测试",
            "unit": "kWh" if category == "electricity" else "m3",
            "active_from": "2026-07-01",
            "enabled": True,
        }
    )
    service = SimpleNamespace(
        save_meter=AsyncMock(
            return_value=MeterResponse(id=uuid4(), version=1, definition=definition)
        )
    )
    app = FastAPI()
    app.include_router(router, prefix="/api/v1/energy")
    app.dependency_overrides[ledger_service] = lambda: service
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        response = await client.post(
            "/api/v1/energy/meter-points",
            json={"definition": definition.model_dump(mode="json")},
        )
        assert response.status_code == 200, response.text
        assert (
            response.json()["data"]["definition"]["reading_type"]
            == definition.reading_type
        )


@pytest.mark.asyncio
async def test_snapshot_import_uses_the_same_period_and_manual_field_boundary() -> None:
    from unittest.mock import Mock

    from app.modules.energy.ledger_imports import preview_snapshot_import
    from app.modules.energy.ledger_schemas import SnapshotImportRequest

    snapshot_id = uuid4()
    snapshot = SimpleNamespace(
        id=snapshot_id,
        sheet_id=uuid4(),
        is_deleted=False,
        row_count=2,
        header_values=[
            "日期",
            "总计",
            "设备甲",
            "设备乙",
            "小计",
            "其他",
            "分表总和",
            "西动力总",
            "东动力总",
        ],
        content_hash="fixture",
    )
    rows = [
        SimpleNamespace(
            row_index=2, values=["7/1", 100, 30, None, 30, 70, 100, 40, 60]
        ),
        SimpleNamespace(row_index=3, values=["7/2", 120, 40, 20, 60, 60, 120, 50, 70]),
    ]
    service = SimpleNamespace(
        actor_id=None,
        require_manager=AsyncMock(),
        repo=SimpleNamespace(serialize_write=AsyncMock()),
        db=SimpleNamespace(
            get=AsyncMock(
                side_effect=[snapshot, SimpleNamespace(title="电量", is_deleted=False)]
            ),
            scalars=AsyncMock(return_value=SimpleNamespace(all=lambda: rows)),
            scalar=AsyncMock(return_value=None),
            add=Mock(),
            flush=AsyncMock(),
        ),
    )
    response = await preview_snapshot_import(
        service,
        SnapshotImportRequest(
            snapshot_ids=[snapshot_id],
            year=2026,
            start_date=date(2026, 7, 1),
            end_date=date(2026, 7, 1),
        ),
    )
    sheet = response.payload["sheets"][0]
    assert sheet["calculated_columns"] == [2, 5, 6, 7]
    assert sheet["input_columns"] == [3, 4, 8, 9]
    assert {cell["column"] for cell in sheet["cells"]} == {3, 8, 9}
    assert {cell["date"] for cell in sheet["cells"]} == {"2026-07-01"}
