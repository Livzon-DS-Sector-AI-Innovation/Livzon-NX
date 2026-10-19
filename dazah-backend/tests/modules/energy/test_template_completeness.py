from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import openpyxl
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.main import app
from app.modules.energy.ledger_imports import inspect_xlsx
from app.modules.energy.ledger_template import extract_template, stable_code
from app.platform.identity.deps import get_current_user
from app.platform.identity.models import User


def preview(workbook: openpyxl.Workbook):
    output = BytesIO()
    workbook.save(output)
    return inspect_xlsx(output.getvalue(), 2026)


def test_empty_named_source_is_preserved_in_subtotal_dependencies() -> None:
    book = openpyxl.Workbook()
    source = book.active
    assert source is not None
    source.title = "电量"
    source.append(["日期", "已填设备", "待核设备", "小计"])
    source.append(["7/1", 10, None, "=SUM(B2:C2)"])
    draft = extract_template(preview(book))
    assert len(draft.meters) == 2
    assert all(not point.enabled for point in draft.meters)
    assert {term.metric_key for term in draft.rules[0].terms} == {
        "point:" + stable_code("电量", 2),
        "point:" + stable_code("电量", 3),
    }
    assert any("待核设备" in issue and "无原始读数" in issue for issue in draft.issues)


def test_unnamed_output_keeps_source_identity_and_downstream_sum() -> None:
    book = openpyxl.Workbook()
    source = book.active
    assert source is not None
    source.title = "饮用水码"
    source.append(["日期", "103-2车间", "103-2车间"])
    source.append(["7/1", 100, 200])
    source.append(["7/2", 110, 220])
    usage = book.create_sheet("饮用水量")
    usage.append(["日期", "103-2车间", None, "分表和"])
    usage.append(
        [
            "7/1",
            "='饮用水码'!B3-'饮用水码'!B2",
            "='饮用水码'!C3-'饮用水码'!C2",
            "=B2+C2",
        ]
    )
    draft = extract_template(preview(book))
    unnamed = next(rule for rule in draft.rules if rule.source_column == 3)
    assert "103-2车间" in unnamed.name and "C列无表头" in unnamed.name
    assert unnamed.terms[0].metric_key == "point:" + stable_code("饮用水码", 3)
    subtotal = next(rule for rule in draft.rules if rule.source_column == 4)
    assert len(subtotal.terms) == 2


def test_counter_conversion_and_concrete_formula_variation_evidence() -> None:
    book = openpyxl.Workbook()
    source = book.active
    assert source is not None
    source.title = "蒸汽码"
    source.append(["日期", "余热回收"])
    source.append(["7/1", 100])
    source.append(["7/2", 110])
    source.append(["7/3", 120])
    usage = book.create_sheet("蒸汽量")
    usage.append(["日期", "余热折算", "排空量"])
    usage.append(["7/1", "=('蒸汽码'!B3-'蒸汽码'!B2)*1000*3600/2756.4/1000", "=B2"])
    usage.append(["7/2", "=('蒸汽码'!B4-'蒸汽码'!B3)*1000*3600/2756.4/1000", "=B2"])
    draft = extract_template(preview(book))
    conversion = next(rule for rule in draft.rules if rule.source_column == 2)
    assert conversion.terms[0].coefficient == Decimal("1.30605137")
    assert "2756.4" in conversion.notes
    assert any(
        "C3" in issue and "2026-07-02" in issue and "=B2" in issue
        for issue in draft.issues
    )


def test_cancelled_tank_columns_do_not_create_configuration_warnings() -> None:
    book = openpyxl.Workbook()
    source = book.active
    assert source is not None
    source.title = "电量"
    source.append(["日期", "101车间"])
    source.append(["7/1", 10])
    output = book.create_sheet("101车间")
    output.append(["日期", "电", "大罐数量"])
    output.append(["7/1", "='电量'!B2", "=B2"])
    draft = extract_template(preview(book))
    assert len(draft.rules) == 1
    assert not any("大罐" in issue for issue in draft.issues)


def test_blank_reading_rows_do_not_create_false_formula_changes() -> None:
    book = openpyxl.Workbook()
    source = book.active
    assert source is not None
    source.title = "饮用水码"
    source.append(["日期", "水表"])
    source.append(["7/1", 100])
    source.append(["7/2", None])
    usage = book.create_sheet("饮用水量")
    usage.append(["日期", "水量"])
    usage.append(["7/1", "='饮用水码'!B3-'饮用水码'!B2"])
    usage.append(["7/2", "='饮用水码'!B4-'饮用水码'!B3"])
    assert not extract_template(preview(book)).issues


def test_manual_output_zero_is_reviewed_and_not_created_as_a_raw_meter() -> None:
    book = openpyxl.Workbook()
    source = book.active
    assert source is not None
    source.title = "蒸汽码"
    source.append(["日期", "污水车间"])
    source.append(["7/1", 100])
    output = book.create_sheet("蒸汽量")
    output.append(["日期", "污水车间"])
    output.append(["7/1", 0])
    draft = extract_template(preview(book))
    assert len(draft.meters) == 1 and not draft.rules
    assert any("B2=0" in issue and "独立调整" in issue for issue in draft.issues)


@pytest.mark.asyncio
async def test_real_preview_and_draft_routes_preserve_empty_source(
    db_session: AsyncSession,
) -> None:
    actor = User(
        id=uuid4(),
        username="draft-" + uuid4().hex,
        name="能源配置管理员",
        role="admin",
        status="active",
        auth_source="local",
    )
    db_session.add(actor)
    await db_session.flush()

    async def user():
        return actor

    async def database():
        yield db_session

    book = openpyxl.Workbook()
    source = book.active
    assert source is not None
    source.title = "电量"
    source.append(["日期", "已填设备", "待核设备", "小计"])
    source.append(["7/1", 10, None, "=SUM(B2:C2)"])
    output = BytesIO()
    book.save(output)
    app.dependency_overrides[get_current_user] = user
    app.dependency_overrides[get_db] = database
    try:
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
            headers={"X-Dazah-Page-Key": "energy:settings"},
        ) as client:
            uploaded = await client.post(
                "/api/v1/energy/imports/preview",
                params={"year": 2026},
                files={
                    "file": (
                        "source.xlsx",
                        output.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert uploaded.status_code == 200, uploaded.text
            batch = uploaded.json()["data"]
            drafted = await client.post(
                f"/api/v1/energy/imports/{batch['id']}/draft-config"
            )
            assert drafted.status_code == 200, drafted.text
            assert len(drafted.json()["data"]["meters"]) == 2
            assert len(drafted.json()["data"]["rules"][0]["terms"]) == 2
            invalid = await client.post(
                "/api/v1/energy/imports/preview",
                params={"year": 2026},
                files={"file": ("bad.xlsx", b"invalid", "application/octet-stream")},
            )
            assert invalid.status_code == 422
    finally:
        app.dependency_overrides.pop(get_current_user, None)
        app.dependency_overrides.pop(get_db, None)
