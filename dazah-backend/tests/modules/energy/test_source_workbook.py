"""Original worksheet formulas must keep their cell/date semantics."""

from datetime import date
from decimal import Decimal

import pytest
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.energy.ledger_source_workbook import SourceWorkbook


def source(*sheets: tuple[str, list[list[str | int | None]]]) -> SourceWorkbook:
    return SourceWorkbook.from_payload(
        {
            "year": 2026,
            "sheets": [
                {
                    "name": name,
                    "headers": rows[0],
                    "source_grid": rows,
                    "row_dates": [
                        {"row": index, "date": f"2026-07-{index - 1:02}"}
                        for index in range(2, len(rows) + 1)
                    ],
                }
                for name, rows in sheets
            ],
        }
    )


def test_original_formula_keeps_previous_day_instead_of_normalizing_date() -> None:
    book = source(
        ("电量", [["日期", "设备"], ["7/1", 100], ["7/2", 200]]),
        ("101车间", [["日期", "电"], ["7/1", "='电量'!B2"], ["7/2", "='电量'!B2"]]),
    )
    cell = book.cell_on("101车间", 2, date(2026, 7, 2))
    assert cell.value == Decimal(100)
    assert cell.provenance["formula"] == "='电量'!B2"
    assert cell.provenance["source_cell"] == "B3"
    assert book.result("电量", "B3").value == Decimal(200)


def test_original_literal_override_is_not_recomputed_as_counter_difference() -> None:
    book = source(
        ("冰水码", [["日期", "表码"], ["7/1", 316278], ["7/2", 2365593]]),
        (
            "冰水量",
            [["日期", "用量"], ["7/1", 28674], ["7/2", "='冰水码'!B3-'冰水码'!B2"]],
        ),
    )
    assert book.result("冰水量", "B2").value == Decimal(28674)
    assert book.result("冰水量", "B2").provenance["source_kind"] == "literal"
    assert book.result("冰水量", "B3").value == Decimal(2049315)


def test_daily_formula_composition_and_negative_results_are_preserved() -> None:
    book = source(
        ("冰水量", [["日期", "7℃", "-5℃"], ["7/1", 10, 20], ["7/2", 30, 40]]),
        (
            "201-2车间",
            [["日期", "7℃冷水"], ["7/1", "='冰水量'!B2"], ["7/2", "='冰水量'!C3-50"]],
        ),
    )
    assert book.result("201-2车间", "B2").value == Decimal(10)
    assert book.result("201-2车间", "B3").value == Decimal(-10)
    assert book.result("201-2车间", "B3").status == "complete"


def test_full_precision_inputs_feed_conversion_and_do_not_use_cached_outputs() -> None:
    book = source(
        ("蒸汽码", [["日期", "余热"], ["7/1", "2421.170"], ["7/2", "2421.174"]]),
        (
            "蒸汽量",
            [["日期", "折算"], ["7/1", "=('蒸汽码'!B3-'蒸汽码'!B2)*3600/2756.4"]],
        ),
    )
    result = book.result("蒸汽量", "B2")
    assert result.value == Decimal("0.005224205485415759686547670874").quantize(
        Decimal(".00000001")
    )
    assert result.provenance["formula"].endswith("*3600/2756.4")


def test_blank_reference_uses_arithmetic_but_blank_cell_stays_blank() -> None:
    book = source(
        ("电量", [["日期", "设备", "空白", "总计"], ["7/1", 10, None, "=SUM(B2:C2)"]])
    )
    assert book.result("电量", "C2").value is None
    assert book.result("电量", "C2").status == "source_blank"
    assert book.result("电量", "D2").value == Decimal(10)
    assert book.cell_on("电量", 2, date(2026, 8, 1)).status == "not_applicable"


@pytest.mark.parametrize(
    "formula", ["=__import__('os')", "=B2**100000", '=INDIRECT("B2")', "='外部'!B2"]
)
def test_unsupported_functions_and_missing_sheets_are_explicit(formula: str) -> None:
    book = source(("电量", [["日期", "输入", "结果"], ["7/1", 10, formula]]))
    result = book.result("电量", "C2")
    assert result.value is None
    assert result.status in {"source_formula_error", "source_missing_sheet"}


def test_cycles_and_divide_by_zero_do_not_become_successful_zero() -> None:
    book = source(("日总量", [["日期", "环", "零分母"], ["7/1", "=B2+1", "=1/0"]]))
    assert book.result("日总量", "B2").status == "source_formula_cycle"
    assert book.result("日总量", "C2").status == "zero_denominator"


def test_reading_overlay_recalculates_all_dependent_dates() -> None:
    book = source(
        ("电量", [["日期", "设备"], ["7/1", 100], ["7/2", 200]]),
        ("101车间", [["日期", "电"], ["7/1", "='电量'!B2"], ["7/2", "='电量'!B2"]]),
    )
    book.set_reading("电量", 2, date(2026, 7, 1), Decimal(123), version=2)
    assert book.result("101车间", "B2").value == Decimal(123)
    assert book.result("101车间", "B3").value == Decimal(123)
    assert book.result("电量", "B3").value == Decimal(200)


def test_duplicate_dates_are_rejected_instead_of_choosing_an_arbitrary_row() -> None:
    with pytest.raises(ValueError, match="重复日期"):
        SourceWorkbook.from_payload(
            {
                "sheets": [
                    {
                        "name": "电量",
                        "headers": ["日期", "电"],
                        "source_grid": [["日期", "电"], ["7/1", 10], ["7/1", 20]],
                        "row_dates": [
                            {"row": 2, "date": "2026-07-01"},
                            {"row": 3, "date": "2026-07-01"},
                        ],
                    }
                ]
            }
        )


def test_xlsx_preview_preserves_literal_overrides_formulas_and_full_source_grid() -> (
    None
):
    from io import BytesIO

    import openpyxl

    from app.modules.energy.ledger_imports import inspect_xlsx

    original = openpyxl.Workbook()
    sheet = original.active
    assert sheet is not None
    sheet.title = "电量"
    sheet.append(["日期", "输入", "结果"])
    sheet.append(["7/1", "10.004", "=B2+1"])
    sheet.append(["7/2", None, 0])
    output = BytesIO()
    original.save(output)
    payload = inspect_xlsx(output.getvalue(), 2026)
    book = SourceWorkbook.from_payload(payload)
    assert book.result("电量", "B2").value == Decimal("10.004")
    assert book.result("电量", "C2").value == Decimal("11.004")
    assert book.result("电量", "C3").value == 0
    assert book.result("电量", "B3").value is None


@pytest.mark.asyncio
async def test_all_21_source_tables_through_real_http_routes_and_export(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from io import BytesIO
    from unittest.mock import AsyncMock

    import openpyxl
    from fastapi import FastAPI
    from httpx import ASGITransport, AsyncClient

    from app.core.exceptions import AppException
    from app.modules.energy import ledger_api
    from app.modules.energy.ledger_workbook import DATASETS, EnergyWorkbookService

    book = source(
        *[
            (item.name, [["日期", "原始或结果"], ["7/1", index], ["7/2", "=B2+1"]])
            for index, item in enumerate(DATASETS, 1)
        ]
    )

    async def load(workbook: EnergyWorkbookService, year: int, month: int) -> None:
        workbook.source_workbook = book
        workbook.source_batch_id = "reviewed-test-batch"
        workbook.loaded = year, month

    monkeypatch.setattr(EnergyWorkbookService, "load", load)
    api = FastAPI()
    api.include_router(ledger_api.router, prefix="/api/v1/energy")
    api.dependency_overrides[ledger_api.ledger_service] = lambda: AsyncMock()

    @api.exception_handler(AppException)
    async def application_error(_request: object, exc: AppException):
        from fastapi.responses import JSONResponse

        return JSONResponse(
            status_code=exc.status_code, content={"message": exc.message}
        )

    async with AsyncClient(
        transport=ASGITransport(app=api), base_url="http://test"
    ) as client:
        for index, item in enumerate(DATASETS, 1):
            response = await client.get(
                "/api/v1/energy/reporting/workbook",
                params={"year": 2026, "month": 7, "dataset": item.name},
            )
            assert response.status_code == 200
            report = response.json()["data"]
            assert Decimal(report["rows"][0]["cells"]["2"]["value"]) == index
            cell = report["rows"][1]["cells"]["2"]
            assert Decimal(cell["value"]) == index + 1
            assert cell["provenance"]["formula"] == "=B2+1"
            assert cell["provenance"]["source_cell"] == "B3"
        unknown = await client.get(
            "/api/v1/energy/reporting/workbook",
            params={"year": 2026, "month": 7, "dataset": "不存在"},
        )
        assert unknown.status_code == 404
        invalid = await client.get(
            "/api/v1/energy/reporting/workbook",
            params={"year": 2026, "month": 13, "dataset": "电量"},
        )
        assert invalid.status_code == 422
        response = await client.get(
            "/api/v1/energy/reporting/workbook/export",
            params={"year": 2026, "month": 7},
        )
        assert response.status_code == 200
        exported = openpyxl.load_workbook(BytesIO(response.content), data_only=True)
        assert len(exported.sheetnames) == 21
        for index, item in enumerate(DATASETS, 1):
            assert exported[item.name].cell(2, 2).value == index
            assert exported[item.name].cell(3, 2).value == index + 1
            assert exported[item.name].cell(4, 2).value is None
        exported.close()


@pytest.mark.asyncio
async def test_database_source_precedence_recalculation(
    db_session: AsyncSession,
) -> None:
    from datetime import UTC, datetime, timedelta
    from uuid import uuid4

    from app.modules.energy.ledger_models import (
        EnergyImportBatch,
        EnergyMeterPoint,
        EnergyReading,
    )
    from app.modules.energy.ledger_schemas import MeterDefinition
    from app.modules.energy.ledger_service import EnergyLedgerService
    from app.modules.energy.ledger_workbook import DATASETS, EnergyWorkbookService
    from app.platform.identity.models import User

    actor = User(
        id=uuid4(),
        username=uuid4().hex,
        name="原表验证",
        role="admin",
        status="active",
        auth_source="local",
    )
    db_session.add(actor)
    await db_session.flush()
    payload = {
        "year": 2026,
        "sheets": [
            {
                "name": item.name,
                "headers": ["日期", "测试字段"],
                "source_grid": [["日期", "测试字段"], ["7/1", 1], ["7/2", 2]],
                "row_dates": [
                    {"row": 2, "date": "2026-07-01"},
                    {"row": 3, "date": "2026-07-02"},
                ],
            }
            for item in DATASETS
        ],
    }
    payload["sheets"][2]["source_grid"][1][1] = "2421.170"
    payload["sheets"][2]["source_grid"][2][1] = "2421.174"
    payload["sheets"][7]["source_grid"][1][1] = "=('蒸汽码'!B3-'蒸汽码'!B2)*3600/2756.4"
    imported_at = datetime.now(UTC)
    batch = EnergyImportBatch(
        id=uuid4(),
        status="imported",
        content_hash=uuid4().hex * 2,
        payload=payload,
        created_by=actor.id,
        created_at=imported_at,
    )
    definition = MeterDefinition(
        code="source_" + uuid4().hex,
        name="测试余热",
        category="steam",
        area="测试",
        department="能源",
        unit="原表单位",
        active_from=date(2026, 7, 1),
        enabled=True,
        source_sheet="蒸汽码",
        source_column=2,
        quantity_kind="heat",
    )
    meter = EnergyMeterPoint(
        id=uuid4(),
        code=definition.code,
        version=1,
        definition=definition.model_dump(mode="json"),
    )
    db_session.add_all([batch, meter])
    await db_session.flush()
    reading = EnergyReading(
        id=uuid4(),
        point_id=meter.id,
        task_id=uuid4(),
        reading_date=date(2026, 7, 2),
        value=Decimal("2421.17"),
        version=1,
        source="import",
        submitted_at=imported_at - timedelta(seconds=1),
    )
    db_session.add(reading)
    await db_session.flush()
    service = EnergyLedgerService(db_session, actor)
    report = await EnergyWorkbookService(service).report(2026, 7, "蒸汽量")
    assert report.rows[0].cells["2"].value == Decimal("0.00522421")
    assert report.rows[0].cells["2"].provenance["source_batch_id"] == str(batch.id)
    reading.value = Decimal("2421.180")
    reading.submitted_at = imported_at + timedelta(seconds=1)
    reading.version = 2
    await db_session.flush()
    report = await EnergyWorkbookService(service).report(2026, 7, "蒸汽量")
    assert report.rows[0].cells["2"].value == (
        Decimal("0.010") * 3600 / Decimal("2756.4")
    ).quantize(Decimal(".00000001"))
