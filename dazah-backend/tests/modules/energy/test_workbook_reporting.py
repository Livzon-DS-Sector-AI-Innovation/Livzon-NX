from datetime import date
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import openpyxl
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.main import app
from app.modules.energy.ledger_imports import inspect_xlsx
from app.modules.energy.ledger_models import EnergyReading, EnergyUsageFact
from app.modules.energy.ledger_schemas import (
    MeterDefinition,
    MeterWrite,
    RuleDefinition,
    RuleTerm,
    RuleWrite,
)
from app.modules.energy.ledger_service import EnergyLedgerService
from app.modules.energy.ledger_template import extract_template
from app.platform.identity.deps import get_current_user
from app.platform.identity.models import User
from tests.modules.energy.workbook_fixtures import seed_meter, seed_rule


def test_equipment_name_temperature_is_not_an_electricity_dimension() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "电量"
    sheet.append(["日期", "7度冰水用电量", "动力用电量", "小计"])
    sheet.append(["7/1", 10, 20, "=SUM(B2:C2)"])
    output = BytesIO()
    workbook.save(output)
    draft = extract_template(inspect_xlsx(output.getvalue(), 2026))
    assert all(meter.temperature == "" for meter in draft.meters)
    assert all(rule.temperature == "" for rule in draft.rules)


@pytest.mark.asyncio
async def test_duplicate_active_source_mapping_is_visible_even_if_one_value_is_missing(
    db_session: AsyncSession,
) -> None:
    from app.modules.energy.ledger_workbook import EnergyWorkbookService

    workbook = EnergyWorkbookService(EnergyLedgerService(db_session))
    workbook.loaded = (2026, 7)
    workbook.headers["饮用水码"] = ["日期", "同一来源列"]
    for code in ["first", "second"]:
        workbook.definitions["point:" + code] = MeterDefinition(
            code=code,
            name=code,
            category="drinking_water",
            area="101",
            department="101车间",
            unit="m3",
            active_from=date(2026, 7, 1),
            enabled=True,
            source_sheet="饮用水码",
            source_column=2,
        )
    workbook.readings[("point:first", date(2026, 7, 1))] = EnergyReading(
        id=uuid4(),
        point_id=uuid4(),
        task_id=uuid4(),
        reading_date=date(2026, 7, 1),
        value=Decimal(0),
        source="import",
        version=1,
    )
    report = await workbook.report(2026, 7, "饮用水码")
    assert report.rows[0].cells["2"].status == "configuration_conflict"
    assert report.rows[0].cells["2"].value is None


@pytest.mark.asyncio
async def test_original_datasets_readings_usage_quality_filters_and_exports(
    db_session: AsyncSession,
) -> None:
    actor = User(
        id=uuid4(),
        username=uuid4().hex,
        name="报表测试",
        role="admin",
        status="active",
        auth_source="local",
    )
    db_session.add(actor)
    await db_session.flush()
    service = EnergyLedgerService(db_session, actor)
    day = date(2026, 7, 1)
    meter = await seed_meter(
        service,
        MeterWrite(
            definition=MeterDefinition(
                code="water_" + uuid4().hex,
                name="101水表",
                category="drinking_water",
                area="101",
                workshop="101车间",
                department="101车间",
                unit="m3",
                active_from=day,
                enabled=True,
                source_sheet="饮用水码",
                source_column=2,
            )
        ),
    )
    await seed_rule(
        service,
        RuleWrite(
            definition=RuleDefinition(
                code="water_usage_" + uuid4().hex,
                name="101用量",
                category="drinking_water",
                department="101车间",
                workshop="101车间",
                unit="m3",
                active_from=day,
                enabled=True,
                terms=[RuleTerm(metric_key="point:" + meter.definition.code)],
                source_sheet="饮用水量",
                source_column=2,
            )
        ),
    )
    for current, value in [
        (day, "100"),
        (date(2026, 7, 2), "110"),
        (date(2026, 7, 3), "0"),
    ]:
        db_session.add(
            EnergyReading(
                id=uuid4(),
                point_id=meter.id,
                task_id=uuid4(),
                reading_date=current,
                value=Decimal(value),
                source="import",
                version=1,
            )
        )
    await db_session.flush()
    await service.rebuild({day, date(2026, 7, 2)})

    async def user():
        return actor

    async def database():
        yield db_session

    app.dependency_overrides[get_current_user] = user
    app.dependency_overrides[get_db] = database
    try:
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
            headers={"X-Dazah-Page-Key": "energy:reports"},
        ) as client:
            catalog = await client.get("/api/v1/energy/reporting/datasets")
            assert catalog.status_code == 200, catalog.text
            assert len(catalog.json()["data"]) == 21
            metrics = await client.get("/api/v1/energy/reporting/metrics")
            assert metrics.status_code == 200, metrics.text
            assert any(
                item["metric_key"] == "point:" + meter.definition.code
                for item in metrics.json()["data"]
            )
            single = await client.get(
                "/api/v1/energy/reporting",
                params={
                    "start_date": "2026-07-01",
                    "end_date": "2026-07-01",
                    "metric_key": "point:" + meter.definition.code,
                },
            )
            assert single.status_code == 200, single.text
            assert single.json()["data"]["total_count"] == 1
            assert Decimal(single.json()["data"]["totals"][0]["value"]) == 10
            params = {"year": 2026, "month": 7, "dataset": "饮用水码"}
            raw = await client.get("/api/v1/energy/reporting/workbook", params=params)
            assert raw.status_code == 200, raw.text
            raw_data = raw.json()["data"]
            assert raw_data["rows"][0]["cells"]["2"]["value"] == "100.00000000"
            assert Decimal(raw_data["rows"][2]["cells"]["2"]["value"]) == 0
            assert raw_data["rows"][3]["cells"]["2"]["status"] == "missing_reading"
            assert "totals" not in raw_data
            usage = await client.get(
                "/api/v1/energy/reporting/workbook",
                params={**params, "dataset": "饮用水量"},
            )
            assert usage.status_code == 200, usage.text
            assert Decimal(usage.json()["data"]["rows"][0]["cells"]["2"]["value"]) == 10
            assert (
                usage.json()["data"]["rows"][1]["cells"]["2"]["status"]
                == "missing_dependency"
            )
            unknown = await client.get(
                "/api/v1/energy/reporting/workbook",
                params={**params, "dataset": "其他表"},
            )
            assert unknown.status_code == 404
            invalid = await client.get(
                "/api/v1/energy/reporting/workbook", params={**params, "month": 13}
            )
            assert invalid.status_code == 422
            exported = await client.get(
                "/api/v1/energy/reporting/workbook/export",
                params={"year": 2026, "month": 7},
            )
            assert exported.status_code == 200, exported.text[:200]
            book = openpyxl.load_workbook(BytesIO(exported.content), data_only=True)
            assert len(book.sheetnames) == 21
            assert book["饮用水码"].cell(2, 2).value == 100
            assert book["饮用水量"].cell(2, 2).value == 10
            assert book["饮用水码"].cell(4, 2).value == 0
            assert book["饮用水码"].cell(5, 2).value is None
            book.close()
            quality = await client.get(
                "/api/v1/energy/data-quality",
                params={
                    "start_date": "2026-07-01",
                    "end_date": "2026-07-03",
                    "workshop": "不存在车间",
                },
                headers={"X-Dazah-Page-Key": "energy:data-quality"},
            )
            assert quality.status_code == 200, quality.text
            assert quality.json()["data"]["total_count"] == 0
            rollback = await client.get(
                "/api/v1/energy/data-quality",
                params={
                    "start_date": "2026-07-01",
                    "end_date": "2026-07-03",
                    "quality_status": "reading_rollback",
                },
                headers={"X-Dazah-Page-Key": "energy:data-quality"},
            )
            assert rollback.status_code == 200
            assert rollback.json()["data"]["total_count"] == 1
            assert rollback.json()["data"]["quality_counts"]["reading_rollback"] == 1
    finally:
        app.dependency_overrides.pop(get_current_user, None)
        app.dependency_overrides.pop(get_db, None)


@pytest.mark.asyncio
async def test_abnormal_nonmissing_result_is_not_a_complete_period(
    db_session: AsyncSession,
) -> None:
    actor = User(
        id=uuid4(),
        username=uuid4().hex,
        name="质量测试",
        role="admin",
        status="active",
        auth_source="local",
    )
    db_session.add(actor)
    await db_session.flush()
    service = EnergyLedgerService(db_session, actor)
    meter = await seed_meter(
        service,
        MeterWrite(
            definition=MeterDefinition(
                code="abnormal_" + uuid4().hex,
                name="异常用量",
                category="drinking_water",
                area="101",
                department="101车间",
                unit="m3",
                active_from=date(2026, 7, 1),
                enabled=True,
                report_scope="detail",
            )
        ),
    )
    db_session.add(
        EnergyUsageFact(
            id=uuid4(),
            metric_key="point:" + meter.definition.code,
            usage_date=date(2026, 7, 1),
            definition=meter.definition.model_dump(mode="json"),
            calculated_value=Decimal(100),
            effective_value=Decimal(100),
            status="abnormal_usage",
            provenance={},
        )
    )
    await db_session.flush()
    report = await service.report(date(2026, 7, 1), date(2026, 7, 1))
    assert report.status == "partial"
    assert report.latest_complete_date is None
    assert report.totals[0].available_count == 0
