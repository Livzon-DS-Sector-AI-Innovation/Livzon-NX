from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import openpyxl
import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.exceptions import AppException
from app.main import app
from app.modules.energy.ledger_calculations import (
    Calculation,
    Observation,
    calculate_point,
    calculate_rule,
    validate_graph,
)
from app.modules.energy.ledger_imports import inspect_xlsx, parse_source_date
from app.modules.energy.ledger_models import (
    EnergyUsageFact,
)
from app.modules.energy.ledger_schemas import (
    MeterDefinition,
    MeterWrite,
    RuleDefinition,
    RuleTerm,
    RuleWrite,
)
from app.modules.energy.ledger_service import EnergyLedgerService
from app.modules.energy.ledger_template import extract_template
from app.platform.identity.deps import get_current_user
from app.platform.identity.models import User
from tests.modules.energy.workbook_fixtures import seed_meter, seed_rule

DAY = date(2026, 7, 1)


@pytest.mark.asyncio
async def test_reporting_basis_through_http_uses_all_roots_before_pagination(
    native_service: EnergyLedgerService,
) -> None:
    service = native_service
    definition = point(
        code="basis_" + uuid4().hex, source_sheet="饮用水码", source_column=14
    )
    await seed_meter(service, MeterWrite(definition=definition))
    total = RuleDefinition(
        code="basis_total_" + uuid4().hex,
        name="饮用水汇总",
        category=definition.category,
        department=definition.department,
        unit=definition.unit,
        active_from=DAY,
        report_scope="plant",
        enabled=True,
        terms=[RuleTerm(metric_key="point:" + definition.code)],
    )
    await seed_rule(service, RuleWrite(definition=total))
    for prefix, config in [("point:", definition), ("rule:", total)]:
        service.db.add(
            EnergyUsageFact(
                metric_key=prefix + config.code,
                usage_date=DAY,
                definition=config.model_dump(mode="json"),
                calculated_value=Decimal(12),
                effective_value=Decimal(12),
                status="complete",
                provenance={
                    "reading_dates": ["2026-07-01", "2026-07-02"],
                    "reading_values": ["100", "112"],
                }
                if prefix == "point:"
                else {},
            )
        )
    await service.db.flush()

    async def user() -> User:
        assert service.actor is not None
        return service.actor

    async def database():
        yield service.db

    app.dependency_overrides[get_current_user] = user
    app.dependency_overrides[get_db] = database
    try:
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
            headers={"X-Dazah-Page-Key": "energy:ledger-overview"},
        ) as client:
            response = await client.get(
                "/api/v1/energy/reporting",
                params={
                    "start_date": str(DAY),
                    "end_date": str(DAY + timedelta(days=2)),
                    "report_scope": "plant",
                    "include_basis": "true",
                    "page_size": 1,
                },
            )
            assert response.status_code == 200, response.text
            basis = response.json()["data"]["totals"][0]["basis"]
            assert basis["usage_date"] == str(DAY)
            assert basis["nodes"][1]["source_column"] == "N"
            assert Decimal(basis["nodes"][1]["value"]) == Decimal(12)
            assert basis["nodes"][1]["reading_values"] == ["100", "112"]
            invalid = await client.get(
                "/api/v1/energy/reporting",
                params={
                    "start_date": "2026-07-02",
                    "end_date": "2026-07-01",
                    "include_basis": "true",
                },
            )
            assert invalid.status_code == 422
    finally:
        app.dependency_overrides.pop(get_current_user, None)
        app.dependency_overrides.pop(get_db, None)


def test_historical_rule_can_preserve_a_not_yet_effective_dependency() -> None:
    rule = RuleDefinition(
        code="historical_total",
        name="历史汇总",
        category="drinking_water",
        department="101车间",
        unit="m3",
        active_from=DAY,
        terms=[RuleTerm(metric_key="rule:future_source")],
        enabled=False,
    )
    with pytest.raises(ValueError, match="来源不存在"):
        validate_graph({}, {"rule:historical_total": rule})
    assert validate_graph({}, {"rule:historical_total": rule}, allow_missing=True) == [
        "rule:historical_total"
    ]
    result = calculate_rule(rule, {})
    assert result.value is None
    assert result.status == "missing_dependency"


def point(**updates: object) -> MeterDefinition:
    return MeterDefinition.model_validate(
        {
            "code": "water",
            "name": "101车间水表",
            "category": "drinking_water",
            "area": "101",
            "department": "101车间",
            "unit": "m3",
            "active_from": DAY,
            "enabled": True,
            "report_scope": "source",
            **updates,
        }
    )


def test_counter_uses_next_day_and_preserves_zero() -> None:
    with pytest.raises(ValueError, match="直接用量按日填写"):
        point(reading_type="direct", frequency="monthly")
    meter = point(multiplier="10")
    result = calculate_point(
        meter,
        DAY,
        {
            DAY: Observation(Decimal(100), 1),
            DAY + timedelta(days=1): Observation(Decimal(100), 2),
        },
    )
    assert result is not None and result.value == 0 and result.status == "complete"
    result = calculate_point(
        meter,
        DAY,
        {
            DAY: Observation(Decimal(100), 1),
            DAY + timedelta(days=1): Observation(Decimal("101.2"), 2),
        },
    )
    assert result is not None and result.value == Decimal(12)
    assert result.provenance["reading_dates"] == ["2026-07-01", "2026-07-02"]


def test_missing_gap_is_not_allocated_to_a_day() -> None:
    result = calculate_point(
        point(),
        DAY,
        {
            DAY: Observation(Decimal(100), 1),
            DAY + timedelta(days=2): Observation(Decimal(120), 1),
        },
    )
    assert (
        result is not None
        and result.value is None
        and result.status == "missing_reading"
    )


def test_monthly_counter_only_emits_actual_month_interval() -> None:
    meter = point(frequency="monthly")
    readings = {
        DAY: Observation(Decimal(100), 1),
        date(2026, 8, 1): Observation(Decimal(160), 1),
    }
    assert calculate_point(meter, DAY, readings) is None
    result = calculate_point(meter, date(2026, 7, 31), readings)
    assert result is not None and result.value == 60
    assert result.provenance["interval_end"] == "2026-08-01"


def test_rollback_requires_meter_event_and_does_not_take_absolute_value() -> None:
    readings = {
        DAY: Observation(Decimal(100), 1),
        DAY + timedelta(days=1): Observation(Decimal(10), 1),
    }
    result = calculate_point(point(), DAY, readings)
    assert (
        result is not None
        and result.value is None
        and result.status == "reading_rollback"
    )
    readings[DAY + timedelta(days=1)] = Observation(
        Decimal(10),
        1,
        {
            "kind": "replacement",
            "previous_final": "110",
            "new_initial": "0",
            "reason": "换表",
        },
    )
    result = calculate_point(point(), DAY, readings)
    assert result is not None and result.value == 20


def test_rule_missing_zero_denominator_and_decimal_sum() -> None:
    rule = RuleDefinition(
        code="total",
        name="全厂水量",
        category="drinking_water",
        unit="m3",
        department="101车间",
        active_from=DAY,
        terms=[RuleTerm(metric_key="point:water", coefficient=Decimal(2))],
    )
    assert calculate_rule(rule, {}).value is None
    assert calculate_rule(
        rule, {"point:water": Calculation(Decimal("0.1"), "complete", {})}
    ).value == Decimal("0.2")
    ratio = rule.model_copy(
        update={
            "quantity_kind": "ratio",
            "operation": "ratio",
            "terms": [
                RuleTerm(metric_key="point:water"),
                RuleTerm(metric_key="point:water"),
            ],
        }
    )
    assert (
        calculate_rule(
            ratio, {"point:water": Calculation(Decimal(0), "complete", {})}
        ).status
        == "zero_denominator"
    )


def test_graph_rejects_cycles_units_and_duplicate_aggregation() -> None:
    meter = point(report_scope="detail")
    rule = RuleDefinition(
        code="total",
        name="车间水量",
        category="drinking_water",
        unit="m3",
        department="101车间",
        active_from=DAY,
        enabled=True,
        terms=[RuleTerm(metric_key="point:water")],
    )
    with pytest.raises(ValueError, match="同一汇总"):
        validate_graph({"point:water": meter}, {"rule:total": rule})
    rule = rule.model_copy(update={"unit": "t"})
    with pytest.raises(ValueError, match="单位"):
        validate_graph({"point:water": point()}, {"rule:total": rule})
    cycle = rule.model_copy(
        update={"enabled": False, "terms": [RuleTerm(metric_key="rule:total")]}
    )
    with pytest.raises(ValueError, match="循环"):
        validate_graph({}, {"rule:total": cycle})


def test_shared_meter_requires_allocation_and_large_calculations_stay_visible() -> None:
    rules = {
        "rule:first": RuleDefinition(
            code="first",
            name="一车间",
            category="drinking_water",
            unit="m3",
            department="101车间",
            active_from=DAY,
            enabled=True,
            terms=[RuleTerm(metric_key="point:water")],
        ),
        "rule:second": RuleDefinition(
            code="second",
            name="二车间",
            category="drinking_water",
            unit="m3",
            department="101车间",
            active_from=DAY,
            enabled=True,
            terms=[RuleTerm(metric_key="point:water")],
        ),
    }
    with pytest.raises(ValueError, match="重复统计"):
        validate_graph({"point:water": point()}, rules)
    for rule in rules.values():
        rule.terms[0].coefficient = Decimal("0.5")
    assert len(validate_graph({"point:water": point()}, rules)) == 2
    large = calculate_point(
        point(multiplier="10000000000000000000"),
        DAY,
        {
            DAY: Observation(Decimal(0), 1),
            DAY + timedelta(days=1): Observation(Decimal(100), 1),
        },
    )
    assert large is not None and large.value is None and large.status == "out_of_range"


def workbook_bytes() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "饮用水码"
    sheet.append(["日期", "101车间", "101车间"])
    sheet.append(["7/1", 100, 200])
    sheet.append(["7/2", 110, 215])
    output = workbook.create_sheet("饮用水量")
    output.append(["日期", "101车间"])
    output.append(["7/1", "='饮用水码'!B3-'饮用水码'!B2"])
    workshop = workbook.create_sheet("101车间")
    workshop.append(["日期", "饮用水", "大罐数量"])
    workshop.append(["7/1", "='饮用水量'!B2", 1])
    target = BytesIO()
    workbook.save(target)
    return target.getvalue()


def test_import_preview_no_cache_and_disabled_template_without_tank() -> None:
    payload = inspect_xlsx(workbook_bytes(), 2026)
    draft = extract_template(payload)
    assert len(draft.meters) == 2
    assert draft.meters[0].code != draft.meters[1].code
    assert all(not meter.enabled for meter in draft.meters)
    assert any(rule.workshop == "101车间" for rule in draft.rules)
    assert all("大罐数量" not in rule.name for rule in draft.rules)
    assert all(rule.terms for rule in draft.rules)
    assert parse_source_date("7/1", 2025) == date(2025, 7, 1)


@pytest.fixture
async def native_service(db_session: AsyncSession) -> EnergyLedgerService:
    actor = User(
        id=uuid4(),
        username="native-" + uuid4().hex,
        name="能源测试管理员",
        role="admin",
        status="active",
        auth_source="local",
    )
    db_session.add(actor)
    await db_session.flush()
    return EnergyLedgerService(db_session, actor)


@pytest.mark.asyncio
async def test_report_missing_dates_remain_visible_without_department_scope(
    native_service: EnergyLedgerService,
) -> None:
    service = native_service
    await seed_meter(
        service, MeterWrite(definition=point(code="missing_" + uuid4().hex))
    )
    report = await service.report(DAY, DAY + timedelta(days=2), report_scope="source")
    assert report.status == "pending" and report.total_count == 3
    assert all(row.effective_value is None for row in report.rows)


@pytest.mark.asyncio
async def test_import_conflict_does_not_overwrite_and_repeat_is_idempotent(
    native_service: EnergyLedgerService,
) -> None:
    from app.modules.energy.ledger_imports import confirm_import, preview_import
    from app.modules.energy.ledger_models import EnergyReading
    from app.modules.energy.ledger_schemas import ImportConfirm, ImportMapping

    service = native_service
    meter = await seed_meter(
        service,
        MeterWrite(
            definition=point(
                code="import_" + uuid4().hex, source_sheet="饮用水码", source_column=2
            )
        ),
    )
    batch = await preview_import(service, workbook_bytes(), 2026)
    request = ImportConfirm(
        mappings=[ImportMapping(sheet="饮用水码", column=2, point_id=meter.id)],
        reason="核对历史抄表",
    )
    imported = await confirm_import(service, batch.id, request)
    assert imported.status == "imported" and imported.payload["imported_count"] == 2
    original_usage = next(
        value
        for key, _, value in await service.evaluate(DAY)
        if key == "point:" + meter.definition.code
    )
    assert original_usage.value == Decimal(10)
    assert original_usage.provenance["late"] is False

    imported_readings = list((await service.db.scalars(select(EnergyReading))).all())
    assert imported_readings and all(
        row.source == "import" for row in imported_readings
    )
    again = await confirm_import(service, batch.id, request)
    assert again.payload["imported_count"] == 2
    readings = list(
        (
            await service.db.scalars(
                select(EnergyReading).where(EnergyReading.point_id == meter.id)
            )
        ).all()
    )
    assert len(readings) == 2
    readings[0].value = Decimal(101)
    await service.db.flush()
    with pytest.raises(AppException) as conflict:
        await confirm_import(
            service, batch.id, request.model_copy(update={"reason": "再次核对"})
        )
    assert conflict.value.status_code == 409 and readings[0].value == 101
