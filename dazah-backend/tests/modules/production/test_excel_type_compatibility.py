from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

from app.modules.production.dr_schedule_api import _pick_latest_sheet
from app.modules.production.schedule_excel_service import parse_workbook_bytes


def test_typed_workbook_keeps_visible_sheet_dates_and_merges():
    workbook = Workbook()
    old = workbook.worksheets[0]
    old.title = "2026.06.01"
    old.sheet_state = "hidden"
    sheet = workbook.create_sheet("2026.07.01")
    sheet.append(["日期", "批次"])
    sheet.append([datetime(2026, 7, 1), "B001"])
    sheet.merge_cells("B2:C2")
    assert _pick_latest_sheet(workbook).title == "2026.07.01"
    output = BytesIO()
    workbook.save(output)
    result = parse_workbook_bytes(output.getvalue())
    assert result["sheet_name"] == "2026.07.01"
    assert result["rows"][1][:2] == ["2026-07-01", "B001"]
    assert result["merges"] == [{"s": {"r": 1, "c": 1}, "e": {"r": 1, "c": 2}}]
