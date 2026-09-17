"""Import/export keeps the workbook boundary and persisted result aligned."""

from __future__ import annotations

from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from openpyxl import Workbook, load_workbook

from app.modules.registration.service import declaration_progress as progress


@pytest.mark.asyncio
async def test_import_and_export_respect_workbook_permission_and_content(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    template = progress.DECLARATION_PROGRESS_SHEET_TEMPLATES[0]
    definition = progress.DeclarationProgressSheetDefinition(
        worksheet_title=template.worksheet_title,
        sheet_key=template.sheet_key,
        sheet_name=template.sheet_name,
        sheet_title=template.sheet_title,
        supports_sub_records=template.supports_sub_records,
        columns=progress._build_columns(template),
    )
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    sheet = workbook.create_sheet(template.worksheet_title)
    sheet.cell(5, 1).value = "旧数据"
    stream = BytesIO()
    workbook.save(stream)
    content = stream.getvalue()
    workbook.close()

    config_path = tmp_path / "configured" / "declaration.xlsx"
    monkeypatch.setattr(progress, "_get_workbook_path", lambda: config_path)
    monkeypatch.setattr(
        progress, "read_upload_secure",
        AsyncMock(return_value=("declaration.xlsx", content)),
    )
    permissions: list[str] = []

    @asynccontextmanager
    async def authorize(_session: object, workbook_key: str, action: str):
        assert workbook_key == "declaration-progress"
        permissions.append(action)
        yield

    monkeypatch.setattr(progress, "authorized_workbook", authorize)

    def read_versions(path: Path):
        assert path.read_bytes() == content
        return [SimpleNamespace(sheet_key=definition.sheet_key, version_number=1)]

    monkeypatch.setattr(progress, "_load_versions_from_workbook_path", read_versions)
    monkeypatch.setattr(
        progress, "_parse_workbook_definitions", lambda: ([definition], None)
    )
    session = SimpleNamespace(commit=AsyncMock(), rollback=AsyncMock())
    service = progress.DeclarationProgressWorkbookService(session)
    service.repository = SimpleNamespace(
        count_versions=AsyncMock(return_value=1),
        replace_all_versions=AsyncMock(),
        list_active_versions_by_sheet=AsyncMock(return_value=[]),
    )

    imported = await service.import_workbook(
        SimpleNamespace(filename="declaration.xlsx")
    )
    assert imported.imported_records == 1
    assert imported.sheet_record_counts == {definition.sheet_key: 1}
    assert config_path.read_bytes() == content
    service.repository.replace_all_versions.assert_awaited_once()
    session.commit.assert_awaited_once()

    export_dir = tmp_path / "export"

    def make_export_dir(*, prefix: str) -> str:
        assert prefix == "declaration-progress-export-"
        export_dir.mkdir()
        return str(export_dir)

    monkeypatch.setattr(progress.tempfile, "mkdtemp", make_export_dir)
    exported_path, name = await service.export_workbook()
    assert name == progress.DECLARATION_PROGRESS_EXPORT_NAME
    exported = load_workbook(exported_path)
    try:
        assert "Sheet1" not in exported.sheetnames
        assert exported[template.worksheet_title].cell(5, 1).value is None
    finally:
        exported.close()
    assert permissions == ["bulk_import", "sensitive_export"]


@pytest.mark.asyncio
async def test_import_rolls_back_failed_replace(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    @asynccontextmanager
    async def authorize(_session: object, _workbook_key: str, _action: str):
        yield

    monkeypatch.setattr(progress, "authorized_workbook", authorize)
    monkeypatch.setattr(
        progress,
        "read_upload_secure",
        AsyncMock(return_value=("declaration.xlsx", b"test")),
    )
    monkeypatch.setattr(
        progress, "_load_versions_from_workbook_path", lambda _path: []
    )
    monkeypatch.setattr(
        progress, "_get_workbook_path", lambda: tmp_path / "unused.xlsx"
    )
    session = SimpleNamespace(commit=AsyncMock(), rollback=AsyncMock())
    service = progress.DeclarationProgressWorkbookService(session)
    service.repository = SimpleNamespace(
        count_versions=AsyncMock(return_value=1),
        replace_all_versions=AsyncMock(side_effect=RuntimeError("write failed")),
    )

    with pytest.raises(RuntimeError, match="write failed"):
        await service.import_workbook(SimpleNamespace(filename="declaration.xlsx"))
    session.rollback.assert_awaited_once()
    session.commit.assert_not_awaited()
