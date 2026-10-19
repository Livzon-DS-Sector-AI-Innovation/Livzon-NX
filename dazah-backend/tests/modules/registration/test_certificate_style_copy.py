from copy import copy

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side

from app.modules.registration.service.certificate import _copy_row_style


def test_certificate_copies_concrete_cell_styles_without_sharing_proxies():
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    source = sheet.cell(1, 1)
    source.font = Font(name="宋体", bold=True)
    source.fill = PatternFill("solid", fgColor="00FF00")
    source.border = Border(bottom=Side(style="thin"))
    source.alignment = Alignment(horizontal="center")
    source.protection = Protection(locked=False)
    source.number_format = "0.000"
    sheet.row_dimensions[1].height = 30
    _copy_row_style(sheet, 1, 2, 1)
    target = sheet.cell(2, 1)
    for attribute in ["font", "fill", "border", "alignment", "protection"]:
        assert copy(getattr(source, attribute)) == copy(getattr(target, attribute))
    assert target.number_format == "0.000"
    assert sheet.row_dimensions[2].height == 30
    target.font = Font(name="Arial")
    assert source.font.name == "宋体"
