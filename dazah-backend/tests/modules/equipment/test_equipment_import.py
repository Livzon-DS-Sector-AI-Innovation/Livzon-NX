"""设备台账 Excel 导入测试。"""

import io
import json
import uuid
from datetime import datetime
from typing import Any

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook

from app.modules.equipment.service.equipment_import import (
    build_ledger_template,
    parse_date_cell,
    parse_tech_params_cell,
)

HEADERS = [
    "序号",
    "设备编号",
    "设备名称",
    "规格型号",
    "技术参数",
    "生产厂家",
    "出厂日期",
    "出厂编号",
    "入厂日期",
    "安装地点",
    "备注",
]


def _ledger_bytes(rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _uid() -> str:
    return uuid.uuid4().hex[:6].upper()


async def _prepare_category_and_location(
    client: AsyncClient, location_name: str | None = None
) -> tuple[str, str]:
    uid = _uid()
    category = await client.post(
        "/api/v1/equipment/categories",
        json={"name": "反应设备", "code": f"IMP-CAT-{uid}"},
    )
    category_id = category.json()["data"]["id"]
    location_id = ""
    if location_name:
        location = await client.post(
            "/api/v1/equipment/locations",
            json={"name": location_name, "code": f"IMP-LOC-{uid}"},
        )
        location_id = location.json()["data"]["id"]
    return category_id, location_id


# ---------- 纯解析函数 ----------


def test_parse_date_cell_accepts_common_formats() -> None:
    assert parse_date_cell(datetime(2024, 5, 1, 8, 30)) == datetime(2024, 5, 1).date()
    assert str(parse_date_cell("2024-05-01")) == "2024-05-01"
    assert str(parse_date_cell("2024/5/1")) == "2024-05-01"
    assert str(parse_date_cell("2024.5.1")) == "2024-05-01"
    assert str(parse_date_cell(20240501)) == "2024-05-01"
    assert str(parse_date_cell("2024年5月1日")) == "2024-05-01"
    assert str(parse_date_cell("2024年5月")) == "2024-05-01"
    assert str(parse_date_cell("2024-05")) == "2024-05-01"
    assert parse_date_cell(None) is None
    assert parse_date_cell("  ") is None


def test_parse_date_cell_accepts_excel_serial_numbers() -> None:
    # Excel 日期序列号：46023 = 2026-01-01，45292 = 2024-01-01
    assert str(parse_date_cell(46023)) == "2026-01-01"
    assert str(parse_date_cell(46023.0)) == "2026-01-01"
    assert str(parse_date_cell(45292)) == "2024-01-01"


def test_parse_date_cell_accepts_year_month_numbers() -> None:
    # 数值型「年.月」：Excel 把 2017.11 存为数字
    assert str(parse_date_cell(2017.11)) == "2017-11-01"
    assert str(parse_date_cell(2020.05)) == "2020-05-01"
    # 尾零月份（2024.10）被存成 2024.1
    assert str(parse_date_cell(2024.1)) == "2024-10-01"
    assert str(parse_date_cell(2017)) == "2017-01-01"  # 纯年份


def test_parse_date_cell_treats_placeholder_as_empty() -> None:
    assert parse_date_cell("-") is None
    assert parse_date_cell("—") is None
    assert parse_date_cell("无") is None
    assert parse_date_cell("/") is None
    # 尾部多余的句点：2012.10.
    assert str(parse_date_cell("2012.10.")) == "2012-10-01"


def test_parse_date_cell_rejects_out_of_range_numbers() -> None:
    for invalid in (1, 123456, 99, 20241301):
        try:
            parse_date_cell(invalid)
        except ValueError:
            continue
        raise AssertionError(f"expected ValueError for {invalid}")


def test_parse_date_cell_rejects_unknown_text() -> None:
    try:
        parse_date_cell("2024年五月")
    except ValueError as error:
        assert "日期格式无法识别" in str(error)
    else:
        raise AssertionError("expected ValueError")


def test_parse_tech_params_cell_splits_pairs() -> None:
    assert parse_tech_params_cell("功率：3kW；容积 = 500L\n设计压力: 0.3MPa") == {
        "功率": "3kW",
        "容积": "500L",
        "设计压力": "0.3MPa",
    }
    assert parse_tech_params_cell(None) is None
    assert parse_tech_params_cell("") is None


def test_parse_tech_params_cell_keeps_separator_free_fragment() -> None:
    assert parse_tech_params_cell("见铭牌") == {"参数1": "见铭牌"}


def test_template_contains_ledger_headers_and_example() -> None:
    sheet = load_workbook(io.BytesIO(build_ledger_template())).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == tuple(HEADERS)
    assert rows[1][1] == "SB-0001"


# ---------- API 导入 ----------


async def test_import_creates_updates_and_reports_failures(client: AsyncClient) -> Any:
    category_id, location_id = await _prepare_category_and_location(
        client, location_name="201二车间"
    )
    excel = _ledger_bytes(
        [
            [
                1, "IMP-001", "发酵罐", "PJ-500L", "功率：3kW",
                "温州设备厂", "2024-05-01", "CC-1", "2024-06-15",
                "201二车间", "第一行",
            ],
            [
                2, "IMP-002", "离心机", None, None, None,
                None, None, None, "未知车间", None,
            ],
            [
                3, "IMP-003", "干燥机", None, None, None,
                "五月", None, None, "201二车间", None,
            ],
            [4, "", "缺编号设备", None, None, None, None, None, None, None, None],
        ]
    )

    response = await client.post(
        "/api/v1/equipment/equipments/import",
        files={
            "file": (
                "ledger.xlsx",
                excel,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"default_category_id": category_id, "default_location_id": location_id},
    )
    assert response.status_code == 200
    summary = response.json()["data"]
    assert summary["total_rows"] == 4
    assert summary["created"] == 4
    assert summary["updated"] == 0
    # 坏日期/空编号不再拒行：导入并标红待修正
    assert summary["failed"] == 0
    assert summary["flagged"] == 2

    # 校验落库内容：匹配到的安装地点、默认地点、技术参数与台账字段
    list_response = await client.get(
        "/api/v1/equipment/equipments",
        params={"keyword": "IMP-", "page": 1, "page_size": 20},
    )
    items = {
        item["equipment_no"]: item for item in list_response.json()["data"]
    }
    assert items["IMP-001"]["technical_params"] == {"功率": "3kW"}
    assert items["IMP-001"]["factory_no"] == "CC-1"
    assert items["IMP-001"]["arrival_date"] == "2024-06-15"
    assert items["IMP-001"]["location_name"] == "201二车间"
    # 匹配不上的非空安装地点自动新建位置（默认地点只兜空值）
    assert items["IMP-002"]["location_name"] == "未知车间"
    assert summary["locations_created"] == 1
    assert items["IMP-001"]["category_names"] == "反应设备"
    # 坏日期：置空并记录原文待修正（【字段】前缀供页面定位单元格标红）
    assert items["IMP-003"]["production_date"] is None
    assert "【出厂日期】原文「五月」无法识别" in items["IMP-003"]["data_issue_note"]
    # 空编号：生成确定性临时编号（WB- 前缀）并记录待修正
    all_items = (
        await client.get(
            "/api/v1/equipment/equipments",
            params={"page": 1, "page_size": 50},
        )
    ).json()["data"]
    empty_no_items = [
        item for item in all_items if item["equipment_no"].startswith("WB-")
    ]
    assert len(empty_no_items) == 1
    assert "【设备编号】为空" in empty_no_items[0]["data_issue_note"]
    assert empty_no_items[0]["name"] == "缺编号设备"

    # 重复导入同一文件 → 全部转为更新（临时编号确定性生成，不产生重复）
    repeat = await client.post(
        "/api/v1/equipment/equipments/import",
        files={
            "file": (
                "ledger.xlsx",
                excel,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"default_category_id": category_id, "default_location_id": location_id},
    )
    repeat_summary = repeat.json()["data"]
    assert repeat_summary["created"] == 0
    assert repeat_summary["updated"] == 4
    assert repeat_summary["failed"] == 0


async def test_import_fills_merged_cells_and_joins_multirow_equipment(
    client: AsyncClient,
) -> Any:
    """真实台账常见形态：编号/名称纵向合并，参数、备注分行填写。"""
    _, location_id = await _prepare_category_and_location(
        client, location_name="201二车间"
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    # 行2~4：同一设备（编号/名称合并），技术参数、备注分行
    sheet.append(
        [1, None, None, None, "功率：3kW", None,
         "2024-05-01", None, None, "201二车间", None]
    )
    sheet.append(
        [None, None, None, None, "转速：3000rpm", None, None, None, None, None, None]
    )
    sheet.append(
        [None, None, None, None, None, None, None, None, None, None, "含备用转子"]
    )
    sheet["B2"] = "IMP-300"
    sheet.merge_cells("B2:B4")  # 设备编号合并（B 列，A 列是序号）
    sheet["C2"] = "离心机组"
    sheet.merge_cells("C2:C4")  # 设备名称合并
    buffer = io.BytesIO()
    workbook.save(buffer)

    response = await client.post(
        "/api/v1/equipment/equipments/import",
        files={
            "file": (
                "ledger.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"default_location_id": location_id},
    )
    assert response.status_code == 200
    summary = response.json()["data"]
    assert summary["created"] == 1
    assert summary["failed"] == 0

    items = (
        await client.get(
            "/api/v1/equipment/equipments",
            params={"keyword": "IMP-300", "page": 1, "page_size": 20},
        )
    ).json()["data"]
    assert len(items) == 1
    equipment = items[0]
    assert equipment["equipment_no"] == "IMP-300"
    assert equipment["name"] == "离心机组"
    # 分行的技术参数拼接进同一设备，备注同理
    assert equipment["technical_params"] == {"功率": "3kW", "转速": "3000rpm"}
    assert equipment["description"] == "含备用转子"
    assert equipment["location_name"] == "201二车间"
    assert str(equipment["production_date"]) == "2024-05-01"


async def test_import_without_category_creates_uncategorized(
    client: AsyncClient,
) -> Any:
    _, location_id = await _prepare_category_and_location(
        client, location_name="201二车间"
    )
    excel = _ledger_bytes(
        [
            [
                1, "IMP-201", "压滤机", None, None, None,
                None, None, None, "201二车间", None,
            ]
        ]
    )
    response = await client.post(
        "/api/v1/equipment/equipments/import",
        files={
            "file": (
                "ledger.xlsx",
                excel,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"default_location_id": location_id},
    )
    assert response.status_code == 200
    summary = response.json()["data"]
    assert summary["created"] == 1
    assert summary["failed"] == 0
    items = {
        item["equipment_no"]: item
        for item in (
            await client.get(
                "/api/v1/equipment/equipments",
                params={"keyword": "IMP-201", "page": 1, "page_size": 20},
            )
        ).json()["data"]
    }
    assert items["IMP-201"]["category_names"] is None
    assert items["IMP-201"]["location_name"] == "201二车间"


async def test_import_auto_creates_unmatched_location(client: AsyncClient) -> Any:
    """安装地点为台账自由文本时自动新建为顶层位置。"""
    category_id, _ = await _prepare_category_and_location(client)
    excel = _ledger_bytes(
        [
            [
                1, "IMP-101", "过滤机", None, None, None,
                None, None, None, "2#萃取滤液罐上安装", None,
            ],
            [
                2, "IMP-102", "离心机", None, None, None,
                None, None, None, "2#萃取滤液罐上安装", None,
            ],
        ]
    )
    response = await client.post(
        "/api/v1/equipment/equipments/import",
        files={
            "file": (
                "ledger.xlsx",
                excel,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"default_category_id": category_id},
    )
    summary = response.json()["data"]
    assert summary["created"] == 2
    assert summary["failed"] == 0
    # 同一地点文本只新建一次
    assert summary["locations_created"] == 1
    items = {
        item["equipment_no"]: item
        for item in (
            await client.get(
                "/api/v1/equipment/equipments",
                params={"keyword": "IMP-10", "page": 1, "page_size": 20},
            )
        ).json()["data"]
    }
    assert items["IMP-101"]["location_name"] == "2#萃取滤液罐上安装"
    assert items["IMP-102"]["location_name"] == "2#萃取滤液罐上安装"
    # 新位置进入位置树，可被后续导入复用
    locations = (
        await client.get("/api/v1/equipment/locations", params={"tree": True})
    ).json()["data"]
    flat_names = json.dumps(locations, ensure_ascii=False)
    assert "2#萃取滤液罐上安装" in flat_names


async def test_import_fails_when_location_empty_and_no_default(
    client: AsyncClient,
) -> Any:
    category_id, _ = await _prepare_category_and_location(client)
    excel = _ledger_bytes(
        [[1, "IMP-103", "干燥机", None, None, None, None, None, None, None, None]]
    )
    response = await client.post(
        "/api/v1/equipment/equipments/import",
        files={
            "file": (
                "ledger.xlsx",
                excel,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"default_category_id": category_id},
    )
    summary = response.json()["data"]
    assert summary["created"] == 0
    assert summary["failed"] == 1
    assert "安装地点为空" in summary["failures"][0]["reason"]


async def test_import_rejects_missing_header(client: AsyncClient) -> Any:
    category_id, _ = await _prepare_category_and_location(client)
    workbook = Workbook()
    workbook.active.append(["序号", "设备名称", "备注"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    response = await client.post(
        "/api/v1/equipment/equipments/import",
        files={
            "file": (
                "ledger.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"default_category_id": category_id},
    )
    assert response.status_code == 400
    assert "设备编号" in response.json()["message"]


async def test_import_template_downloads_xlsx(client: AsyncClient) -> Any:
    response = await client.get("/api/v1/equipment/equipments/import-template")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None
    assert list(sheet.iter_rows(values_only=True))[0][1] == "设备编号"
