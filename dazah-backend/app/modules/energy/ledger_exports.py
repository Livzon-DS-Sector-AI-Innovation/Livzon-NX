"""Export monthly dataset cells without conflating blank, zero or quality state."""

from collections.abc import Sequence
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils.cell import get_column_letter

from app.modules.energy.ledger_schemas import WorkbookReport

STATUS_LABELS = {
    "complete": "完整",
    "adjusted": "已调整",
    "unconfigured": "待确认配置",
    "not_applicable": "不适用",
    "missing_reading": "缺少读数",
    "missing_dependency": "计算缺项",
    "reading_rollback": "读数回退",
    "configuration_conflict": "配置冲突",
    "abnormal_usage": "用量超出阈值",
    "abnormal_dependency": "来源存在异常",
    "zero_denominator": "分母为零",
    "out_of_range": "超出计量范围",
    "invalid_meter_event": "换表记录无效",
}


def build_workbook_export(reports: list[WorkbookReport]) -> bytes:
    workbook = Workbook(write_only=True)
    for report in reports:
        sheet = workbook.create_sheet(report.dataset)
        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = 14
        sheet.row_dimensions[1].height = 48
        for index, column in enumerate(report.columns):
            sheet.column_dimensions[get_column_letter(2 + index * 2)].width = 30
            sheet.column_dimensions[get_column_letter(3 + index * 2)].width = 24
        sheet.auto_filter.ref = (
            f"A1:{get_column_letter(1 + len(report.columns) * 2)}{len(report.rows) + 1}"
        )

        def append(
            values: Sequence[str | float | datetime | None], *, header: bool = False
        ) -> None:
            cells = []
            for index, value in enumerate(values):
                cell = WriteOnlyCell(sheet, value=value)
                cell.alignment = Alignment(
                    horizontal="center"
                    if header or (index > 0 and index % 2 == 0)
                    else "right"
                    if index > 0
                    else "left",
                    vertical="center",
                    wrap_text=header,
                )
                if header:
                    cell.font = Font(bold=True)
                if isinstance(value, str):
                    cell.data_type = "s"
                elif isinstance(value, datetime):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(value, float):
                    cell.number_format = (
                        "#,##0" if value.is_integer() else "#,##0.########"
                    )
                cells.append(cell)
            sheet.append(cells)

        append(
            ["日期"]
            + [
                label
                for column in report.columns
                for label in [
                    f"{column.label}（{column.unit}）",
                    f"{column.label}·状态",
                ]
            ],
            header=True,
        )
        for row in report.rows:
            values: list[str | float | datetime | None] = [
                datetime.combine(row.date, datetime.min.time())
            ]
            for column in report.columns:
                cell = row.cells[column.key]
                number: Decimal | None = cell.value
                values.extend(
                    [
                        None
                        if number is None
                        else float(number)
                        if len(number.normalize().as_tuple().digits) <= 15
                        else str(number),
                        STATUS_LABELS.get(cell.status, cell.status),
                    ]
                )
            append(values)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
