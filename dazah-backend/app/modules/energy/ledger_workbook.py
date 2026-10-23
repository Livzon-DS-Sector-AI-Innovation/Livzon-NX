"""Monthly views of the original energy datasets, backed by native facts."""

from calendar import monthrange
from datetime import date, datetime, timedelta

from openpyxl.utils.cell import get_column_letter
from pydantic import JsonValue
from sqlalchemy import select

from app.core.exceptions import AppException
from app.modules.energy.ledger_models import EnergyImportBatch, EnergyReading
from app.modules.energy.ledger_schemas import (
    MeterDefinition,
    RuleDefinition,
    UsageReview,
    UsageRow,
    WorkbookCell,
    WorkbookColumn,
    WorkbookDataset,
    WorkbookDay,
    WorkbookReport,
)
from app.modules.energy.ledger_service import EnergyLedgerService
from app.modules.energy.ledger_source_policy import manual_input_allowed
from app.modules.energy.ledger_source_workbook import SourceWorkbook
from app.modules.energy.ledger_template import stable_code

DATASETS = [
    WorkbookDataset(name=name, kind=kind)
    for name, kind in [
        ("饮用水码", "reading"),
        ("冰水码", "reading"),
        ("蒸汽码", "reading"),
        ("循环水码", "reading"),
        ("空气码", "reading"),
        ("电量", "mixed"),
        ("饮用水量", "usage"),
        ("蒸汽量", "usage"),
        ("冰水量", "usage"),
        ("空气量", "usage"),
        ("循环水量", "usage"),
        ("日总量", "usage"),
        ("101车间", "usage"),
        ("102-1车间", "usage"),
        ("103-1车间", "usage"),
        ("103-2车间", "usage"),
        ("201-1车间", "usage"),
        ("201-2车间", "usage"),
        ("203-1车间", "usage"),
        ("203-2车间", "usage"),
        ("203-3车间", "usage"),
    ]
]


class EnergyWorkbookService:
    def __init__(self, service: EnergyLedgerService) -> None:
        self.service = service
        self.facts: dict[tuple[str, date], UsageRow] = {}
        self.readings: dict[tuple[str, date], EnergyReading] = {}
        self.definitions: dict[str, MeterDefinition | RuleDefinition] = {}
        self.headers: dict[str, list[str]] = {}
        self.loaded: tuple[int, int] | None = None
        self.reviews: list[UsageReview] = []
        self.source_workbook: SourceWorkbook | None = None
        self.source_batch_id: str | None = None
        self.source_created_at: datetime | None = None

    async def load(self, year: int, month: int) -> None:
        if self.loaded == (year, month):
            return
        self.facts.clear()
        self.readings.clear()
        self.definitions.clear()
        self.headers.clear()
        self.source_workbook = None
        self.source_batch_id = None
        self.source_created_at = None
        self.reviews = await self.service.usage_reviews()
        start, end = date(year, month, 1), date(year, month, monthrange(year, month)[1])
        meters = await self.service.meters()
        rules = await self.service.rules()
        for meter in meters:
            self.definitions["point:" + meter.definition.code] = meter.definition
        for rule in rules:
            self.definitions["rule:" + rule.definition.code] = rule.definition
        records = (
            await self.service.db.scalars(
                select(EnergyReading).where(
                    EnergyReading.point_id.in_([meter.id for meter in meters]),
                    EnergyReading.reading_date.between(
                        start - timedelta(days=32), end + timedelta(days=32)
                    ),
                    EnergyReading.is_deleted.is_(False),
                )
            )
        ).all()
        codes = {meter.id: "point:" + meter.definition.code for meter in meters}
        for reading in records:
            self.readings[(codes[reading.point_id], reading.reading_date)] = reading
        # Full-grid imports preserve source formulas and literal overrides separately
        # from native meter inputs. Legacy imports retain the native-fact behavior.
        batches = (
            await self.service.db.scalars(
                select(EnergyImportBatch)
                .where(
                    EnergyImportBatch.is_deleted.is_(False),
                    EnergyImportBatch.status == "imported",
                )
                .order_by(EnergyImportBatch.created_at.desc())
                .limit(20)
            )
        ).all()
        for batch in reversed(batches):
            try:
                source = SourceWorkbook.from_payload(batch.payload)
            except ValueError as exc:
                raise AppException(409, "原表快照无效，请重新导入完整原表") from exc
            if source.covers(year, month):
                self.source_workbook = source
                self.source_batch_id = str(batch.id)
                self.source_created_at = batch.created_at
            sheets = batch.payload.get("sheets")
            if not isinstance(sheets, list):
                continue
            for sheet in sheets:
                if isinstance(sheet, dict):
                    headers = sheet.get("headers")
                    if isinstance(headers, list):
                        self.headers[str(sheet["name"])] = [str(x) for x in headers]
        if self.source_workbook is not None:
            for (key, day), reading in self.readings.items():
                definition = self.definitions[key]
                if self.source_created_at is not None and (
                    reading.submitted_at is None
                    or reading.submitted_at < self.source_created_at
                ):
                    # An older, rounded import must not replace the newer source's
                    # full-precision reading. Later submissions still recalculate.
                    continue
                if definition.source_column and reading.value is not None:
                    try:
                        self.source_workbook.set_reading(
                            definition.source_sheet,
                            definition.source_column,
                            day,
                            reading.value,
                            version=reading.version,
                        )
                    except ValueError as exc:
                        raise AppException(
                            409, "原始读数来源与原表公式冲突，请核对导入版本"
                        ) from exc
        if self.source_workbook is None or any(
            item.name not in self.source_workbook.sheets for item in DATASETS
        ):
            report = await self.service.report(
                start, end, report_scope="all", include_all=True
            )
            self.facts = {(row.metric_key, row.usage_date): row for row in report.rows}
        self.loaded = (year, month)

    def source_report(self, year: int, month: int, dataset: str) -> WorkbookReport:
        """Keep original formulas, source dates and literals in the 21-table view."""
        book = self.source_workbook
        if book is None or dataset not in book.sheets:
            raise AppException(409, "原表快照缺少工作表，请重新导入完整原表")
        source = book.sheets[dataset]
        columns: list[WorkbookColumn] = []
        for column in range(2, max((len(row) for row in source.grid), default=0) + 1):
            label = source.headers[column - 1] if column <= len(source.headers) else ""
            if "大罐数量" in label:
                continue
            if not label:
                if not any(
                    len(row) >= column and row[column - 1] is not None
                    for row in source.grid
                ):
                    continue
                label = f"原表{get_column_letter(column)}列"
            definition = next(
                (
                    item
                    for item in self.definitions.values()
                    if item.source_sheet == dataset and item.source_column == column
                ),
                None,
            )
            columns.append(
                WorkbookColumn(
                    key=str(column),
                    label=label,
                    unit=definition.unit if definition else "待确认",
                    value_kind="reading" if dataset.endswith("码") else "usage",
                    quantity_kind=definition.quantity_kind if definition else "",
                )
            )
        rows: list[WorkbookDay] = []
        available, expected = 0, 0
        for day in range(1, monthrange(year, month)[1] + 1):
            current = date(year, month, day)
            cells: dict[str, WorkbookCell] = {}
            for field in columns:
                cell = book.cell_on(dataset, int(field.key), current)
                cell.provenance["source_batch_id"] = self.source_batch_id
                evidence: list[JsonValue] = [
                    item
                    for review in self.reviews
                    if review.dataset == dataset and review.column == int(field.key)
                    for item in review.evidence_on(current)
                ]
                if evidence:
                    # Review notes remain evidence; the source value is retained
                    # for inspection but no longer counts as a complete value.
                    cell = cell.model_copy(
                        update={
                            "status": "pending_review"
                            if cell.status == "complete"
                            else cell.status,
                            "provenance": {**cell.provenance, "review": evidence},
                        }
                    )
                cells[field.key] = cell
                expected += cell.status != "not_applicable"
                available += cell.value is not None and cell.status == "complete"
            rows.append(WorkbookDay(date=current, cells=cells))
        return WorkbookReport(
            year=year,
            month=month,
            dataset=dataset,
            columns=columns,
            rows=rows,
            expected_count=expected,
            available_count=available,
        )

    def matches(
        self,
        key: str,
        definition: MeterDefinition | RuleDefinition,
        sheet: str,
        column: int,
    ) -> bool:
        if definition.source_sheet:
            return (
                definition.source_sheet == sheet and definition.source_column == column
            )
        # Compatibility with older extracted configurations.
        code = key.split(":", 1)[1]
        original = stable_code(sheet, column)
        return code == original or code.endswith("_" + original)

    async def report(self, year: int, month: int, dataset: str) -> WorkbookReport:
        if dataset not in {item.name for item in DATASETS}:
            raise AppException(404, "能源数据表不存在")
        await self.load(year, month)
        if self.source_workbook is not None and dataset in self.source_workbook.sheets:
            return self.source_report(year, month, dataset)
        headers = self.headers.get(dataset, [])
        kind = next(item.kind for item in DATASETS if item.name == dataset)
        last_column = max(
            [len(headers), 200 if not headers else 0]
            + [
                definition.source_column or 0
                for definition in self.definitions.values()
                if definition.source_sheet == dataset
            ]
        )
        columns: list[WorkbookColumn] = []
        candidates: dict[str, list[str]] = {}
        for column in range(2, last_column + 1):
            keys = [
                key
                for key, definition in self.definitions.items()
                if self.matches(key, definition, dataset, column)
            ]
            label = headers[column - 1].strip() if column <= len(headers) else ""
            if not label and keys:
                label = self.definitions[keys[0]].name
            if not label or "大罐数量" in label:
                continue
            if sum(header.strip() == label for header in headers) > 1:
                label = f"{label}〔{get_column_letter(column)}列〕"
            definition = next(
                (
                    self.definitions[key]
                    for key in keys
                    if self.definitions[key].enabled
                ),
                self.definitions[keys[0]] if keys else None,
            )
            col_key = str(column)
            candidates[col_key] = keys
            columns.append(
                WorkbookColumn(
                    key=col_key,
                    label=label,
                    unit=definition.unit if definition else "待确认",
                    value_kind="reading"
                    if kind == "reading"
                    or (keys and all(key.startswith("point:") for key in keys))
                    else "usage",
                    quantity_kind=definition.quantity_kind if definition else "",
                )
            )
        rows: list[WorkbookDay] = []
        expected, available = 0, 0
        for day in range(1, monthrange(year, month)[1] + 1):
            current = date(year, month, day)
            cells: dict[str, WorkbookCell] = {}
            for field in columns:
                choices: list[WorkbookCell] = []
                for key in candidates[field.key]:
                    definition = self.definitions[key]
                    if key.startswith("point:"):
                        if isinstance(
                            definition, MeterDefinition
                        ) and not manual_input_allowed(
                            definition.source_sheet,
                            definition.name,
                            definition.reading_type,
                            definition.category,
                        ):
                            if definition.enabled:
                                choices.append(
                                    WorkbookCell(status="configuration_conflict")
                                )
                            continue
                        reading = self.readings.get((key, current))
                        if reading and reading.value is not None:
                            choices.append(
                                WorkbookCell(
                                    value=reading.value,
                                    status="complete",
                                    reading_id=reading.id,
                                    provenance={
                                        "source": reading.source,
                                        "version": reading.version,
                                        "reading_date": str(current),
                                    },
                                )
                            )
                        elif (
                            definition.enabled
                            and definition.active_from <= current
                            and (
                                definition.active_to is None
                                or current <= definition.active_to
                            )
                        ):
                            if (
                                definition.frequency != "monthly"
                                or day == monthrange(year, month)[1]
                            ):
                                choices.append(WorkbookCell(status="missing_reading"))
                    elif fact := self.facts.get((key, current)):
                        choices.append(
                            WorkbookCell(
                                value=fact.effective_value,
                                calculated_value=fact.calculated_value,
                                status=fact.status,
                                trace_id=None
                                if fact.provenance.get("unmaterialized")
                                else fact.id,
                                provenance=fact.provenance,
                            )
                        )
                populated = [cell for cell in choices if cell.value is not None]
                if len(choices) > 1:
                    cell = WorkbookCell(status="configuration_conflict")
                elif populated:
                    cell = populated[0]
                elif choices:
                    cell = choices[0]
                else:
                    cell = WorkbookCell(
                        status="unconfigured"
                        if not candidates[field.key]
                        or not any(
                            self.definitions[key].enabled
                            for key in candidates[field.key]
                        )
                        else "not_applicable"
                    )
                if field.value_kind == "usage":
                    evidence = [
                        item
                        for review in self.reviews
                        if review.dataset == dataset and review.column == int(field.key)
                        for item in review.evidence_on(current)
                    ]
                    if evidence:
                        cell = cell.model_copy(
                            update={
                                "value": None,
                                "status": "pending_review",
                                "provenance": {**cell.provenance, "review": evidence},
                            }
                        )
                cells[field.key] = cell
                expected += cell.status != "not_applicable"
                available += cell.value is not None and cell.status in {
                    "complete",
                    "adjusted",
                }
            rows.append(WorkbookDay(date=current, cells=cells))
        return WorkbookReport(
            year=year,
            month=month,
            dataset=dataset,
            columns=columns,
            rows=rows,
            expected_count=expected,
            available_count=available,
        )
