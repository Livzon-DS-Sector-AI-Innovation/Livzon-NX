"""Extract editable, disabled configuration from a reviewed spreadsheet preview."""

import hashlib
import re
from datetime import date
from decimal import Decimal

from openpyxl.utils.cell import column_index_from_string, get_column_letter
from pydantic import BaseModel, JsonValue

from app.modules.energy.ledger_calculations import validate_graph
from app.modules.energy.ledger_imports import (
    SOURCE_CATEGORIES,
    calculated_columns,
    input_column_allowed,
)
from app.modules.energy.ledger_schemas import (
    EnergyCategory,
    MeterDefinition,
    RuleDefinition,
    RuleTerm,
)

REFERENCE = re.compile(
    r"(?:(?P<sheet>'(?:[^']|'')+'|[\w.]+)!)?"
    r"\$?(?P<column>[A-Z]{1,3})\$?(?P<row>\d+)"
)
CATEGORIES: dict[str, EnergyCategory] = {
    "饮用水码": "drinking_water",
    "冰水码": "chilled_water",
    "蒸汽码": "steam",
    "循环水码": "circulating_water",
    "空气码": "air",
    "电量": "electricity",
}


class TemplateDraftResponse(BaseModel):
    meters: list[MeterDefinition]
    rules: list[RuleDefinition]
    issues: list[str]


def stable_code(sheet: str, column: int) -> str:
    return "xlsx_" + hashlib.sha256(sheet.encode()).hexdigest()[:10] + "_" + str(column)


def reference_sheet(match: re.Match[str], default: str) -> str:
    source = match["sheet"] or default
    return source[1:-1].replace("''", "'") if source.startswith("'") else source


def formula_expression(formula: str) -> str:
    # Keep spaces and escaped apostrophes inside quoted worksheet names.
    return re.sub(
        r"'(?:[^']|'')+'|\s+",
        lambda match: "" if match[0].isspace() else match[0],
        formula.lstrip("="),
    )


def temperature(label: str) -> str:
    if re.search(r"-5\s*[°度℃]", label):
        return "-5℃"
    return "7℃" if re.search(r"7\s*[°度℃]", label) else ""


def electricity_fallback(
    headers: list[JsonValue], row: int, computed: set[int]
) -> dict[int, str]:
    """Rebuild only the four confirmed aggregate relationships, never cached values."""
    labels = [str(label).strip() for label in headers]

    def unique(*names: str) -> int | None:
        matches = [i + 1 for i, label in enumerate(labels) if label in names]
        return matches[0] if len(matches) == 1 else None

    def cell(column: int) -> str:
        return f"{get_column_letter(column)}{row}"

    total = unique("总计", "合计")
    subtotal = unique("小计")
    difference = unique("其他", "差额", "差值")
    summed = unique("分表总和")
    west, east = unique("西动力总", "西动力总表"), unique("东动力总", "东动力总表")
    result: dict[int, str] = {}
    if total and west and east:
        result[total] = f"={cell(west)}+{cell(east)}"
    if subtotal:
        devices = [
            column
            for column in range(2, subtotal)
            if input_column_allowed("电量", headers, column, computed)
            and column not in {west, east}
        ]
        if devices:
            result[subtotal] = "=" + "+".join(cell(column) for column in devices)
    if difference and total and subtotal:
        result[difference] = f"={cell(total)}-{cell(subtotal)}"
    if summed and subtotal and difference:
        result[summed] = f"={cell(subtotal)}+{cell(difference)}"
    return result


def counter_factor(expression: str) -> Decimal | None:
    """Accept a bounded counter difference scaled only by numeric factors."""
    if expression == "R-R":
        return Decimal(1)
    if not re.fullmatch(r"\(R-R\)(?:[*/]\d+(?:\.\d+)?){1,12}", expression):
        return None
    factor = Decimal(1)
    for operator, number in re.findall(r"([*/])(\d+(?:\.\d+)?)", expression):
        value = Decimal(number)
        if len(number) > 20 or (operator == "/" and value == 0):
            raise ValueError("折算常量超限或除数为零")
        factor = factor * value if operator == "*" else factor / value
        if abs(factor) >= Decimal("1e20"):
            raise ValueError("折算系数超限")
    return factor.quantize(Decimal(".00000001"))


def extract_template(payload: dict[str, JsonValue]) -> TemplateDraftResponse:
    raw_sheets = payload.get("sheets")
    raw_formulas = payload.get("formulas")
    if not isinstance(raw_sheets, list) or not isinstance(raw_formulas, list):
        raise ValueError("缺少工作簿预览")
    sheets = {str(item["name"]): item for item in raw_sheets if isinstance(item, dict)}
    formulas: dict[tuple[str, int], list[dict[str, JsonValue]]] = {}
    for item in raw_formulas:
        if isinstance(item, dict):
            column_match = re.match(r"([A-Z]+)", str(item["cell"]))
            if column_match:
                formulas.setdefault(
                    (str(item["sheet"]), column_index_from_string(column_match[1])), []
                ).append(item)
    points: dict[tuple[str, int], MeterDefinition] = {}
    rules: dict[tuple[str, int], RuleDefinition] = {}
    issues: list[str] = []
    for name, sheet in sheets.items():
        cells, headers = sheet.get("cells"), sheet.get("headers")
        if name in SOURCE_CATEGORIES or not isinstance(cells, list):
            continue
        manual: dict[int, list[dict[str, JsonValue]]] = {}
        for cell in cells:
            column = cell.get("column") if isinstance(cell, dict) else None
            if (
                isinstance(cell, dict)
                and isinstance(column, int)
                and isinstance(headers, list)
                and column <= len(headers)
                and "大罐数量" not in str(headers[column - 1])
            ):
                manual.setdefault(column, []).append(cell)
        for column, values in manual.items():
            first = values[0]
            issues.append(
                f"{name}!{get_column_letter(column)} 有 {len(values)} 个手填用量，"
                f"例如 {first['date']} {first['cell']}={first['value']}；"
                "须核对为独立调整，不作为原始表底或自动零值导入"
            )
    row_dates: dict[tuple[str, int], date] = {}
    for name, sheet in sheets.items():
        dated_rows = sheet.get("row_dates")
        if isinstance(dated_rows, list):
            for item in dated_rows:
                row_number = item.get("row") if isinstance(item, dict) else None
                if isinstance(item, dict) and isinstance(row_number, int):
                    row_dates[name, row_number] = date.fromisoformat(str(item["date"]))
    for name, sheet in sheets.items():
        headers = sheet.get("headers")
        if name != "电量" or not isinstance(headers, list):
            continue
        dated = sorted(
            (row, day) for (source, row), day in row_dates.items() if source == name
        )
        cells = sheet.get("cells")
        if not dated and isinstance(cells, list):
            days = sorted(
                {str(cell["date"]) for cell in cells if isinstance(cell, dict)}
            )
            dated = [(row, date.fromisoformat(day)) for row, day in enumerate(days, 2)]
            row_dates.update({(name, row): day for row, day in dated})
        if not dated:
            continue
        row, day = dated[0]
        for column, formula in electricity_fallback(
            headers, row, calculated_columns(name, headers)
        ).items():
            if (name, column) not in formulas:
                formulas[name, column] = [
                    {
                        "sheet": name,
                        "cell": f"{get_column_letter(column)}{row}",
                        "formula": formula,
                        "date": day.isoformat(),
                    }
                ]
                issues.append(
                    f"电量!{get_column_letter(column)} 未保留公式；"
                    "按已确认的自动汇总关系生成停用草案，范围和缺项仍须核对"
                )
    for name, sheet in sheets.items():
        if name not in SOURCE_CATEGORIES:
            continue
        cells, headers = sheet.get("cells"), sheet.get("headers")
        if not isinstance(cells, list) or not isinstance(headers, list):
            continue
        columns: set[int] = set(range(2, len(headers) + 1))
        populated: set[int] = set()
        for cell in cells:
            if isinstance(cell, dict):
                source_column = cell.get("column")
                if isinstance(source_column, int) and not isinstance(
                    source_column, bool
                ):
                    columns.add(source_column)
                    populated.add(source_column)
        dates = [
            date.fromisoformat(str(item["date"]))
            for item in cells
            if isinstance(item, dict)
        ]
        dates.extend(day for (source, _), day in row_dates.items() if source == name)
        if not dates:
            issues.append(f"{name} 没有可识别的业务日期，未生成计量点")
            continue
        for column in sorted(columns):
            raw_computed = sheet.get("calculated_columns", [])
            computed = (
                {value for value in raw_computed if isinstance(value, int)}
                if isinstance(raw_computed, list)
                else set()
            )
            if (name, column) in formulas or not input_column_allowed(
                name, headers, column, computed
            ):
                if (
                    column <= len(headers)
                    and str(headers[column - 1]).strip()
                    and (name, column) not in formulas
                    and column not in calculated_columns(name, headers)
                ):
                    issues.append(
                        f"{name}!{get_column_letter(column)} 来源字段待确认，"
                        "不生成原始计量点或自动补零"
                    )
                continue
            label = str(headers[column - 1]).strip()
            if not label or "大罐数量" in label:
                continue
            if column not in populated:
                issues.append(
                    f"{name}!{get_column_letter(column)} {label} 无原始读数；"
                    "保留未启用计量点，不按零处理"
                )
            quantity = (
                "electricity"
                if name == "电量"
                else "mass"
                if name == "蒸汽码"
                else "heat"
                if "热值" in label
                else "volume"
            )
            duplicate = headers.count(headers[column - 1]) > 1
            points[(name, column)] = MeterDefinition(
                code=stable_code(name, column),
                name=f"{label}〔{get_column_letter(column)}列〕"
                if duplicate
                else label,
                category=CATEGORIES[name],
                area=label,
                department="待确认",
                unit="待确认",
                workshop=label if "车间" in label else "",
                reading_type="direct" if name == "电量" else "counter",
                frequency="monthly" if "月底抄表" in label else "daily",
                active_from=min(dates),
                enabled=False,
                quantity_kind=quantity,
                source_sheet=name,
                source_column=column,
                temperature=temperature(label) if name == "冰水码" else "",
                notes="从原表提取的未启用草案；请确认部门、单位、倍率和统计范围。"
                + ("名称含倍率提示，请核对原公式是否已换算。" if "倍" in label else ""),
            )
    visiting: set[tuple[str, int]] = set()

    def resolve(sheet_name: str, column: int) -> str:
        key = (sheet_name, column)
        if key in points:
            return "point:" + points[key].code
        if key in rules:
            return "rule:" + rules[key].code
        if key in visiting or key not in formulas:
            raise ValueError(f"{sheet_name}!{get_column_letter(column)} 来源缺失或循环")
        visiting.add(key)
        headers = sheets[sheet_name].get("headers", [])
        output_label = (
            str(headers[column - 1])
            if isinstance(headers, list) and column <= len(headers)
            else ""
        )
        expected_temperature = temperature(output_label)
        candidates: list[dict[str, JsonValue]] = []
        for candidate in formulas[key]:
            expression = formula_expression(str(candidate["formula"]))
            candidate_refs = list(REFERENCE.finditer(expression))
            if (
                REFERENCE.sub("R", expression) == "R"
                and candidate_refs
                and reference_sheet(candidate_refs[0], sheet_name) in SOURCE_CATEGORIES
                and reference_sheet(candidate_refs[0], sheet_name) != "电量"
            ):
                issues.append(
                    f"{sheet_name}!{candidate['cell']} 错将累计表底作为用量，"
                    "排除错误样本；按其他日期已确认的差分关系生成停用草案"
                )
                continue
            wrong_temperature = False
            for match in REFERENCE.finditer(expression):
                source = reference_sheet(match, sheet_name)
                source_headers = sheets.get(source, {}).get("headers")
                source_column = column_index_from_string(match["column"])
                if isinstance(source_headers, list) and source_column <= len(
                    source_headers
                ):
                    source_temperature = temperature(
                        str(source_headers[source_column - 1])
                    )
                    wrong_temperature |= bool(
                        expected_temperature
                        and source_temperature
                        and expected_temperature != source_temperature
                    )
            if wrong_temperature:
                issues.append(
                    f"{sheet_name}!{candidate['cell']} 温度引用不一致："
                    f"{candidate['formula']}；排除该错误样本，缺少正确来源时保留待确认"
                )
            else:
                candidates.append(candidate)
        if not candidates:
            raise ValueError("温度不一致，没有已确认的正确计算来源，保留待确认")
        compositions = {
            REFERENCE.sub(
                lambda match: f"{reference_sheet(match, sheet_name)}!{match['column']}",
                formula_expression(str(candidate["formula"])),
            )
            for candidate in candidates
        }
        if len(compositions) > 1:
            evidence = "；".join(
                f"{item.get('date')} {item['cell']} {item['formula']}"
                for item in (candidates[0], candidates[-1])
            )
            raise ValueError(
                f"计算组成变化，设备归属或分摊待确认，未生成规则：{evidence}"
            )
        sample = candidates[0]
        expression = formula_expression(str(sample["formula"]))
        refs = list(REFERENCE.finditer(expression))
        terms: list[RuleTerm] = []
        operation = "sum"
        if expression.startswith("SUM("):
            sum_match = re.fullmatch(r"SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)", expression)
            if not sum_match or sum_match[2] != sum_match[4]:
                raise ValueError("仅提取同日横向 SUM")
            for source_column in range(
                column_index_from_string(sum_match[1]),
                column_index_from_string(sum_match[3]) + 1,
            ):
                if (sheet_name, source_column) in points or (
                    sheet_name,
                    source_column,
                ) in formulas:
                    terms.append(
                        RuleTerm(metric_key=resolve(sheet_name, source_column))
                    )
                else:
                    raise ValueError("合计来源字段缺失或未确认，不能按零跳过")
        else:
            transformed = REFERENCE.sub("R", expression)
            source_keys = [
                (
                    reference_sheet(ref, sheet_name),
                    column_index_from_string(ref["column"]),
                )
                for ref in refs
            ]
            factor = counter_factor(transformed)
            if (
                len(refs) == 2
                and factor is not None
                and source_keys[0] == source_keys[1]
                and source_keys[0] in points
                and points[source_keys[0]].reading_type == "counter"
            ):
                if int(refs[0]["row"]) != int(refs[1]["row"]) + 1:
                    raise ValueError("累计读数不是相邻次日减当日，须核对日期")
                target_day = (
                    date.fromisoformat(str(sample["date"]))
                    if sample.get("date")
                    else None
                )
                source_days = [
                    row_dates.get((source_keys[0][0], int(ref["row"]))) for ref in refs
                ]
                before_day, after_day = source_days
                if (
                    target_day
                    and before_day
                    and after_day
                    and (after_day != target_day or (before_day - target_day).days != 1)
                ):
                    raise ValueError("累计读数日期不是次日减当日，保留待确认")
                terms = [
                    RuleTerm(metric_key=resolve(*source_keys[0]), coefficient=factor)
                ]
            else:
                target_day = (
                    date.fromisoformat(str(sample["date"]))
                    if sample.get("date")
                    else None
                )
                if target_day and any(
                    (source_day := row_dates.get((source[0], int(ref["row"]))))
                    and source_day != target_day
                    for source, ref in zip(source_keys, refs, strict=True)
                ):
                    issues.append(
                        f"{sheet_name}!{sample['cell']} 跨日引用已按同日业务日期"
                        "生成停用草案，请核对后启用"
                    )
                if not re.fullmatch(r"R(?:[+\-/]R)*", transformed):
                    raise ValueError("含未支持的常量、函数或表达式，需人工配置")
                if any(
                    key in points and points[key].reading_type == "counter"
                    for key in source_keys
                ):
                    raise ValueError("直接引用累计表底不能当作用量，须配置相邻读数差值")
                if "/" in transformed:
                    if transformed != "R/R":
                        raise ValueError("仅提取单一分子除以分母")
                    operation = "ratio"
                signs = [Decimal(1)] + [
                    Decimal(-1) if sign == "-" else Decimal(1)
                    for sign in re.findall(r"[+\-/]", transformed)
                ]
                terms = [
                    RuleTerm(metric_key=resolve(*source), coefficient=sign)
                    for source, sign in zip(source_keys, signs, strict=True)
                ]
        if not terms:
            raise ValueError("没有可用来源")
        definitions = {"point:" + item.code: item for item in points.values()} | {
            "rule:" + item.code: item for item in rules.values()
        }
        first = definitions[terms[0].metric_key]
        headers = sheets[sheet_name].get("headers", [])
        label = (
            str(headers[column - 1])
            if isinstance(headers, list) and column <= len(headers)
            else ""
        )
        if not label:
            label = f"{first.name}（原表{get_column_letter(column)}列无表头）"
            issues.append(
                f"{sheet_name}!{get_column_letter(column)} 无表头，"
                f"按明确公式来源保留：{first.name}"
            )
        if "大罐数量" in label:
            raise ValueError("没有有效指标名称")
        report_scope = (
            "plant"
            if sheet_name == "日总量"
            else "detail"
            if "车间" in sheet_name
            else "source"
        )
        if any(term in label for term in ("差值", "差额", "排空", "占比")) or (
            sheet_name == "电量" and label.strip() == "其他"
        ):
            report_scope = "check"
        rules[key] = RuleDefinition(
            code=stable_code(sheet_name, column),
            name=f"{sheet_name}·{label}",
            category=first.category,
            unit="比例" if operation == "ratio" else first.unit,
            workshop=sheet_name if "车间" in sheet_name else "",
            department="待确认",
            operation=operation,
            frequency=first.frequency,
            terms=terms,
            quantity_kind="ratio" if operation == "ratio" else first.quantity_kind,
            temperature=first.temperature,
            report_scope=report_scope,
            active_from=date.fromisoformat(str(sample["date"]))
            if sample.get("date")
            else first.active_from,
            enabled=False,
            notes=(
                f"原表 {sheet_name}!{sample['cell']}：{sample['formula']}。"
                "须确认统计范围、单位及特殊日期引用后启用。"
            ),
            source_sheet=sheet_name,
            source_column=column,
        )
        visiting.remove(key)
        # Detect changed row offsets and changed source compositions.
        signatures: dict[str, dict[str, JsonValue]] = {}
        dated = all(
            formula.get("date")
            and all(
                (reference_sheet(match, sheet_name), int(match["row"])) in row_dates
                for match in REFERENCE.finditer(str(formula["formula"]))
            )
            for formula in formulas[key]
        )
        for formula in formulas[key]:
            row_match = re.search(r"\d+$", str(formula["cell"]))
            target_row = int(row_match[0]) if row_match else 0
            target_date = (
                date.fromisoformat(str(formula["date"]))
                if formula.get("date")
                else None
            )

            def reference_signature(match: re.Match[str]) -> str:
                source = reference_sheet(match, sheet_name)
                source_day = row_dates.get((source, int(match["row"])))
                offset = (
                    f"日期差{(source_day - target_date).days}"
                    if dated and source_day and target_date
                    else f"行差{int(match['row']) - target_row}"
                )
                return f"{source}!{match['column']}[{offset}]"

            signature = REFERENCE.sub(reference_signature, str(formula["formula"]))
            signatures.setdefault(signature, formula)
        if len(signatures) > 1:
            evidence = "；".join(
                f"{item.get('date') or '无日期'} {item['cell']} {item['formula']}"
                for item in signatures.values()
            )
            issues.append(
                f"{sheet_name}·{label} 存在跨行引用，"
                f"草案按同日业务日期计算，须核对：{evidence}"
            )
            rules[key].notes = (rules[key].notes + "差异：" + evidence)[:2000]
        return "rule:" + rules[key].code

    for key in formulas:
        headers = sheets[key[0]].get("headers")
        if (
            isinstance(headers, list)
            and key[1] <= len(headers)
            and "大罐数量" in str(headers[key[1] - 1])
        ):
            continue
        try:
            resolve(*key)
        except (ValueError, KeyError) as exc:
            issues.append(f"{key[0]}!{get_column_letter(key[1])}：{exc}")
            visiting.clear()
    # Parent and child totals must not both be additive in a reporting scope.
    by_code = {"rule:" + item.code: item for item in rules.values()}
    for rule in rules.values():
        if rule.operation == "sum":
            for term in rule.terms:
                child = by_code.get(term.metric_key)
                if (
                    child
                    and child.report_scope == rule.report_scope
                    and rule.report_scope in {"plant", "detail"}
                ):
                    child.report_scope = "source"
    point_defs = {"point:" + item.code: item for item in points.values()}
    order = validate_graph(point_defs, by_code)
    return TemplateDraftResponse(
        meters=list(points.values()),
        rules=[by_code[key] for key in order],
        issues=issues,
    )
