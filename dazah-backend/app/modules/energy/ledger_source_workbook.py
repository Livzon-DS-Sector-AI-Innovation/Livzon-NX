"""Bounded evaluation of an imported source grid, preserving exact cell references.

Source formulas are historical reporting evidence. They do not become editable
meter points, change permission checks, or authorize writes to native readings.
"""

import ast
import re
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation, localcontext

from openpyxl.utils.cell import column_index_from_string, get_column_letter
from pydantic import JsonValue

from app.modules.energy.ledger_schemas import WorkbookCell

REFERENCE = re.compile(
    r"(?:(?P<sheet>'(?:[^']|'')+'|[\w.]+)!)?"
    r"\$?(?P<column>[A-Z]{1,3})\$?(?P<row>\d+)"
)
ADDRESS = re.compile(r"([A-Z]{1,3})([1-9]\d{0,3})")
MAX_CELLS = 100_000


@dataclass(frozen=True)
class SourceSheet:
    headers: list[str]
    grid: list[list[JsonValue]]
    dates: dict[date, int]


class SourceFormulaError(ValueError):
    def __init__(self, status: str) -> None:
        self.status = status
        super().__init__(status)


class SourceWorkbook:
    def __init__(self, sheets: dict[str, SourceSheet]) -> None:
        self.sheets = sheets
        self._cache: dict[tuple[str, str], WorkbookCell] = {}
        self._readings: dict[tuple[str, str], tuple[Decimal, int]] = {}

    @classmethod
    def from_payload(cls, payload: dict[str, JsonValue]) -> "SourceWorkbook":
        source = payload.get("sheets")
        if not isinstance(source, list) or len(source) > 21:
            raise ValueError("原表快照最多包含21张工作表")
        sheets: dict[str, SourceSheet] = {}
        count = 0
        for item in source:
            if not isinstance(item, dict) or "source_grid" not in item:
                continue
            name, grid, headers = (
                item.get("name"),
                item.get("source_grid"),
                item.get("headers"),
            )
            if (
                not isinstance(name, str)
                or not name
                or name in sheets
                or not isinstance(grid, list)
                or len(grid) > 5000
                or not isinstance(headers, list)
                or len(headers) > 200
            ):
                raise ValueError("原表快照结构无效或包含重复表名")
            rows: list[list[JsonValue]] = []
            for row in grid:
                if not isinstance(row, list) or len(row) > 200:
                    raise ValueError("原表快照每行最多200列")
                if any(
                    isinstance(value, (dict, list))
                    or isinstance(value, str)
                    and len(value) > 4096
                    for value in row
                ):
                    raise ValueError("原表单元格不是有效的标量")
                count += len(row)
                if count > MAX_CELLS:
                    raise ValueError("原表快照超过单次处理上限")
                rows.append(row)
            dates: dict[date, int] = {}
            dated = item.get("row_dates", [])
            if not isinstance(dated, list):
                raise ValueError("原表缺少日期索引")
            for entry in dated:
                if not isinstance(entry, dict):
                    raise ValueError("原表日期索引无效")
                raw_day, row_index = entry.get("date"), entry.get("row")
                if (
                    not isinstance(raw_day, str)
                    or type(row_index) is not int
                    or not 1 <= row_index <= len(rows)
                ):
                    raise ValueError("原表日期索引超出范围")
                day = date.fromisoformat(raw_day)
                if day in dates:
                    raise ValueError("原表存在重复日期，不能任意选择日期行")
                dates[day] = row_index
            sheets[name] = SourceSheet([str(x) for x in headers], rows, dates)
        return cls(sheets)

    def covers(self, year: int, month: int) -> bool:
        return any(
            day.year == year and day.month == month
            for sheet in self.sheets.values()
            for day in sheet.dates
        )

    def set_reading(
        self, sheet: str, column: int, day: date, value: Decimal, *, version: int
    ) -> None:
        source = self.sheets.get(sheet)
        if source is None or day not in source.dates:
            return
        address = f"{get_column_letter(column)}{source.dates[day]}"
        original = self._raw(sheet, address)
        if isinstance(original, str) and original.startswith("="):
            raise ValueError("原表公式列不能被人工读数覆盖")
        if not value.is_finite() or abs(value) >= Decimal("1e20"):
            raise ValueError("原始读数超出计量范围")
        self._readings[sheet, address] = value, version
        self._cache.clear()

    def cell_on(self, sheet: str, column: int, day: date) -> WorkbookCell:
        source = self.sheets.get(sheet)
        if source is None or day not in source.dates:
            return WorkbookCell(status="not_applicable")
        return self.result(sheet, f"{get_column_letter(column)}{source.dates[day]}")

    def _raw(self, sheet: str, address: str) -> JsonValue:
        source = self.sheets.get(sheet)
        if source is None:
            raise SourceFormulaError("source_missing_sheet")
        match = ADDRESS.fullmatch(address)
        if match is None:
            raise SourceFormulaError("source_formula_error")
        row, column = int(match[2]), column_index_from_string(match[1])
        if row > 5000 or column > 200:
            raise SourceFormulaError("source_formula_error")
        # A known worksheet's empty cells have Excel blank-reference semantics.
        if row > len(source.grid) or column > len(source.grid[row - 1]):
            return None
        return source.grid[row - 1][column - 1]

    @staticmethod
    def _number(value: JsonValue) -> Decimal:
        if isinstance(value, bool) or value is None:
            raise SourceFormulaError("source_non_numeric")
        try:
            result = Decimal(str(value).strip())
        except InvalidOperation as exc:
            raise SourceFormulaError("source_non_numeric") from exc
        if not result.is_finite() or abs(result) >= Decimal("1e20"):
            raise SourceFormulaError("out_of_range")
        return result

    def result(
        self, sheet: str, address: str, stack: tuple[tuple[str, str], ...] = ()
    ) -> WorkbookCell:
        key = sheet, address
        if key in stack or len(stack) >= 64:
            return WorkbookCell(status="source_formula_cycle")
        if key in self._cache:
            return self._cache[key].model_copy(deep=True)
        provenance: dict[str, JsonValue] = {
            "source_kind": "literal",
            "source_sheet": sheet,
            "source_cell": address,
        }
        try:
            raw = self._raw(sheet, address)
            if key in self._readings:
                value, version = self._readings[key]
                provenance.update(source_kind="reading", version=version)
            elif isinstance(raw, str) and raw.startswith("="):
                provenance.update(source_kind="formula", formula=raw)
                value, sources = self._formula(sheet, raw, stack + (key,))
                provenance["sources"] = sources
            elif raw is None or raw == "":
                return WorkbookCell(status="source_blank", provenance=provenance)
            else:
                value = self._number(raw)
            if not value.is_finite() or abs(value) >= Decimal("1e20"):
                raise SourceFormulaError("out_of_range")
            with localcontext() as context:
                context.prec = 60
                value = value.quantize(Decimal(".00000001"))
            cell = WorkbookCell(
                value=value,
                calculated_value=value,
                status="complete",
                provenance=provenance,
            )
        except SourceFormulaError as exc:
            cell = WorkbookCell(status=exc.status, provenance=provenance)
        except (ArithmeticError, SyntaxError, RecursionError):
            cell = WorkbookCell(status="source_formula_error", provenance=provenance)
        self._cache[key] = cell
        return cell.model_copy(deep=True)

    def _formula(
        self, sheet: str, formula: str, stack: tuple[tuple[str, str], ...]
    ) -> tuple[Decimal, list[JsonValue]]:
        if len(formula) > 2048:
            raise SourceFormulaError("source_formula_error")
        references: dict[str, tuple[str, str]] = {}
        evidence: list[JsonValue] = []

        def replace(match: re.Match[str]) -> str:
            source = match["sheet"] or sheet
            if source.startswith("'"):
                source = source[1:-1].replace("''", "'")
            token = f"cell_{len(references)}"
            references[token] = source, match["column"] + match["row"]
            return token

        expression = REFERENCE.sub(replace, formula[1:])

        def expand(match: re.Match[str]) -> str:
            left, right = references[match[1]], references[match[2]]
            a, b = ADDRESS.fullmatch(left[1]), ADDRESS.fullmatch(right[1])
            if left[0] != right[0] or a is None or b is None:
                raise SourceFormulaError("source_formula_error")
            tokens: list[str] = []
            for row in range(int(a[2]), int(b[2]) + 1):
                for column in range(
                    column_index_from_string(a[1]), column_index_from_string(b[1]) + 1
                ):
                    if len(references) >= 1000:
                        raise SourceFormulaError("source_formula_error")
                    token = f"cell_{len(references)}"
                    references[token] = left[0], f"{get_column_letter(column)}{row}"
                    tokens.append(token)
            return ",".join(tokens)

        expression = re.sub(r"(cell_\d+):(cell_\d+)", expand, expression)
        tree = ast.parse(expression, mode="eval")
        if sum(1 for _ in ast.walk(tree)) > 2000:
            raise SourceFormulaError("source_formula_error")

        def run(node: ast.AST, *, in_sum: bool = False) -> Decimal:
            if isinstance(node, ast.Expression):
                return run(node.body)
            if isinstance(node, ast.Constant) and type(node.value) in (int, float):
                return self._number(str(node.value))
            if isinstance(node, ast.Name) and node.id in references:
                source, address = references[node.id]
                cell = self.result(source, address, stack)
                evidence.append(
                    {
                        "sheet": source,
                        "cell": address,
                        "value": str(cell.value) if cell.value is not None else None,
                        "status": cell.status,
                    }
                )
                if cell.status == "source_blank" or (
                    in_sum and cell.status == "source_non_numeric"
                ):
                    return Decimal(0)
                if cell.value is None:
                    raise SourceFormulaError(cell.status)
                return cell.value
            if isinstance(node, ast.UnaryOp) and isinstance(
                node.op, (ast.USub, ast.UAdd)
            ):
                value = run(node.operand)
                return -value if isinstance(node.op, ast.USub) else value
            if isinstance(node, ast.BinOp):
                left, right = run(node.left), run(node.right)
                if isinstance(node.op, ast.Add):
                    return left + right
                if isinstance(node.op, ast.Sub):
                    return left - right
                if isinstance(node.op, ast.Mult):
                    return left * right
                if isinstance(node.op, ast.Div):
                    if not right:
                        raise SourceFormulaError("zero_denominator")
                    return left / right
            if (
                isinstance(node, ast.Call)
                and isinstance(node.func, ast.Name)
                and node.func.id.upper() == "SUM"
                and not node.keywords
            ):
                return sum((run(arg, in_sum=True) for arg in node.args), Decimal(0))
            raise SourceFormulaError("source_formula_error")

        with localcontext() as context:
            context.prec = 60
            return run(tree), evidence
