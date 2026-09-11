"""OOT 限度告知单（docx）导入与导出。

导入：解析「XX年 XX 产品OOT限度通知单」格式的 Word 告知单，得到产品信息与
限度明细（一级项目/项目/标准/OOT限度）。已存在同产品名+年份的产品时整表替换
其明细，否则新建产品（编码自动生成，可在页面修改）。
导出：以内置模板（由真实告知单原件裁剪数据行而来）克隆数据行填充，保证版式
与原件一致；产品含一级项目时使用分组版式模板。
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from copy import deepcopy
from dataclasses import asdict, dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

from sqlalchemy import delete as sa_delete
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AppException
from app.modules.quality.models.oot_limit import OotLimitItem, OotLimitProduct

logger = logging.getLogger(__name__)

TEMPLATES_DIR = Path(__file__).resolve().parent.parent / "templates"
SIMPLE_TEMPLATE_PATH = TEMPLATES_DIR / "oot_limit_notice_template.docx"
GROUPED_TEMPLATE_PATH = TEMPLATES_DIR / "oot_limit_notice_grouped_template.docx"

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
ZIP_MEDIA_TYPE = "application/zip"

MAX_IMPORT_FILES = 20

# 标题形如「2026 年 洛伐他汀 产品OOT限度通知单」（空格多少不固定）
_TITLE_RE = re.compile(
    r"^(?P<pre>\s*)(?P<year>20\d{2})(?P<mid>\s*年\s*)(?P<name>.+?)"
    r"(?P<post>\s*产品\s*OOT\s*限度通知单\s*)$"
)
_YEAR_RE = re.compile(r"20\d{2}")


@dataclass
class NoticeItemData:
    """告知单中的一行限度明细。"""

    display_order: int
    item_group: str | None
    item_name: str
    standard_value: str
    oot_limit_value: str


@dataclass
class NoticeParseResult:
    """单个告知单文件的解析结果。"""

    source_file_name: str
    document_title: str
    document_year: int | None
    product_name: str
    items: list[NoticeItemData] = field(default_factory=list)
    row_errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def _clean_text(value: str | None) -> str:
    """合并连续空白并去除首尾空白（保留单元格内的换行合并为空格）。"""
    return re.sub(r"\s+", " ", value or "").strip()


def _logical_cells(row: Any) -> list[tuple[int, int, str]]:
    """按 grid 几何取一行的逻辑单元格：[(起始列, grid_span, 文本)]。

    Word 中横向合并的单元格会在 row.cells 里重复出现（同一底层 tc），
    按 tc 去重后才是真实列结构。
    """
    seen: set[int] = set()
    cells: list[tuple[int, int, str]] = []
    for col, cell in enumerate(row.cells):
        tc = cell._tc
        if id(tc) in seen:
            continue
        seen.add(id(tc))
        cells.append((col, int(tc.grid_span or 1), cell.text))
    return cells


def _find_notice_table(document: Any) -> Any | None:
    """定位限度表：表头含「OOT限度」的第一个表格（告知单第一个表是签批表）。"""
    for table in document.tables:
        if not table.rows:
            continue
        header_texts = [cell.text for cell in table.rows[0].cells]
        if any("OOT限度" in (text or "") for text in header_texts):
            return table
    return None


def parse_notice_docx(content: bytes, source_file_name: str) -> NoticeParseResult:
    """解析单个告知单 docx；文件无法打开时抛 AppException(400)。"""
    import docx

    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001
        raise AppException(
            message="无法解析 Word 文件，请确认上传的是有效 .docx 文件",
            status_code=400,
        ) from exc

    warnings: list[str] = []
    document_title = ""
    product_name = ""
    document_year: int | None = None
    for paragraph in document.paragraphs:
        raw = (paragraph.text or "").strip()
        match = _TITLE_RE.match(raw)
        if match:
            document_title = _clean_text(raw)
            document_year = int(match.group("year"))
            product_name = _clean_text(match.group("name"))
            break

    if not product_name:
        # 正文无标准标题段时按文件名兜底推断
        stem = re.sub(r"\.docx$", "", source_file_name or "", flags=re.IGNORECASE)
        year_match = _YEAR_RE.search(stem)
        if year_match:
            document_year = int(year_match.group())
        name_guess = re.sub(
            r"年|产品|OOT|限度|告知单|通知单|（|\(|）|\)|\s",
            "",
            _YEAR_RE.sub("", stem),
        )
        if name_guess:
            product_name = name_guess
            document_title = _clean_text(stem)
            warnings.append(
                "未在正文识别到标准标题段，已按文件名推断产品与年份，请核对"
            )

    result = NoticeParseResult(
        source_file_name=source_file_name,
        document_title=document_title or source_file_name,
        document_year=document_year,
        product_name=product_name,
        warnings=warnings,
    )
    if not product_name:
        result.row_errors.append("无法识别产品名称（正文标题段与文件名均未命中）")
        return result

    table = _find_notice_table(document)
    if table is None:
        result.row_errors.append("未找到限度表（表头缺少「OOT限度」列）")
        return result

    # 分组版式（多拉菌素）为 5 grid 列（项目占2列+标准占2列+OOT限度），
    # 普通版式为 3 grid 列；表头逻辑单元格两种版式都是 3 个，不能按其判断
    grouped_layout = len(table.rows[0].cells) >= 4
    for row_no, row in enumerate(table.rows[1:], start=1):
        texts = [_clean_text(text) for _, _, text in _logical_cells(row)]
        if not any(texts):
            continue
        item_group: str | None
        if grouped_layout:
            if len(texts) >= 4:
                item_group, item_name, standard, oot = (
                    texts[0],
                    texts[1],
                    texts[2],
                    texts[3],
                )
            elif len(texts) == 3:
                item_group, item_name, standard, oot = (
                    None,
                    texts[0],
                    texts[1],
                    texts[2],
                )
            else:
                result.row_errors.append(f"第{row_no}行：列数不足，无法解析")
                continue
        else:
            item_group = None
            if len(texts) >= 3:
                item_name, standard, oot = texts[0], texts[1], texts[2]
            elif len(texts) == 2:
                # 标准与 OOT 限度合并为一个单元格（两者相同）
                item_name, standard, oot = texts[0], texts[1], texts[1]
            else:
                result.row_errors.append(f"第{row_no}行：列数不足，无法解析")
                continue
        if not item_name or not standard or not oot:
            missing = "、".join(
                label
                for label, value in (
                    ("项目", item_name),
                    ("标准值", standard),
                    ("OOT限度", oot),
                )
                if not value
            )
            result.row_errors.append(f"第{row_no}行：缺少{missing}")
            continue
        result.items.append(
            NoticeItemData(
                display_order=len(result.items) + 1,
                item_group=item_group or None,
                item_name=item_name,
                standard_value=standard,
                oot_limit_value=oot,
            )
        )
    return result


async def _find_notice_product(
    db: AsyncSession, parsed: NoticeParseResult
) -> OotLimitProduct | None:
    """匹配已有产品：优先 产品名+年份，标题被改过时按源文件名兜底。"""
    stmt = select(OotLimitProduct).where(OotLimitProduct.is_deleted.is_(False))
    if parsed.document_year is not None:
        stmt = stmt.where(
            OotLimitProduct.product_name == parsed.product_name,
            OotLimitProduct.document_year == parsed.document_year,
        )
    else:
        stmt = stmt.where(OotLimitProduct.product_name == parsed.product_name)
    result = await db.execute(stmt.order_by(OotLimitProduct.created_at.asc()).limit(1))
    product = result.scalar_one_or_none()
    if product is not None:
        return product
    fallback = await db.execute(
        select(OotLimitProduct)
        .where(
            OotLimitProduct.is_deleted.is_(False),
            OotLimitProduct.source_file_name == parsed.source_file_name,
        )
        .order_by(OotLimitProduct.created_at.asc())
        .limit(1)
    )
    return fallback.scalar_one_or_none()


async def _generate_product_code(db: AsyncSession, year: int | None) -> str:
    prefix = f"OOT-{year if year is not None else 'NA'}-"
    result = await db.execute(
        select(OotLimitProduct.product_code).where(
            OotLimitProduct.product_code.like(f"{prefix}%")
        )
    )
    existing = {row[0] for row in result.all()}
    seq = 1
    while f"{prefix}{seq:02d}" in existing:
        seq += 1
    return f"{prefix}{seq:02d}"


async def preview_oot_limit_notice_import(
    db: AsyncSession, files: list[tuple[str, bytes]]
) -> dict[str, Any]:
    """逐文件解析并标记 新建/更新，不写库。"""
    file_entries: list[dict[str, Any]] = []
    for filename, content in files:
        try:
            parsed = parse_notice_docx(content, filename)
        except AppException as exc:
            file_entries.append(
                {"filename": filename, "status": "error", "error": exc.message}
            )
            continue
        product = await _find_notice_product(db, parsed)
        file_entries.append(
            {
                "filename": filename,
                "status": "ok",
                "mode": "update" if product is not None else "create",
                "matched_product_code": product.product_code if product else None,
                "product_name": parsed.product_name,
                "document_title": parsed.document_title,
                "document_year": parsed.document_year,
                "item_count": len(parsed.items),
                "items": [asdict(item) for item in parsed.items],
                "row_errors": parsed.row_errors,
                "warnings": parsed.warnings,
            }
        )
    return {"files": file_entries}


async def confirm_oot_limit_notice_import(
    db: AsyncSession, files: list[tuple[str, bytes]]
) -> dict[str, Any]:
    """逐文件导入：命中已有产品则更新并整表替换明细，否则新建；单文件独立提交。"""
    created_count = 0
    update_count = 0
    error_details: list[dict[str, str]] = []
    for filename, content in files:
        try:
            parsed = parse_notice_docx(content, filename)
            if not parsed.items:
                raise AppException(
                    message="；".join(parsed.row_errors) or "未解析到限度数据",
                    status_code=400,
                )
            product = await _find_notice_product(db, parsed)
            if product is None:
                product = OotLimitProduct(
                    product_code=await _generate_product_code(db, parsed.document_year),
                    product_name=parsed.product_name,
                    document_title=parsed.document_title,
                    document_year=parsed.document_year,
                    source_file_name=filename,
                    is_active=True,
                )
                db.add(product)
                created_count += 1
            else:
                product.product_name = parsed.product_name
                product.document_title = parsed.document_title
                if parsed.document_year is not None:
                    product.document_year = parsed.document_year
                product.source_file_name = filename
                # 明细是告知单的全量派生数据：物理删除后按文件重建
                # （display_order 唯一约束覆盖软删行，无法软删后重建同序号）
                await db.execute(
                    sa_delete(OotLimitItem).where(OotLimitItem.product_id == product.id)
                )
                update_count += 1
            await db.flush()
            for item_data in parsed.items:
                db.add(
                    OotLimitItem(
                        product_id=product.id,
                        display_order=item_data.display_order,
                        item_group=item_data.item_group,
                        item_name=item_data.item_name,
                        standard_value=item_data.standard_value,
                        oot_limit_value=item_data.oot_limit_value,
                        # 旧版列与标准列同步写入，保持旧读取兼容
                        specification=item_data.standard_value,
                        oot_limit=item_data.oot_limit_value,
                    )
                )
            await db.commit()
        except Exception as exc:  # noqa: BLE001
            await db.rollback()
            message = exc.message if isinstance(exc, AppException) else str(exc)
            logger.warning("OOT notice import failed for %s: %s", filename, message)
            error_details.append({"filename": filename, "error": message})
    return {
        "created_count": created_count,
        "update_count": update_count,
        "error_count": len(error_details),
        "error_details": error_details,
    }


def notice_export_filename(product: Any) -> str:
    candidate = (getattr(product, "source_file_name", None) or "").strip()
    if not candidate:
        title = getattr(product, "document_title", None)
        candidate = f"{title or product.product_name}.docx"
    if not candidate.lower().endswith(".docx"):
        candidate = f"{candidate}.docx"
    return candidate


def _set_paragraph_text(paragraph: Any, text: str) -> None:
    """改写段落文本并保留首 run 的字体格式。"""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _replace_notice_title(document: Any, product: Any) -> None:
    """按模板标题的原始空格结构替换年份与产品名。"""
    product_name = getattr(product, "product_name", "") or ""
    for paragraph in document.paragraphs:
        raw = (paragraph.text or "").strip()
        match = _TITLE_RE.match(raw)
        if not match:
            continue
        year = getattr(product, "document_year", None) or match.group("year")
        new_text = (
            f"{match.group('pre')}{year}{match.group('mid')}"
            f"{product_name}{match.group('post')}"
        )
        _set_paragraph_text(paragraph, new_text)
        return


def _set_cell_text(cell_elem: Any, text: str) -> None:
    """在 XML 层写单元格文本，保留单元格与 run 的原有格式。"""
    from docx.oxml.ns import qn

    paragraph_elems = cell_elem.findall(qn("w:p"))
    if not paragraph_elems:
        paragraph_elem = cell_elem.makeelement(qn("w:p"), {})
        cell_elem.append(paragraph_elem)
        paragraph_elems = [paragraph_elem]
    first = paragraph_elems[0]
    for extra in paragraph_elems[1:]:
        extra.getparent().remove(extra)

    run_elems = first.findall(qn("w:r"))
    if run_elems:
        keep = run_elems[0]
        for extra in run_elems[1:]:
            first.remove(extra)
    else:
        keep = first.makeelement(qn("w:r"), {})
        first.append(keep)
    for text_elem in keep.findall(qn("w:t")):
        keep.remove(text_elem)
    text_elem = keep.makeelement(qn("w:t"), {})
    text_elem.text = text
    text_elem.set(qn("xml:space"), "preserve")
    keep.append(text_elem)


def export_notice_docx(product: Any, items: list[Any]) -> bytes:
    """导出单个产品告知单；items 须按 display_order 升序。

    产品含一级项目时用分组版式模板（多拉菌素样式），否则用普通版式
    （美伐他汀样式）。数据行通过 deepcopy 模板样例行填充，保留原件的
    字体/边框/列宽/签批表。
    """
    import docx
    from docx.oxml.ns import qn

    has_group = any(getattr(item, "item_group", None) for item in items)
    template_path = GROUPED_TEMPLATE_PATH if has_group else SIMPLE_TEMPLATE_PATH
    if not template_path.exists():
        raise AppException(message="导出模板缺失，请联系管理员", status_code=500)
    document = docx.Document(str(template_path))

    _replace_notice_title(document, product)

    table = _find_notice_table(document)
    if table is None:
        raise AppException(message="导出模板缺少限度表，请联系管理员", status_code=500)

    table_element = table._tbl
    row_elements = table_element.findall(qn("w:tr"))
    grouped_layout = len(table.rows[0].cells) >= 4
    if grouped_layout:
        # 分组模板保留两行样例：普通行（3 逻辑格）+ 分组行（4 逻辑格）
        plain_sample = deepcopy(row_elements[1])
        group_sample = deepcopy(row_elements[2])
        for row_elem in row_elements[1:]:
            table_element.remove(row_elem)
    else:
        plain_sample = deepcopy(row_elements[1])
        group_sample = None
        for row_elem in row_elements[1:]:
            table_element.remove(row_elem)

    for item in items:
        item_group = getattr(item, "item_group", None)
        if grouped_layout and item_group:
            row_elem = deepcopy(group_sample)
            values = [
                item_group,
                item.item_name,
                item.standard_value,
                item.oot_limit_value,
            ]
        else:
            row_elem = deepcopy(plain_sample)
            values = [item.item_name, item.standard_value, item.oot_limit_value]
        table_element.append(row_elem)
        for tc_elem, text in zip(row_elem.findall(qn("w:tc")), values):
            _set_cell_text(tc_elem, str(text or ""))

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output.getvalue()


def build_notice_zip(entries: list[tuple[str, bytes]]) -> bytes:
    """把多个告知单 docx 打成 zip，文件名重复时自动追加序号。"""
    buffer = BytesIO()
    used: set[str] = set()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for filename, data in entries:
            name = filename
            counter = 1
            while name in used:
                stem, dot, ext = filename.rpartition(".")
                name = f"{stem}({counter}).{ext}" if dot else f"{filename}({counter})"
                counter += 1
            used.add(name)
            archive.writestr(name, data)
    return buffer.getvalue()


__all__ = [
    "DOCX_MEDIA_TYPE",
    "MAX_IMPORT_FILES",
    "NoticeItemData",
    "NoticeParseResult",
    "ZIP_MEDIA_TYPE",
    "build_notice_zip",
    "confirm_oot_limit_notice_import",
    "export_notice_docx",
    "notice_export_filename",
    "parse_notice_docx",
    "preview_oot_limit_notice_import",
]
