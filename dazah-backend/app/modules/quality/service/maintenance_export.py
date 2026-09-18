"""设备维护保养记录按月汇总导出（xlsx，复刻用户桌面模板格式）。

模板：app/modules/quality/templates/设备年度预防性维护保养汇总表-模板.xlsx
（由用户桌面 2026年7月 表格经 LibreOffice 转换入库，保留字体/列宽/行高/
合并/边框）。导出时删除模板示例数据行，按所选月份的维保镜像数据重填并
复制示例行样式；末尾保留合并的备注说明行。

列布局复刻模板现状：第 2 列表头为「设备编号」但实际填设备名称、
第 3 列表头为「设备名称」但实际填设备编号（用户手工表即如此，保持一致）。
维护日期输出文本 YYYY.MM.DD（如 2026.07.12），与手工表填写方式一致。
"""

from __future__ import annotations

import logging
from copy import copy
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AppException
from app.modules.quality.service.inspection_instrument_mirror import (
    list_instrument_mirror,
)
from app.modules.quality.service.instruments_dashboard import (
    _cell_text,
    parse_feishu_date,
)

logger = logging.getLogger(__name__)

_TEMPLATE_PATH = (
    Path(__file__).resolve().parent.parent
    / "templates"
    / "设备年度预防性维护保养汇总表-模板.xlsx"
)
# 模板结构：第1行标题(合并A1:J1)、第2行表头、第3~5行示例数据、第6行备注(合并A6:J6)
_DATA_FIRST_ROW = 3
_TEMPLATE_DATA_ROWS = 3
_TEMPLATE_FOOTER_ROW = 6
_DATA_COLUMNS = 10
_PAGE_SIZE = 500


def _format_dot_date(value: Any) -> str:
    """维护日期 -> 文本 YYYY.MM.DD（与用户手工表填写方式一致）。"""
    parsed = parse_feishu_date(value)
    if parsed is None:
        return _cell_text(value)
    return f"{parsed.year}.{parsed.month:02d}.{parsed.day:02d}"


def _month_bounds(month: str) -> tuple[int, int]:
    try:
        year_str, month_str = month.split("-")
        year, month_num = int(year_str), int(month_str)
        if not (1900 < year < 3000) or not (1 <= month_num <= 12):
            raise ValueError
    except (ValueError, AttributeError) as exc:
        raise AppException(
            message="月份格式不正确，应为 YYYY-MM（如 2026-07）", status_code=400
        ) from exc
    return year, month_num


def _remark_value(item: dict[str, Any]) -> str:
    """备注列：Base 里列名在「备注」「维备注」间有过改动，两者都兼容。"""
    for key in ("备注", "维备注"):
        text = _cell_text(item.get(key))
        if text:
            return text
    return ""


def _row_values(index: int, item: dict[str, Any]) -> list[Any]:
    """镜像行 -> 模板一行（列序复刻用户手工表：名称在编号前）。"""
    remark = _remark_value(item)
    return [
        index,
        _cell_text(item.get("设备名称")),
        _cell_text(item.get("设备编号")),
        _cell_text(item.get("安装位置")),
        _cell_text(item.get("维护内容")),
        _format_dot_date(item.get("维护日期")),
        _cell_text(item.get("维护人")),
        _cell_text(item.get("复核人")),
        _cell_text(item.get("维保类型")),
        remark if remark else "-",
    ]


async def export_maintenance_summary(
    db: AsyncSession, month: str
) -> tuple[bytes, str]:
    """导出所选月份的维保汇总 xlsx，返回 (文件字节, 文件名)。"""
    year, month_num = _month_bounds(month)
    if not _TEMPLATE_PATH.exists():
        raise AppException(
            message="导出模板缺失，请联系管理员检查质量模块模板目录",
            status_code=500,
        )

    page = await list_instrument_mirror(
        db, "qc_instr_maintenance", page=1, page_size=_PAGE_SIZE
    )
    rows = [
        item
        for item in page.get("items") or []
        if (parsed := parse_feishu_date(item.get("维护日期"))) is not None
        and parsed.year == year
        and parsed.month == month_num
    ]
    rows.sort(
        key=lambda item: (
            parse_feishu_date(item.get("维护日期")) or 0,
            str(item.get("record_id") or ""),
        )
    )

    workbook = load_workbook(_TEMPLATE_PATH)
    sheet = workbook.active
    # 摘示例行样式与行高（逐列），删除示例数据后按数据量重插并回填
    style_refs = [
        copy(sheet.cell(_DATA_FIRST_ROW, col)._style)
        for col in range(1, _DATA_COLUMNS + 1)
    ]
    number_formats = [
        sheet.cell(_DATA_FIRST_ROW, col).number_format
        for col in range(1, _DATA_COLUMNS + 1)
    ]
    data_row_height = sheet.row_dimensions[_DATA_FIRST_ROW].height
    footer_row_height = sheet.row_dimensions[_TEMPLATE_FOOTER_ROW].height
    footer_text = sheet.cell(_TEMPLATE_FOOTER_ROW, 1).value

    # openpyxl 增删行不搬合并区域与行高：先解除末尾备注合并，行数调整后
    # 再把备注文字/行高放到新位置并重新合并
    sheet.unmerge_cells(
        start_row=_TEMPLATE_FOOTER_ROW, start_column=1,
        end_row=_TEMPLATE_FOOTER_ROW, end_column=_DATA_COLUMNS,
    )
    sheet.delete_rows(_DATA_FIRST_ROW, _TEMPLATE_DATA_ROWS)
    if rows:
        sheet.insert_rows(_DATA_FIRST_ROW, len(rows))
    for offset, item in enumerate(rows):
        excel_row = _DATA_FIRST_ROW + offset
        for col, value in enumerate(_row_values(offset + 1, item), start=1):
            cell = sheet.cell(excel_row, col)
            cell.value = value
            cell._style = copy(style_refs[col - 1])  # noqa: SLF001
            cell.number_format = number_formats[col - 1]
        if data_row_height is not None:
            sheet.row_dimensions[excel_row].height = data_row_height

    footer_row = _DATA_FIRST_ROW + len(rows)
    footer_cell = sheet.cell(footer_row, 1)
    if footer_cell.value in (None, ""):
        footer_cell.value = footer_text
    if footer_row_height is not None:
        sheet.row_dimensions[footer_row].height = footer_row_height
    sheet.merge_cells(
        start_row=footer_row, start_column=1,
        end_row=footer_row, end_column=_DATA_COLUMNS,
    )
    # 清掉增删行后残留的旧行高定义（模板行号失效产生的空行样式）
    for stale in [
        index
        for index in sheet.row_dimensions
        if index > footer_row
    ]:
        del sheet.row_dimensions[stale]

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    filename = f"{year}年{month_num}月  设备年度预防性维护保养汇总表.xlsx"
    logger.info(
        "维保汇总导出 month=%s rows=%d filename=%s", month, len(rows), filename
    )
    return buffer.read(), filename
