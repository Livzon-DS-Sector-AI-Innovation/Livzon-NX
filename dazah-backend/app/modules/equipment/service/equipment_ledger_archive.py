"""设备台账 Excel 上传存档服务。

台账导入成功后保存一份只读存档：解析后的全量表格快照
（rows/merges/col_widths，与前端渲染协议一致）、原件相对路径
和导入结果摘要；供台账页回看历次上传、下载原件。删除存档
只删快照，不影响已导入的台账数据。
"""

from __future__ import annotations

import uuid
from datetime import date, datetime, time
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AppException
from app.modules.equipment.models.equipment_ledger_archive import (
    EquipmentLedgerArchive,
)


def _serialize_cell(value: Any) -> Any:
    """openpyxl 单元格值 → JSON 安全且可读的展示值。

    日期/时间转可读字符串（前端直接展示，避免出现 Excel 序列号）；
    其余标量原样，未知类型转字符串。
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        if (
            value.hour == 0
            and value.minute == 0
            and value.second == 0
            and value.microsecond == 0
        ):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def parse_ledger_snapshot(data: bytes) -> dict[str, Any]:
    """解析台账 .xlsx 首个可见工作表为回看快照（全量二维数组）。"""
    try:
        workbook = load_workbook(BytesIO(data), data_only=True)
    except Exception as error:  # openpyxl 对损坏文件抛多种异常
        raise AppException(
            status_code=400, message=f"无法读取 Excel 文件：{error}"
        ) from error
    try:
        visible = [ws for ws in workbook.worksheets if ws.sheet_state == "visible"]
        worksheet = visible[0] if visible else workbook.worksheets[0]
        rows: list[list[Any]] = [
            [_serialize_cell(cell.value) for cell in row]
            for row in worksheet.iter_rows()
        ]
        merges = [
            {
                "s": {"r": rng.min_row - 1, "c": rng.min_col - 1},
                "e": {"r": rng.max_row - 1, "c": rng.max_col - 1},
            }
            for rng in worksheet.merged_cells.ranges
        ]
        col_widths: list[int] = []
        for index in range(1, worksheet.max_column + 1):
            dimension = worksheet.column_dimensions.get(_column_letter(index))
            width = dimension.width if dimension and dimension.width else 80
            col_widths.append(int(width))
        return {
            "sheet_name": worksheet.title,
            "rows": rows,
            "merges": merges,
            "col_widths": col_widths,
            "row_count": len(rows),
            "col_count": worksheet.max_column,
        }
    finally:
        workbook.close()


def _column_letter(index: int) -> str:
    """1-based 列号 → Excel 列字母（A、B、…、AA）。"""
    letters = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


async def create_archive(
    session: AsyncSession,
    *,
    workshop_code: str,
    file_name: str,
    sheet_name: str,
    original_path: str,
    rows: list[Any],
    merges: list[Any],
    col_widths: list[int],
    row_count: int,
    col_count: int,
    import_summary: dict[str, Any] | None = None,
    created_by: uuid.UUID | None = None,
) -> EquipmentLedgerArchive:
    archive = EquipmentLedgerArchive(
        workshop_code=workshop_code,
        file_name=file_name,
        sheet_name=sheet_name,
        original_path=original_path,
        rows=rows,
        merges=merges,
        col_widths=col_widths,
        row_count=row_count,
        col_count=col_count,
        import_summary=import_summary,
        created_by=created_by,
    )
    session.add(archive)
    await session.flush()
    return archive


def serialize_archive(
    archive: EquipmentLedgerArchive,
    *,
    created_by_name: str | None = None,
) -> dict[str, Any]:
    """ORM → API 响应（含完整表格快照）。"""
    return {
        "id": str(archive.id),
        "workshop_code": archive.workshop_code,
        "file_name": archive.file_name,
        "sheet_name": archive.sheet_name,
        "rows": archive.rows,
        "merges": archive.merges,
        "col_widths": archive.col_widths,
        "row_count": archive.row_count,
        "col_count": archive.col_count,
        "import_summary": archive.import_summary,
        "created_by_name": created_by_name,
        "created_at": (archive.created_at.isoformat() if archive.created_at else None),
        "updated_at": (archive.updated_at.isoformat() if archive.updated_at else None),
    }


def serialize_archive_summary(
    archive: EquipmentLedgerArchive,
    *,
    created_by_name: str | None = None,
) -> dict[str, Any]:
    """列表项（不含 rows/merges/col_widths 大字段）。"""
    payload = serialize_archive(archive, created_by_name=created_by_name)
    payload.pop("rows", None)
    payload.pop("merges", None)
    payload.pop("col_widths", None)
    return payload


async def list_archives(
    session: AsyncSession,
    *,
    page: int,
    page_size: int,
    workshop_code: str | None = None,
) -> tuple[list[tuple[EquipmentLedgerArchive, str | None]], int]:
    """分页列表（新在前），附带上传人姓名。"""
    from app.platform.identity.models import User

    conditions = [
        EquipmentLedgerArchive.is_deleted == False,  # noqa: E712
    ]
    if workshop_code:
        conditions.append(EquipmentLedgerArchive.workshop_code == workshop_code)
    total = await session.scalar(
        select(func.count(EquipmentLedgerArchive.id)).where(*conditions)
    )
    query = (
        select(EquipmentLedgerArchive, User.name)
        .outerjoin(User, User.id == EquipmentLedgerArchive.created_by)
        .where(*conditions)
        .order_by(EquipmentLedgerArchive.created_at.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    result = await session.execute(query)
    items = [(archive, name) for archive, name in result.all()]
    return items, int(total or 0)


async def get_user_name(session: AsyncSession, user_id: uuid.UUID | None) -> str | None:
    if user_id is None:
        return None
    from app.platform.identity.models import User

    name = await session.scalar(select(User.name).where(User.id == user_id))
    return name if isinstance(name, str) else None


async def get_archive(
    session: AsyncSession, archive_id: uuid.UUID
) -> EquipmentLedgerArchive | None:
    result = await session.execute(
        select(EquipmentLedgerArchive).where(
            EquipmentLedgerArchive.id == archive_id,
            EquipmentLedgerArchive.is_deleted.is_(False),
        )
    )
    return result.scalar_one_or_none()


async def delete_archive(
    session: AsyncSession,
    archive: EquipmentLedgerArchive,
    *,
    deleted_by: uuid.UUID | None = None,
) -> None:
    """软删存档：只删快照记录，不影响已导入的台账数据与原件文件。"""
    archive.is_deleted = True
    archive.updated_by = deleted_by
    await session.flush()
