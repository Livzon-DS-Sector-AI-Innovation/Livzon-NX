"""危险源辨识台账：车间 Excel 批量导入解析服务。

支持四类车间表格（一个工作簿内通常按岗位/区域拆分为多个子表 sheet）：

- ``activity``  作业活动安全风险分级管控措施清单（LEC 评价，含个体防护/应急/管控主体列）
- ``facility``  设备设施安全风险分级管控措施清单（LSR 评价，序号/名称大量纵向合并）
- ``key_risk``  作业活动重点风险管控措施清单（无 LEC 数值，含 AI 检索列）
- ``legacy``    危险源识别评价表（.xls 双行表头，时态/状态为勾选列，尾部附图例竖条）

解析为纯结构化函数（不依赖数据库），便于单测；导入落库由
``HazardLedgerImportService`` 承接。所有解析基于单元格位置锚点与表头关键字，
不调用 LLM。
"""

import hashlib
import io
import logging
import os
import re
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any

import openpyxl  # type: ignore[import-untyped]
import xlrd
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.safety.repository import SafetyRepository
from app.modules.safety.schemas import RISK_LEVELS
from app.platform.audit.service import record_audit_log

logger = logging.getLogger(__name__)

# ── 格式常量 ──

FORMAT_ACTIVITY = "activity"
FORMAT_FACILITY = "facility"
FORMAT_KEY_RISK = "key_risk"
FORMAT_LEGACY = "legacy"

FORMAT_LABELS: dict[str, str] = {
    FORMAT_ACTIVITY: "作业活动分级管控清单",
    FORMAT_FACILITY: "设备设施分级管控清单",
    FORMAT_KEY_RISK: "重点风险管控措施清单",
    FORMAT_LEGACY: "危险源辨识评价表",
}

# 风险等级原文 → 平台等级 key（level_1~level_4）
LEVEL_ALIASES: dict[str, str] = {
    "低风险": "level_4",
    "四级": "level_4",
    "较低风险": "level_4",
    "一般风险": "level_3",
    "三级": "level_3",
    "较大风险": "level_2",
    "二级": "level_2",
    "重大风险": "level_1",
    "一级": "level_1",
    "特大风险": "level_1",
}

LEVEL_LABEL_BY_KEY: dict[str, str] = {
    item["key"]: item["label"] for item in RISK_LEVELS
}
LEVEL_CONTROL_BY_KEY: dict[str, str] = {
    item["key"]: item["control_level"] for item in RISK_LEVELS
}
LEVEL_PERSON_BY_KEY: dict[str, str] = {
    item["key"]: item["responsible_person"] for item in RISK_LEVELS
}

MAX_UPLOAD_SIZE = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {".xlsx", ".xls"}

# 右侧信息列按表头标签映射（不同车间表格尾部列集合不同，禁止盲取固定偏移）
_RIGHT_LABELS: dict[str, str] = {
    "工程措施": "engineering",
    "管理措施": "management",
    "个体防护": "ppe",
    "应急措施": "emergency",
    "管控层级": "control_level",
    "责任单位": "responsible_unit",
    "责任人员": "responsible_person",
    "检查频次": "check_frequency",
    "检查标准": "check_standard",
}

# 数据库列长上限（超长截断防御，避免整批导入失败）
_FIELD_LIMITS: dict[str, int] = {
    "control_level": 50,
    "responsible_unit": 100,
    "responsible_person": 100,
    "check_frequency": 100,
    "tense": 50,
    "state": 50,
    "is_major": 10,
}


def _fit(field: str, value: str | None) -> str | None:
    """按数据库列长截断字段值。"""
    if value is None:
        return None
    limit = _FIELD_LIMITS.get(field)
    if limit and len(value) > limit:
        return value[: limit - 1] + "…"
    return value

# legacy 勾选列
_CHECK_MARKS = {"√", "✓", "是", "Y"}


@dataclass
class SheetMatrix:
    """归一化后的工作表矩阵（0-based，含合并单元格展开）。"""

    name: str
    rows: list[list[Any]] = field(default_factory=list)
    merged: list[tuple[int, int, int, int]] = field(default_factory=list)


@dataclass
class ParsedRow:
    """单行危险源解析结果。"""

    row_index: int
    category: str | None = None
    name: str | None = None
    step: str | None = None
    location: str | None = None
    target: str | None = None  # 设备设施「辨识对象」
    check_standard: str | None = None
    hazard_desc: str | None = None
    accident: str | None = None
    tense: str | None = None
    state: str | None = None
    l_value: float | None = None
    e_value: float | None = None
    c_value: float | None = None
    d_value: float | None = None
    level_raw: str | None = None
    level_key: str | None = None
    engineering: str | None = None
    management: str | None = None
    ppe: str | None = None
    emergency: str | None = None
    control_level: str | None = None
    responsible_unit: str | None = None
    responsible_person: str | None = None
    check_frequency: str | None = None
    is_major: str | None = None
    # 重点风险清单 AI 检索列
    effective: str | None = None
    tech_update: str | None = None
    case: str | None = None
    extra_recommendation: str | None = None


# ── 工作簿加载 ──


def load_workbook_matrices(data: bytes, filename: str) -> list[SheetMatrix]:
    """把 xlsx/xls 工作簿加载为矩阵列表；扩展名不匹配时抛 ValueError。"""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"仅支持 {', '.join(sorted(ALLOWED_EXTENSIONS))} 格式文件")
    if ext == ".xls":
        return _load_xls(data)
    return _load_xlsx(data)


def _load_xlsx(data: bytes) -> list[SheetMatrix]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    matrices: list[SheetMatrix] = []
    for ws in workbook.worksheets:
        rows = [[cell.value for cell in row] for row in ws.iter_rows()]
        merged = [
            (rng.min_row - 1, rng.max_row - 1, rng.min_col - 1, rng.max_col - 1)
            for rng in ws.merged_cells.ranges
        ]
        matrices.append(_build_matrix(ws.title, rows, merged))
    workbook.close()
    return matrices


def _load_xls(data: bytes) -> list[SheetMatrix]:
    workbook = xlrd.open_workbook(file_contents=data, formatting_info=True)
    matrices: list[SheetMatrix] = []
    for sheet in workbook.sheets():
        rows = [
            [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            for r in range(sheet.nrows)
        ]
        # xlrd 的 merged_cells 上界为开区间，统一转为闭区间
        merged = [
            (rlo, rhi - 1, clo, chi - 1) for rlo, rhi, clo, chi in sheet.merged_cells
        ]
        matrices.append(_build_matrix(sheet.name, rows, merged))
    return matrices


def _build_matrix(
    name: str, rows: list[list[Any]], merged: list[tuple[int, int, int, int]]
) -> SheetMatrix:
    """裁齐行宽并把合并单元格的值展开到整个区域。"""
    width = max((len(r) for r in rows), default=0)
    grid = [list(r) + [None] * (width - len(r)) for r in rows]
    for r1, r2, c1, c2 in merged:
        if not grid or r1 >= len(grid):
            continue
        value = grid[r1][c1] if c1 < width else None
        for r in range(r1, min(r2 + 1, len(grid))):
            for c in range(c1, min(c2 + 1, width)):
                grid[r][c] = value
    return SheetMatrix(name=name, rows=grid, merged=merged)


# ── 通用工具 ──


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = re.sub(r"[^\d.\-]", "", str(value))
    if not cleaned or cleaned in {"-", ".", "-."}:
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def _normalize_level(raw: Any) -> str | None:
    if not raw:
        return None
    key = _text(raw).replace(" ", "")
    if key in LEVEL_ALIASES:
        return LEVEL_ALIASES[key]
    # 兼容「一级（红色）」这类带后缀写法
    for alias, level in LEVEL_ALIASES.items():
        if key.startswith(alias):
            return level
    return None


def guess_department(filename: str) -> str:
    """从文件名推断车间/部门默认值，供前端预填。"""
    stem = os.path.splitext(os.path.basename(filename or ""))[0]
    range_match = re.search(r"(\d{2,4})\s*[-－~～至]\s*(\d{2,4})", stem)
    if range_match:
        return f"{range_match.group(1)}-{range_match.group(2)}车间"
    single_match = re.search(r"(\d{2,4}\s*[一二三四五六七八九十]?)车间", stem)
    if single_match:
        return re.sub(r"\s+", "", single_match.group(1)) + "车间"
    return ""


# ── 格式识别 ──


def detect_format(matrix: SheetMatrix) -> str | None:
    """依据标题行与表头关键字识别子表格式。"""
    head_text = "".join(_text(cell) for row in matrix.rows[:6] for cell in row[:8])
    if not head_text.strip():
        return None
    if "识别评价表" in head_text or "辨识评价表" in head_text:
        return FORMAT_LEGACY
    if "设备设施" in head_text:
        return FORMAT_FACILITY
    if "重点风险" in head_text:
        return FORMAT_KEY_RISK
    if "作业活动" in head_text:
        return FORMAT_ACTIVITY
    # 标题缺失时按表头特征兜底
    if "辨识对象" in head_text:
        return FORMAT_FACILITY
    if "暴露频次" in head_text or "作业步骤" in head_text:
        return FORMAT_ACTIVITY
    return None


def _find_header_row(matrix: SheetMatrix) -> int:
    """定位明细表头行（含「序号」且含任一业务列名）。"""
    for idx, row in enumerate(matrix.rows[:8]):
        cells = {_text(c) for c in row}
        if "序号" in cells and cells & {"作业步骤", "辨识对象", "工程措施", "危险源"}:
            return idx
    return -1


def _is_footer_or_note(row: list[Any]) -> bool:
    """识别表尾（编制/审核）与模板说明行。"""
    first_cell = _text(row[0]) if row else ""
    if first_cell.startswith(("编制", "审核", "批准", "单位名称")):
        return True
    joined = "".join(_text(c) for c in row)
    return "提问式检查表" in joined or joined.startswith("设备分类")


# ── 锚点式行解析（activity / facility）──


@dataclass
class _WalkConfig:
    """从「风险等级」锚点向左回溯的槽位配置。"""

    numeric_count: int
    left_slots: list[str]  # 数值区之后依次消费的槽位名（自右向左）


_WALK_CONFIGS: dict[str, _WalkConfig] = {
    FORMAT_ACTIVITY: _WalkConfig(
        numeric_count=4, left_slots=["state", "tense", "accident", "hazard"]
    ),
    FORMAT_FACILITY: _WalkConfig(
        numeric_count=3, left_slots=["accident", "standard", "hazard"]
    ),
}


def _build_header_map(
    matrix: SheetMatrix, header_idx: int
) -> tuple[dict[str, int], int | None]:
    """从明细表头行提取右侧信息列位置与「风险等级」列位置。

    个别列（如跨两行的「检查频次」）标签可能只在组表头行上，明细行为空，
    此时回退到组表头行取标签。
    """
    header_map: dict[str, int] = {}
    level_col: int | None = None
    group_row = matrix.rows[header_idx - 1] if header_idx >= 1 else []
    for col, cell in enumerate(matrix.rows[header_idx]):
        label = re.sub(r"\s+", "", _text(cell))
        if not label and col < len(group_row):
            label = re.sub(r"\s+", "", _text(group_row[col]))
        if label == "风险等级" and level_col is None:
            level_col = col
            continue
        field_name = _RIGHT_LABELS.get(label)
        if field_name and field_name not in header_map:
            header_map[field_name] = col
    return header_map, level_col


def _parse_row_by_anchor(
    row: list[Any],
    row_idx: int,
    fmt: str,
    header_map: dict[str, int],
    level_col: int | None,
) -> ParsedRow | None:
    """按「风险等级」锚点解析单行。

    车间实际表格存在整行左移（省略「名称/作业步骤」列）的错位子表，
    因此数值与左侧文本以等级标签为锚点回溯；右侧信息列按表头标签定位，
    错位行再按「表头等级列 − 行内锚点列」的偏移量整体平移。
    """
    config = _WALK_CONFIGS[fmt]
    parsed = ParsedRow(row_index=row_idx)

    anchor = -1
    for col, cell in enumerate(row):
        if _normalize_level(cell):
            anchor = col
            parsed.level_raw = _text(cell)
            parsed.level_key = _normalize_level(cell)
            break
    if anchor < 0:
        return None

    # 右侧信息列：表头标签定位 + 错位平移
    delta = (level_col - anchor) if level_col is not None else 0
    for field_name, col in header_map.items():
        source = col - delta
        if 0 <= source < len(row):
            value = _text(row[source])
            if value:
                setattr(parsed, field_name, _fit(field_name, value))

    # 左侧数值区（L/E/C/D 或 L/S/R），自锚点向左连续取数值
    numbers: list[float] = []
    col = anchor - 1
    while col >= 0 and len(numbers) < config.numeric_count:
        num = _to_float(row[col])
        if num is None:
            break
        numbers.append(num)
        col -= 1
    numbers.reverse()  # 恢复自左向右
    if fmt == FORMAT_ACTIVITY:
        if len(numbers) >= 4:
            (
                parsed.l_value,
                parsed.e_value,
                parsed.c_value,
                parsed.d_value,
            ) = numbers[-4:]
        elif len(numbers) == 3:
            parsed.l_value, parsed.c_value, parsed.d_value = numbers
        elif len(numbers) == 2:
            parsed.c_value, parsed.d_value = numbers
        elif numbers:
            parsed.d_value = numbers[0]
    else:
        if len(numbers) >= 3:
            parsed.l_value, parsed.c_value, parsed.d_value = numbers[-3:]
        elif len(numbers) == 2:
            parsed.c_value, parsed.d_value = numbers
        elif numbers:
            parsed.d_value = numbers[0]

    # 数值区之后按槽位消费（跳过空单元格）
    slots: dict[str, str] = {}
    for slot in config.left_slots:
        while col >= 0 and not _text(row[col]):
            col -= 1
        if col < 0:
            break
        slots[slot] = _text(row[col])
        col -= 1
    parsed.accident = slots.get("accident")
    parsed.hazard_desc = slots.get("hazard")
    parsed.tense = slots.get("tense")
    parsed.state = slots.get("state")
    # facility 走左侧槽位取检查标准；activity 的检查标准来自表头映射，勿覆盖
    if "standard" in slots:
        parsed.check_standard = slots["standard"]

    # 其余左侧文本（剔除纯数字序号），恢复自左向右
    leftovers: list[str] = []
    while col >= 0:
        value = _text(row[col])
        if value and _to_float(value) is None:
            leftovers.append(value)
        col -= 1
    leftovers.reverse()
    if fmt == FORMAT_ACTIVITY:
        if len(leftovers) >= 2:
            parsed.category = leftovers[0]
            parsed.name = leftovers[1]
            parsed.step = "；".join(leftovers[2:]) or None
        elif leftovers:
            # 错位行仅剩单个文本时，视作作业活动名称
            parsed.name = leftovers[0]
    else:
        if len(leftovers) == 1:
            parsed.name = leftovers[0]
        else:
            facility_keys = ("category", "name", "location", "target")
            for key, value in zip(facility_keys, leftovers):
                setattr(parsed, key, value or None)
    return parsed


# ── 重点风险管控措施清单（key_risk）──

_KEY_RISK_LABELS: dict[str, str] = {
    "类别": "category",
    "名称": "name",
    "作业步骤": "step",
    "危险源辨识": "hazard_desc",
    "可能发生的事故类型及后果": "accident",
    "工程措施": "engineering",
    "管理措施": "management",
    "个体防护": "ppe",
    "应急措施": "emergency",
    "有效措施": "effective",
    "技术设备更新": "tech_update",
    "案例": "case",
}


def _parse_key_risk_sheet(
    matrix: SheetMatrix,
) -> tuple[list[ParsedRow], list[dict[str, Any]]]:
    header_idx = _find_header_row(matrix)
    if header_idx < 0:
        return [], [{"row": 0, "reason": "未找到表头行"}]
    group_row = matrix.rows[header_idx - 1] if header_idx >= 1 else []
    col_map: dict[str, int] = {}
    for col, cell in enumerate(matrix.rows[header_idx]):
        label = _text(cell) or (_text(group_row[col]) if col < len(group_row) else "")
        field_name = _KEY_RISK_LABELS.get(label)
        if field_name and field_name not in col_map:
            col_map[field_name] = col

    rows: list[ParsedRow] = []
    errors: list[dict[str, Any]] = []
    for idx in range(header_idx + 1, len(matrix.rows)):
        row = matrix.rows[idx]
        if not any(_text(c) for c in row) or _is_footer_or_note(row):
            continue
        parsed = ParsedRow(row_index=idx)
        for attr, col in col_map.items():
            if col < len(row):
                value = _text(row[col]) or None
                setattr(parsed, attr, value)
        if not parsed.hazard_desc and not parsed.name:
            errors.append({"row": idx + 1, "reason": "危险源与名称均为空"})
            continue
        parts = []
        if "effective" in col_map:
            parts.append(f"有效措施：{parsed.effective or '-'}")
        if "tech_update" in col_map:
            parts.append(f"技术设备更新：{parsed.tech_update or '-'}")
        if "case" in col_map:
            parts.append(f"同业案例：{parsed.case or '-'}")
        parsed.extra_recommendation = "；".join(parts) or None
        rows.append(parsed)
    return rows, errors


# ── 危险源识别评价表（legacy, .xls 双行勾选表头）──

_LEGACY_MAIN_LABELS: list[tuple[str, str]] = [
    ("工序", "process_col"),
    ("作业活动", "activity_col"),
    ("工作活动名称", "activity_col"),
    ("危险源", "hazard_col"),
    ("可能导致", "accident_col"),
    ("危险等级", "level_col"),
    ("重要危险源", "major_col"),
    ("防范措施", "measure_col"),
    ("控制措施", "measure_col"),
]


def _parse_legacy_sheet(
    matrix: SheetMatrix,
) -> tuple[list[ParsedRow], list[dict[str, Any]]]:
    header_idx = -1
    for idx, row in enumerate(matrix.rows[:6]):
        if any(_text(c) == "序号" for c in row):
            header_idx = idx
            break
    if header_idx < 0 or header_idx + 1 >= len(matrix.rows):
        return [], [{"row": 0, "reason": "未找到表头行"}]
    main_header = matrix.rows[header_idx]
    sub_header = matrix.rows[header_idx + 1]

    tense_cols: dict[int, str] = {}
    state_cols: dict[int, str] = {}
    col_map: dict[str, int] = {}
    width = max(len(main_header), len(sub_header))
    for col in range(width):
        main_label = _text(main_header[col]) if col < len(main_header) else ""
        sub_label = _text(sub_header[col]) if col < len(sub_header) else ""
        if sub_label in {"过去", "现在", "将来"}:
            tense_cols[col] = sub_label
        elif sub_label in {"正常", "异常", "紧急"}:
            state_cols[col] = sub_label
        elif sub_label in {"L", "E", "C", "D"}:
            col_map[sub_label.lower() + "_col"] = col
        # 表头可能带换行（如「危\n险\n等\n级」），按去空白后的文本匹配
        compact_main = re.sub(r"\s+", "", main_label)
        for label, key in _LEGACY_MAIN_LABELS:
            if label in compact_main:
                col_map.setdefault(key, col)

    rows: list[ParsedRow] = []
    errors: list[dict[str, Any]] = []
    for idx in range(header_idx + 2, len(matrix.rows)):
        row = matrix.rows[idx]
        if not any(_text(c) for c in row) or _is_footer_or_note(row):
            continue
        parsed = ParsedRow(row_index=idx)

        def cell(key: str) -> str | None:
            col = col_map.get(key)
            if col is None or col >= len(row):
                return None
            return _text(row[col]) or None

        parsed.name = cell("process_col")
        parsed.step = cell("activity_col")
        parsed.hazard_desc = cell("hazard_col")
        parsed.accident = cell("accident_col")
        parsed.management = cell("measure_col")
        parsed.is_major = _fit("is_major", cell("major_col"))
        parsed.level_raw = cell("level_col")
        parsed.level_key = _normalize_level(parsed.level_raw)
        for key in ("l", "e", "c", "d"):
            num_col = col_map.get(key + "_col")
            if num_col is not None and num_col < len(row):
                setattr(parsed, f"{key}_value", _to_float(row[num_col]))
        tense_values = [
            label
            for col, label in tense_cols.items()
            if col < len(row) and _text(row[col]) in _CHECK_MARKS
        ]
        state_values = [
            label
            for col, label in state_cols.items()
            if col < len(row) and _text(row[col]) in _CHECK_MARKS
        ]
        parsed.tense = "、".join(tense_values) or None
        parsed.state = "、".join(state_values) or None
        if (
            parsed.d_value is None
            and parsed.l_value is not None
            and parsed.e_value is not None
            and parsed.c_value is not None
        ):
            parsed.d_value = parsed.l_value * parsed.e_value * parsed.c_value
        if not parsed.hazard_desc and not parsed.step:
            if any(
                getattr(parsed, f"{key}_value") is not None
                for key in ("l", "e", "c", "d")
            ):
                errors.append({"row": idx + 1, "reason": "危险源描述为空"})
            continue
        rows.append(parsed)
    return rows, errors


# ── 子表统一解析入口 ──


def parse_sheet(
    matrix: SheetMatrix,
) -> tuple[str | None, list[ParsedRow], list[dict[str, Any]]]:
    """解析单个子表，返回 (格式, 有效行, 行错误)。"""
    fmt = detect_format(matrix)
    rows: list[ParsedRow] = []
    errors: list[dict[str, Any]] = []
    if fmt is None:
        return None, [], [{"row": 0, "reason": "无法识别的表格格式"}]
    if fmt == FORMAT_KEY_RISK:
        rows, errors = _parse_key_risk_sheet(matrix)
        return fmt, rows, errors
    if fmt == FORMAT_LEGACY:
        rows, errors = _parse_legacy_sheet(matrix)
        return fmt, rows, errors

    header_idx = _find_header_row(matrix)
    if header_idx < 0:
        return fmt, [], [{"row": 0, "reason": "未找到表头行"}]
    header_map, level_col = _build_header_map(matrix, header_idx)
    for idx in range(header_idx + 1, len(matrix.rows)):
        row = matrix.rows[idx]
        if not any(_text(c) for c in row) or _is_footer_or_note(row):
            continue
        parsed = _parse_row_by_anchor(row, idx, fmt, header_map, level_col)
        if parsed is None or (not parsed.hazard_desc and parsed.d_value is None):
            if any(_to_float(c) is not None for c in row if _text(c)):
                errors.append({"row": idx + 1, "reason": "缺少风险等级或危险源描述"})
            continue
        rows.append(parsed)
    return fmt, rows, errors


def _row_preview(parsed: ParsedRow) -> dict[str, Any]:
    return {
        "name": parsed.name,
        "step": parsed.step or parsed.target,
        "hazard_desc": parsed.hazard_desc,
        "accident": parsed.accident,
        "l": parsed.l_value,
        "e": parsed.e_value,
        "c": parsed.c_value,
        "d": parsed.d_value,
        "level": LEVEL_LABEL_BY_KEY.get(parsed.level_key or "", parsed.level_raw),
        "engineering": (parsed.engineering or "")[:50] or None,
    }


@dataclass
class SheetAnalysis:
    """单个子表的完整解析结果。"""

    matrix: SheetMatrix
    fmt: str | None
    rows: list[ParsedRow]
    errors: list[dict[str, Any]]


def analyze_workbook(matrices: list[SheetMatrix]) -> list[SheetAnalysis]:
    analyses: list[SheetAnalysis] = []
    for matrix in matrices:
        fmt, rows, errors = parse_sheet(matrix)
        analyses.append(
            SheetAnalysis(matrix=matrix, fmt=fmt, rows=rows, errors=errors)
        )
    return analyses


def parse_workbook(data: bytes, filename: str) -> dict[str, Any]:
    """解析整个工作簿，返回预览结构（不写库）。"""
    matrices = load_workbook_matrices(data, filename)
    analyses = analyze_workbook(matrices)
    sheets: list[dict[str, Any]] = []
    total_rows = 0
    for analysis in analyses:
        non_empty = any(any(_text(c) for c in row) for row in analysis.matrix.rows)
        if not non_empty:
            sheets.append(
                {"name": analysis.matrix.name, "format": None, "format_label": None,
                 "row_count": 0, "skipped": True, "skip_reason": "空子表",
                 "sample_rows": [], "errors": []}
            )
            continue
        if analysis.fmt is None:
            sheets.append(
                {"name": analysis.matrix.name, "format": None, "format_label": None,
                 "row_count": 0, "skipped": True, "skip_reason": "无法识别的表格格式",
                 "sample_rows": [], "errors": analysis.errors[:5]}
            )
            continue
        if not analysis.rows:
            sheets.append(
                {"name": analysis.matrix.name, "format": analysis.fmt,
                 "format_label": FORMAT_LABELS[analysis.fmt], "row_count": 0,
                 "skipped": True, "skip_reason": "未解析出有效数据行",
                 "sample_rows": [], "errors": analysis.errors[:5]}
            )
            continue
        total_rows += len(analysis.rows)
        sheets.append(
            {
                "name": analysis.matrix.name,
                "format": analysis.fmt,
                "format_label": FORMAT_LABELS[analysis.fmt],
                "row_count": len(analysis.rows),
                "skipped": False,
                "skip_reason": None,
                "sample_rows": [_row_preview(r) for r in analysis.rows[:3]],
                "errors": analysis.errors[:5],
            }
        )
    return {
        "filename": os.path.basename(filename or ""),
        "department_guess": guess_department(filename),
        "sheets": sheets,
        "total_valid_rows": total_rows,
    }


# ── 记录构建与导入服务 ──


def build_hazard_records(
    parsed_rows: list[ParsedRow],
    *,
    department: str,
    position: str,
    source_label: str,
    is_facility: bool = False,
) -> list[dict[str, Any]]:
    """把解析行转换为 HazardIdentification 建库字段。

    同内容行按内容指纹去重，保证重复导入同一文件幂等跳过。
    """
    records: list[dict[str, Any]] = []
    seen_nos: set[str] = set()
    for parsed in parsed_rows:
        level_key = parsed.level_key
        if (
            level_key is None
            and parsed.d_value is not None
            and parsed.e_value is not None
        ):
            # 仅 LEC 口径可按 D 值兜底；LSR 的 R 值量纲不同，不做换算
            level_key = next(
                (
                    item["key"]
                    for item in RISK_LEVELS
                    if item["min_d"] <= parsed.d_value <= item["max_d"]
                ),
                None,
            )
        fingerprint = "|".join(
            [
                department,
                position,
                parsed.name or "",
                parsed.step or "",
                parsed.hazard_desc or "",
                parsed.accident or "",
                str(parsed.l_value),
                str(parsed.e_value),
                str(parsed.c_value),
                str(parsed.d_value),
            ]
        )
        hazard_id_no = "IMP-" + hashlib.sha1(
            fingerprint.encode("utf-8")
        ).hexdigest()[:12].upper()
        if hazard_id_no in seen_nos:
            continue
        seen_nos.add(hazard_id_no)

        notes_parts = [f"来源：导入《{source_label}》"]

        equipment: str | None = None
        if is_facility:
            equipment = "；".join(
                part for part in (parsed.name, parsed.location) if part
            ) or None

        record: dict[str, Any] = {
            "hazard_id_no": hazard_id_no,
            "department": department or "未指定",
            "position": position or "未分配",
            "production_step": parsed.name or parsed.step or position or "未命名",
            "specific_activity": parsed.step or parsed.target,
            "equipment_facilities": equipment,
            "hazard_type": parsed.category,
            "possible_accident": parsed.accident,
            "unsafe_behavior": parsed.hazard_desc,
            "l_inherent": parsed.l_value,
            "e_inherent": parsed.e_value,
            "c_inherent": parsed.c_value,
            "d_inherent": parsed.d_value,
            "inherent_risk_level": level_key,
            "inherent_risk_label": (
                LEVEL_LABEL_BY_KEY.get(level_key) if level_key else None
            ),
            "existing_engineering_controls": parsed.engineering,
            "existing_management_controls": parsed.management,
            "existing_ppe": parsed.ppe,
            "existing_emergency_measures": parsed.emergency,
            "control_level": parsed.control_level
            or (LEVEL_CONTROL_BY_KEY.get(level_key) if level_key else None),
            "responsible_person": parsed.responsible_person
            or (LEVEL_PERSON_BY_KEY.get(level_key) if level_key else None),
            # 车间表头专用字段（台账页面按 Excel 列分组展示）
            "tense": parsed.tense,
            "state": parsed.state,
            "responsible_unit": parsed.responsible_unit,
            "check_frequency": parsed.check_frequency,
            "check_standard": parsed.check_standard,
            "is_major_hazard": parsed.is_major,
            "needs_recommendation": "是" if parsed.extra_recommendation else "否",
            "recommendation_type": (
                "重点风险AI检索" if parsed.extra_recommendation else None
            ),
            "recommendation_content": parsed.extra_recommendation,
            "ai_node_progress": "completed",
            "overall_status": "completed",
            "notes": "；".join(notes_parts),
        }
        for script_no in range(1, 8):
            record[f"script{script_no}_review_status"] = "approved"
        records.append(record)
    return records


class HazardLedgerImportService:
    """危险源辨识台账 Excel 导入：预览解析与落库。"""

    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.repo = SafetyRepository(session)

    async def preview(self, filename: str, content: bytes) -> dict[str, Any]:
        if len(content) > MAX_UPLOAD_SIZE:
            raise ValueError("文件大小超过 10MB 限制")
        return parse_workbook(content, filename)

    async def commit(
        self,
        filename: str,
        content: bytes,
        department: str,
        selected_sheets: list[str] | None,
        user_id: uuid.UUID | None,
    ) -> dict[str, Any]:
        if len(content) > MAX_UPLOAD_SIZE:
            raise ValueError("文件大小超过 10MB 限制")
        matrices = load_workbook_matrices(content, filename)
        analyses = analyze_workbook(matrices)
        selected = set(selected_sheets) if selected_sheets else None
        source_filename = os.path.basename(filename or "")
        effective_department = department or guess_department(filename) or "未指定"

        imported = 0
        skipped_existing = 0
        duplicate_in_file = 0
        row_errors: list[dict[str, Any]] = []
        failed_sheets: list[dict[str, Any]] = []

        archive_path = self._archive_file(source_filename, content)

        for analysis in analyses:
            if analysis.fmt is None or not analysis.rows:
                if selected is not None and analysis.matrix.name in selected:
                    reason = (
                        "无法识别的表格格式"
                        if analysis.fmt is None
                        else "未解析出有效数据行"
                    )
                    failed_sheets.append(
                        {"sheet": analysis.matrix.name, "reason": reason}
                    )
                continue
            if selected is not None and analysis.matrix.name not in selected:
                continue
            records = build_hazard_records(
                analysis.rows,
                department=effective_department,
                position=analysis.matrix.name,
                source_label=f"{source_filename}/{analysis.matrix.name}",
                is_facility=analysis.fmt == FORMAT_FACILITY,
            )
            duplicate_in_file += len(analysis.rows) - len(records)
            if not records:
                continue
            for error in analysis.errors[:5]:
                row_errors.append({"sheet": analysis.matrix.name, **error})
            nos = [record["hazard_id_no"] for record in records]
            existing = await self.repo.get_existing_hazard_id_nos(nos)
            for record in records:
                if record["hazard_id_no"] in existing:
                    skipped_existing += 1
                    continue
                if user_id is not None:
                    record["created_by"] = user_id
                    record["updated_by"] = user_id
                record["attachment_path"] = archive_path
                record["attachment_original_name"] = source_filename
                await self.repo.create_hazard_identification(record)
                imported += 1

        await self._audit_import(
            user_id,
            filename=source_filename,
            imported=imported,
            skipped_existing=skipped_existing,
        )
        return {
            "imported": imported,
            "skipped_existing": skipped_existing,
            "duplicate_in_file": duplicate_in_file,
            "failed_sheets": failed_sheets,
            "row_errors": row_errors[:10],
            "archive_path": archive_path,
        }

    def _archive_file(self, filename: str, content: bytes) -> str:
        upload_dir = os.path.join("uploads", "safety", "hazard-import")
        os.makedirs(upload_dir, exist_ok=True)
        safe_name = re.sub(
            r'[\\/:*?"<>|]', "_", os.path.basename(filename) or "import.xlsx"
        )
        stamped = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{safe_name}"
        path = os.path.join(upload_dir, stamped)
        with open(path, "wb") as handle:
            handle.write(content)
        return path

    async def _audit_import(
        self,
        user_id: uuid.UUID | None,
        *,
        filename: str,
        imported: int,
        skipped_existing: int,
    ) -> None:
        try:
            await record_audit_log(
                self.session,
                action="import",
                user_id=user_id,
                resource_type="hazard_identification",
                extra={
                    "filename": filename,
                    "imported": imported,
                    "skipped_existing": skipped_existing,
                },
            )
        except Exception:  # pragma: no cover - 审计失败不阻塞导入
            logger.exception("危险源台账导入审计写入失败")
