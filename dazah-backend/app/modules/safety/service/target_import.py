"""目标责任与绩效：责任书签订登记与季度考核结果的 Excel 批量导入。

纸质签批的责任书与责任制考核表通过系统生成的模板 Excel 回填后整本上传：

- 责任书模板：单 sheet 台账，预填年度/部门/岗位/分类目录，回填甲方/乙方/
  签订日期即可批量登记签订状态；
- 考核模板：按考核表文档模板一个部门一个 sheet（含考核内容/标准分预填，
  得分列为公式展示），每季度填扣减分后上传，服务端按 标准分-扣减分 计分并
  自动评定等级。

解析为纯结构化函数（不依赖数据库），落库复用 SafetyTargetAgreementService /
SafetyDutyAssessmentService；不调用 LLM。
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import AppException
from app.modules.safety.models import SafetyTargetAgreement
from app.modules.safety.schemas.targets import (
    SafetyDutyAssessmentCreate,
    SafetyDutyAssessmentItemInput,
    SafetyDutyAssessmentUpdate,
    SafetyTargetAgreementCreate,
    SafetyTargetAgreementUpdate,
)
from app.modules.safety.service.target import (
    SafetyDutyAssessmentService,
    SafetyTargetAgreementService,
)

MAX_IMPORT_SIZE = 10 * 1024 * 1024

_HEADER_FILL = PatternFill("solid", fgColor="E6E0F5")
_HEADER_FONT = Font(bold=True)
_LABEL_FONT = Font(bold=True)

_AGREEMENT_HEADERS = (
    "年度",
    "责任部门/单位",
    "岗位/角色",
    "分类",
    "甲方（考核方）",
    "乙方（签订人）",
    "签订日期",
    "有效期至",
    "备注",
)

_ASSESSMENT_HEADERS = (
    "序号", "考核内容", "考核标准", "标准分", "扣减分", "考核得分", "备注"
)


def _check_size(content: bytes) -> None:
    if len(content) > MAX_IMPORT_SIZE:
        raise AppException(status_code=400, message="文件大小超过 10MB 限制")


def _as_date(value: Any) -> date | None:
    """宽松解析日期单元格：datetime/date/常用分隔的字符串。"""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace("年", "-").replace("月", "-").replace("日", "")
    text = re.sub(r"[./]", "-", text)
    match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", text)
    if not match:
        return None
    try:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def _as_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d+(?:\.\d+)?", str(value))
    if not match:
        return None
    return float(match.group(0))


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _style_header_row(ws: Any, row_idx: int, width: int) -> None:
    for col in range(1, width + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


# ==================== 责任书签订登记导入 ====================


def build_agreement_template(rows: list[SafetyTargetAgreement]) -> bytes:
    """生成责任书签订登记模板：预填现有台账目录。"""
    wb = openpyxl.Workbook()
    # openpyxl 存根把 active 标为可空联合类型，此处恒为新建工作表
    ws: Any = wb.active
    ws.title = "责任书签订登记"
    ws.append(list(_AGREEMENT_HEADERS))
    _style_header_row(ws, 1, len(_AGREEMENT_HEADERS))
    for item in rows:
        ws.append(
            [
                item.year,
                item.department,
                item.position,
                item.category,
                item.party_a,
                item.party_b,
                item.signed_date or "",
                item.expires_date or "",
                item.remark,
            ]
        )
    widths = (8, 22, 22, 16, 16, 12, 12, 12, 24)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@dataclass
class AgreementImportRow:
    year: int
    department: str
    position: str
    category: str
    party_a: str
    party_b: str
    signed_date: date | None
    expires_date: date | None
    remark: str


def parse_agreement_template(content: bytes) -> list[AgreementImportRow]:
    """解析签订登记模板：按表头列名取列，容忍列序变化与空行。"""
    _check_size(content)
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    # openpyxl 存根把 active 标为可空联合类型，xlsx 恒有默认工作表
    ws: Any = wb.active
    header_row: dict[str, int] = {}
    for cell in ws[1]:
        label = _cell_text(cell.value)
        if label:
            header_row[label] = cell.column
    missing = [
        name
        for name in ("年度", "责任部门/单位", "岗位/角色")
        if name not in header_row
    ]
    if missing:
        raise AppException(
            status_code=400, message=f"模板表头缺少列：{'、'.join(missing)}"
        )

    def col(name: str) -> int:
        return header_row.get(name, 0)

    rows: list[AgreementImportRow] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        values = list(row) + [""] * len(_AGREEMENT_HEADERS)
        year = _as_number(values[col("年度") - 1])
        department = _cell_text(values[col("责任部门/单位") - 1])
        position = _cell_text(values[col("岗位/角色") - 1])
        if not department or not position or year is None:
            continue
        signed_date = _as_date(values[col("签订日期") - 1]) if col("签订日期") else None
        expires_col = col("有效期至")
        expires_date = _as_date(values[expires_col - 1]) if expires_col else None
        category = (
            _cell_text(values[col("分类") - 1]) if col("分类") else "岗位员工"
        ) or "岗位员工"
        party_a_col = col("甲方（考核方）")
        party_b_col = col("乙方（签订人）")
        rows.append(
            AgreementImportRow(
                year=int(year),
                department=department,
                position=position,
                category=category,
                party_a=(
                    _cell_text(values[party_a_col - 1]) if party_a_col else ""
                ),
                party_b=(
                    _cell_text(values[party_b_col - 1]) if party_b_col else ""
                ),
                signed_date=signed_date,
                expires_date=expires_date,
                remark=_cell_text(values[col("备注") - 1]) if col("备注") else "",
            )
        )
    if not rows:
        raise AppException(status_code=400, message="模板中没有可导入的数据行")
    return rows


@dataclass
class AgreementImportReport:
    updated: int = 0
    created: int = 0
    errors: list[dict[str, Any]] = field(default_factory=list)


class AgreementImportService:
    """责任书签订登记批量导入：存在即更新（含签订登记），缺失则新建。"""

    def __init__(self, db: AsyncSession) -> None:
        self.db = db
        self.service = SafetyTargetAgreementService(db)

    async def commit(
        self, rows: list[AgreementImportRow], user_id: Any
    ) -> AgreementImportReport:
        report = AgreementImportReport()
        for index, row in enumerate(rows, start=2):
            try:
                existing = await self.service.get_by_key(
                    row.year, row.department, row.position
                )
                payload: dict[str, Any] = {
                    "party_a": row.party_a,
                    "party_b": row.party_b,
                    "category": row.category,
                    "remark": row.remark,
                }
                if row.expires_date is not None:
                    payload["expires_date"] = row.expires_date
                if row.signed_date is not None:
                    payload["signed_date"] = row.signed_date
                    payload["status"] = "signed"
                if existing is None:
                    await self.service.create_agreement(
                        SafetyTargetAgreementCreate(
                            year=row.year,
                            department=row.department,
                            position=row.position,
                            status="signed" if row.signed_date else "unsigned",
                            signed_date=row.signed_date,
                            **payload,
                        ),
                        user_id=user_id,
                    )
                    report.created += 1
                else:
                    await self.service.update_agreement(
                        existing.id,
                        SafetyTargetAgreementUpdate(**payload),
                        user_id=user_id,
                    )
                    report.updated += 1
            except Exception as exc:  # noqa: BLE001 —— 单行失败记录后继续，不中断整批
                report.errors.append(
                    {
                        "row": index,
                        "target": f"{row.department}/{row.position}",
                        "message": str(exc),
                    }
                )
        return report


# ==================== 季度考核批量导入 ====================


def build_assessment_template(year: int, quarter: int) -> bytes:
    """生成季度考核模板：每个考核表模板一个 sheet，考核内容/标准分预填。"""
    from app.modules.safety.service.target import _load_responsibility_seed

    templates = _load_responsibility_seed().get("assessment_templates", [])
    wb = openpyxl.Workbook()
    # openpyxl 存根把 active 标为可空联合类型，此处恒为新建工作表
    guide: Any = wb.active
    guide.title = "填写说明"
    guide_lines = [
        "安全生产责任制季度考核批量导入模板",
        f"考核期间：{year} 年第 {quarter} 季度（上传时需选择相同年度/季度）",
        "",
        "填写说明：",
        "1. 每个 sheet 对应一个被考核部门/负责人（与 2026 责任制考核表一致）；",
        "2. 只需填写「扣减分」列（不扣的项留空即得满分），并填写表头的"
        "被考核人/考核单位/考核日期；",
        "3. 「考核得分」列为公式展示，无需填写，导入时按 标准分-扣减分（下限 0）"
        "重新计算；",
        "4. 扣减分为空的 sheet 不会导入；导入按 年度+季度+部门+被考核人 覆盖更新；",
        "5. 单元格扣分可填数字（如 2）或带说明（如 未按要求执行扣2分），取数字部分。",
    ]
    for line in guide_lines:
        guide.append([line])
    guide.column_dimensions["A"].width = 80
    guide["A1"].font = Font(bold=True, size=13)

    for tpl in templates:
        name = str(tpl.get("name", "")).strip()[:31]
        if not name:
            continue
        ws: Any = wb.create_sheet(title=name)
        ws["A1"] = "被考核部门"
        ws["B1"] = name
        ws["A2"] = "被考核人"
        ws["B2"] = ""
        ws["D1"] = "考核单位"
        ws["E1"] = ""
        ws["D2"] = "考核日期"
        ws["E2"] = ""
        for label_cell in ("A1", "A2", "D1", "D2"):
            ws[label_cell].font = _LABEL_FONT
        ws.append([])
        ws.append(list(_ASSESSMENT_HEADERS))
        header_row_idx = ws.max_row
        _style_header_row(ws, header_row_idx, len(_ASSESSMENT_HEADERS))
        for order, item in enumerate(tpl.get("items", []), start=1):
            row_idx = ws.max_row + 1
            standard = float(item.get("standard_score") or 0)
            ws.append(
                [
                    order,
                    item.get("content", ""),
                    item.get("rule", ""),
                    standard,
                    None,
                    f"=MAX(0,D{row_idx}-E{row_idx})",
                    "",
                ]
            )
        total_row = ws.max_row + 1
        ws.append(
            [
                "合计",
                "",
                "",
                f"=SUM(D5:D{total_row - 1})",
                f"=SUM(E5:E{total_row - 1})",
                f"=SUM(F5:F{total_row - 1})",
                "",
            ]
        )
        ws.cell(row=total_row, column=1).font = _HEADER_FONT
        widths = (6, 60, 30, 8, 8, 10, 20)
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A5"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@dataclass
class AssessmentSheetData:
    sheet_name: str
    department: str
    person_name: str
    assessor: str
    assessment_date: date | None
    items: list[SafetyDutyAssessmentItemInput]


def parse_assessment_template(
    content: bytes,
) -> tuple[list[AssessmentSheetData], list[str]]:
    """解析考核导入工作簿：返回有效 sheet 与被跳过的 sheet 名列表。"""
    _check_size(content)
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheets: list[AssessmentSheetData] = []
    skipped: list[str] = []

    for sheet in wb.worksheets:
        # openpyxl 存根的 worksheet 联合类型过宽，统一按 Any 处理
        ws: Any = sheet
        if ws.title in ("填写说明", "Sheet", "Worksheet"):
            continue
        header_labels: dict[str, tuple[int, int]] = {}
        for row in ws.iter_rows(min_row=1, max_row=6):
            for cell in row:
                label = _cell_text(cell.value)
                if label in ("被考核部门", "部门", "被考核人", "考核单位", "考核日期"):
                    header_labels.setdefault(label, (cell.row, cell.column))

        def header_value(label: str) -> str:
            pos = header_labels.get(label)
            if pos is None:
                return ""
            r, c = pos
            return _cell_text(ws.cell(row=r, column=c + 1).value)

        # 表头行：同时包含 考核内容 与 扣减分 的行
        header_row_idx = None
        columns: dict[str, int] = {}
        for row in ws.iter_rows(min_row=1, max_row=12):
            row_cols = {
                _cell_text(cell.value): cell.column
                for cell in row
                if _cell_text(cell.value)
            }
            if "考核内容" in row_cols and "扣减分" in row_cols:
                header_row_idx = row[0].row
                columns = row_cols
                break
        if header_row_idx is None:
            skipped.append(ws.title)
            continue

        items: list[SafetyDutyAssessmentItemInput] = []
        has_deduction = False
        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            values = list(row) + [""] * 8
            content_text = _cell_text(values[columns["考核内容"] - 1])
            if not content_text or content_text.startswith(("合计", "考核总得分")):
                break
            standard = _as_number(values[columns.get("标准分", 4) - 1]) or 0.0
            deduction = _as_number(values[columns["扣减分"] - 1]) or 0.0
            if deduction > 0:
                has_deduction = True
            rule = (
                _cell_text(values[columns["考核标准"] - 1])
                if "考核标准" in columns
                else ""
            )
            remark_col = columns.get("备注")
            remark = _cell_text(values[remark_col - 1]) if remark_col else ""
            items.append(
                SafetyDutyAssessmentItemInput(
                    content=content_text,
                    rule=rule,
                    standard_score=standard,
                    deduction_score=deduction,
                    remark=remark,
                )
            )
        if not items:
            skipped.append(ws.title)
            continue
        if not has_deduction:
            skipped.append(ws.title)
            continue

        department = header_value("被考核部门") or header_value("部门") or ws.title
        sheets.append(
            AssessmentSheetData(
                sheet_name=ws.title,
                department=department,
                person_name=header_value("被考核人"),
                assessor=header_value("考核单位"),
                assessment_date=_as_date(header_value("考核日期")),
                items=items,
            )
        )

    if not sheets:
        raise AppException(
            status_code=400,
            message="工作簿中没有可导入的考核数据（只有填了扣减分的 sheet 才会导入）",
        )
    return sheets, skipped


@dataclass
class AssessmentImportReport:
    imported: int = 0
    updated: int = 0
    skipped_sheets: list[str] = field(default_factory=list)
    errors: list[dict[str, Any]] = field(default_factory=list)


class DutyAssessmentImportService:
    """季度考核批量导入：按 年度+季度+部门+被考核人 覆盖更新。"""

    def __init__(self, db: AsyncSession) -> None:
        self.db = db
        self.service = SafetyDutyAssessmentService(db)

    async def commit(
        self,
        sheets: list[AssessmentSheetData],
        *,
        year: int,
        quarter: int,
        default_assessor: str,
        skipped_sheets: list[str] | None = None,
        user_id: Any,
    ) -> AssessmentImportReport:
        report = AssessmentImportReport(skipped_sheets=list(skipped_sheets or []))
        for sheet in sheets:
            try:
                existing = await self.service.get_by_key(
                    year, quarter, sheet.department, sheet.person_name
                )
                payload = {
                    "assessor": sheet.assessor or default_assessor,
                    "assessment_date": sheet.assessment_date,
                    "items": [item.model_dump() for item in sheet.items],
                }
                if existing is None:
                    await self.service.create_assessment(
                        SafetyDutyAssessmentCreate(
                            year=year,
                            quarter=quarter,
                            department=sheet.department,
                            person_name=sheet.person_name,
                            remark="",
                            **payload,
                        ),
                        user_id=user_id,
                    )
                    report.imported += 1
                else:
                    await self.service.update_assessment(
                        existing.id,
                        SafetyDutyAssessmentUpdate(**payload),
                        user_id=user_id,
                    )
                    report.updated += 1
            except Exception as exc:  # noqa: BLE001 —— 单表失败记录后继续
                report.errors.append(
                    {
                        "row": sheet.sheet_name,
                        "target": sheet.department,
                        "message": str(exc),
                    }
                )
        return report
