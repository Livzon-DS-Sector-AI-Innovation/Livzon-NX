"""排产计划 Excel 存档：解析、保存与查询。

解析协议与前端 scheduling 页一致：
- 单元格值为标量或空字符串（空值统一为 ""）；
- 合并单元格用 0-based {s: {r, c}, e: {r, c}}；
- 行列不截断，rows 为完整二维数组。
"""
from __future__ import annotations

import uuid
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from typing import Any

import openpyxl  # type: ignore[import-untyped]
import xlrd  # type: ignore[import-untyped]
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.modules.production import fermentation_board_service
from app.modules.production.schedule_excel_models import ScheduleExcelArchive


def _serialize_cell(value: Any) -> Any:
    """把 openpyxl 单元格值转成 JSON 安全且可读的展示值。

    日期/时间转可读字符串（前端本地解析时日期会显示为序列号，
    这里统一转成可读文本更符合“不转换字段”的业务直觉）；其余原样。
    """
    if value is None:
        return ""
    if isinstance(value, datetime):
        # Excel 无独立 date 类型，date 读出为零点 datetime → 显示为日期
        if (
            value.hour == 0
            and value.minute == 0
            and value.second == 0
            and value.microsecond == 0
        ):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def parse_workbook_bytes(data: bytes) -> dict[str, Any]:
    """解析 Excel 首个可见工作表，返回与前端渲染协议一致的结构。

    按文件魔数分派：.xlsx/.xlsm（ZIP）走 openpyxl；.xls（OLE 复合文档）
    走 xlrd（openpyxl 不支持 BIFF 格式）。两者产出的 rows/merges 协议一致。
    多周期排产文件常把历史 Sheet 隐藏保留；首表可能命中隐藏表，
    因此优先取第一个可见工作表，全部隐藏时回退首表。
    """
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return _parse_xls_bytes(data)
    return _parse_xlsx_bytes(data)


def _parse_xlsx_bytes(data: bytes) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(BytesIO(data), data_only=True)
    try:
        visible = [ws for ws in workbook.worksheets if ws.sheet_state == "visible"]
        worksheet = visible[0] if visible else workbook.worksheets[0]
    finally:
        workbook.close()

    rows: list[list[Any]] = []
    for row in worksheet.iter_rows():  # 空行也保留，与前端 blankrows 一致
        rows.append([_serialize_cell(cell.value) for cell in row])

    merges = [
        {
            "s": {"r": rng.min_row - 1, "c": rng.min_col - 1},
            "e": {"r": rng.max_row - 1, "c": rng.max_col - 1},
        }
        for rng in worksheet.merged_cells.ranges
    ]

    from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

    col_widths: list[int] = []
    for index in range(1, worksheet.max_column + 1):
        dimension = worksheet.column_dimensions.get(get_column_letter(index))
        width = dimension.width if dimension and dimension.width else 80
        col_widths.append(int(width))

    return {
        "sheet_name": worksheet.title,
        "rows": rows,
        "merges": merges,
        "col_widths": col_widths,
        "row_count": len(rows),
        "col_count": worksheet.max_column,
    }


def _xls_text_units(text: str) -> int:
    """显示宽度单位：中文等全宽字符记 2，其余记 1。"""
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in text)


def _parse_xls_bytes(data: bytes) -> dict[str, Any]:
    """xlrd 解析 .xls：单元格序列化与 _serialize_cell 口径一致。"""
    book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    visible = [sheet for sheet in book.sheets() if sheet.visibility == 0]
    worksheet = visible[0] if visible else book.sheet_by_index(0)

    def _cell_value(r: int, c: int) -> Any:
        cell = worksheet.cell(r, c)
        if cell.ctype == xlrd.XL_CELL_DATE:
            value = xlrd.xldate_as_datetime(cell.value, book.datemode)
            if (value.year, value.month, value.day) == (1899, 12, 31):
                return value.strftime("%H:%M:%S")  # 纯时刻（小数时间）
            if (
                value.hour == 0
                and value.minute == 0
                and value.second == 0
                and value.microsecond == 0
            ):
                return value.strftime("%Y-%m-%d")
            return value.strftime("%Y-%m-%d %H:%M:%S")
        if cell.ctype == xlrd.XL_CELL_NUMBER:
            return cell.value
        if cell.ctype == xlrd.XL_CELL_BOOLEAN:
            return bool(cell.value)
        if cell.ctype == xlrd.XL_CELL_TEXT:
            return str(cell.value)
        return ""  # 空/错误/空白

    rows = [
        [_cell_value(r, c) for c in range(worksheet.ncols)]
        for r in range(worksheet.nrows)
    ]

    # xlrd 合并区间为半开 (rlo, rhi, clo, chi)，转 0-based 闭区间端点
    merges = [
        {
            "s": {"r": rlo, "c": clo},
            "e": {"r": rhi - 1, "c": chi - 1},
        }
        for rlo, rhi, clo, chi in worksheet.merged_cells
    ]

    # 列宽按内容自适应：非合并格按所在列计；跨列合并格（标题/备注横跨
    # 多列）渲染时占整行宽，其显示宽度均摊到覆盖列——单列无需独立放下。
    # 宽度按序列化后的显示值计（换行符当空格、中文按 2 倍字符计）。
    span_starts: dict[tuple[int, int], int] = {}
    covered: set[tuple[int, int]] = set()
    for rlo, rhi, clo, chi in worksheet.merged_cells:
        if chi - clo <= 1:
            continue
        span_starts[(rlo, clo)] = chi - clo
        for r in range(rlo, rhi):
            for c in range(clo, chi):
                if (r, c) != (rlo, clo):
                    covered.add((r, c))

    col_units = [0.0] * worksheet.ncols
    for r in range(len(rows)):
        for c in range(worksheet.ncols):
            if (r, c) in covered:
                continue
            value = rows[r][c] if c < len(rows[r]) else ""
            if value in (None, ""):
                continue
            flat = str(value).replace("\n", " ").strip()
            if not flat:
                continue
            share = _xls_text_units(flat) / span_starts.get((r, c), 1)
            for cc in range(c, min(c + span_starts.get((r, c), 1), worksheet.ncols)):
                col_units[cc] = max(col_units[cc], share)
    col_widths = [
        min(300, max(28, round(units * 8 + 12))) for units in col_units
    ]

    return {
        "sheet_name": worksheet.name,
        "rows": rows,
        "merges": merges,
        "col_widths": col_widths,
        "row_count": len(rows),
        "col_count": worksheet.ncols,
    }


def serialize_archive(
    archive: ScheduleExcelArchive,
    *,
    created_by_name: str | None = None,
) -> dict[str, Any]:
    """ORM → API 响应（mode=json 友好的 dict）。"""
    return {
        "id": str(archive.id),
        "product_code": archive.product_code,
        "file_name": archive.file_name,
        "sheet_name": archive.sheet_name,
        "original_path": archive.original_path,
        "rows": archive.rows,
        "merges": archive.merges,
        "col_widths": archive.col_widths,
        "row_count": archive.row_count,
        "col_count": archive.col_count,
        "created_by_name": created_by_name,
        "created_at": (
            archive.created_at.isoformat() if archive.created_at else None
        ),
        "updated_at": (
            archive.updated_at.isoformat() if archive.updated_at else None
        ),
    }


def serialize_archive_summary(
    archive: ScheduleExcelArchive,
    *,
    created_by_name: str | None = None,
) -> dict[str, Any]:
    """列表项（不含 rows，避免大字段传输）。"""
    payload = serialize_archive(archive, created_by_name=created_by_name)
    payload.pop("rows", None)
    payload.pop("merges", None)
    payload.pop("col_widths", None)
    payload.pop("original_path", None)
    return payload


# 排产表按北京时间维护，容器多为 UTC：冻结基准日统一取北京日期
_BEIJING_TZ = timezone(timedelta(hours=8))


async def create_archive(
    session: AsyncSession,
    *,
    product_code: str = "FA",
    file_name: str,
    sheet_name: str,
    original_path: str,
    rows: list[list[Any]],
    merges: list[dict[str, Any]],
    col_widths: list[int],
    row_count: int,
    col_count: int,
    created_by: uuid.UUID | None = None,
    history_fix: bool = False,
    history_fix_reason: str = "",
) -> tuple[ScheduleExcelArchive, dict[str, Any]]:
    # 同产品重复存档时冻结历史：今天之前的日列沿用当前最新存档，
    # 避免重发的排产改动/漏带历史放罐记录覆盖看板历史口径。
    # history_fix=True 时以新文件修正历史（须填原因，由 API 层校验权限）。
    latest = await fermentation_board_service.load_latest_archive(
        session, product_code=product_code
    )
    report: dict[str, Any] = {
        "recognized": True,
        "matched_blocks": 0,
        "frozen_columns": 0,
        "discarded_changes": [],
        "truncated": False,
        "corrected": False,
    }
    if latest is not None and latest.rows:
        merged, report = (
            fermentation_board_service.merge_schedule_rows_for_product(
                rows,
                latest.rows,
                datetime.now(_BEIJING_TZ).date(),
                product_code,
                freeze_past=not history_fix,
            )
        )
        if not report["recognized"]:
            example = fermentation_board_service.SHEET_FORMAT_EXAMPLES.get(
                product_code, "标准扎帐周期块标题（对照最近一次正常存档的排产表）"
            )
            raise ValueError(
                f"未能在新文件中识别任何排产周期块，已取消存档。"
                f"{product_code} 排产表块标题格式应为：{example}"
            )
        rows = merged
        if report["discarded_changes"] and history_fix:
            if not history_fix_reason.strip():
                raise ValueError("应用历史修正必须填写修正原因")
            report["corrected"] = True
        elif report["discarded_changes"]:
            report["warning"] = (
                f"检测到 {len(report['discarded_changes'])} 处「今天之前」的"
                "历史改动，已按冻结规则保留原存档"
                + ("（仅展示前 200 处）" if report["truncated"] else "")
            )
    archive = ScheduleExcelArchive(
        product_code=product_code,
        file_name=file_name,
        sheet_name=sheet_name,
        original_path=original_path,
        rows=rows,
        merges=merges,
        col_widths=col_widths,
        row_count=row_count,
        col_count=col_count,
        created_by=created_by,
    )
    session.add(archive)
    await session.commit()
    await session.refresh(archive)
    return archive, report


async def count_history_fixes(session: AsyncSession, product_code: str) -> int:
    """统计某产品累计历史修正次数（audit.logs 为唯一事实源）。

    修正记录量级极小,取回后在 Python 侧按 extra.product_code 过滤,
    避免 JSON/JSONB 方言差异（extra 列是通用 JSON,无 .astext）。
    """
    from app.platform.audit.models import AuditLog

    result = await session.execute(
        select(AuditLog.extra).where(
            AuditLog.action == "schedule_excel_history_fix",
            AuditLog.resource_type == "schedule_excel_archive",
        )
    )
    return sum(
        1
        for (extra,) in result
        if isinstance(extra, dict)
        and extra.get("product_code") == product_code
    )


async def list_history_fixes(
    session: AsyncSession,
    archive_ids: list[uuid.UUID],
) -> dict[uuid.UUID, list[dict[str, Any]]]:
    """批量查询存档的历史修正记录（展示用，audit.logs 仍为唯一事实源）。

    返回 {存档 id: [修正记录]}；每条含操作时间/操作人/原因与逐格改动
    （changes 最多返回 50 条，changes_total 记实际总数）。
    """
    if not archive_ids:
        return {}
    # 平台审计模型在服务层懒加载，避免模块导入期耦合（同 hr/service 模式）
    from app.platform.audit.models import AuditLog
    from app.platform.identity.models import User

    rows = (
        await session.execute(
            select(AuditLog, User.name)
            .outerjoin(User, User.id == AuditLog.user_id)
            .where(
                AuditLog.action == "schedule_excel_history_fix",
                AuditLog.resource_type == "schedule_excel_archive",
                AuditLog.resource_id.in_(archive_ids),
            )
            .order_by(AuditLog.created_at.asc())
        )
    ).all()
    fixes: dict[uuid.UUID, list[dict[str, Any]]] = {}
    for log, user_name in rows:
        if log.resource_id is None:
            continue
        extra = log.extra or {}
        changes = (log.new_value or {}).get("changes") or []
        fixes.setdefault(log.resource_id, []).append(
            {
                "fixed_at": (
                    log.created_at.isoformat() if log.created_at else None
                ),
                "fixed_by_name": user_name,
                "reason": extra.get("reason"),
                "product_code": extra.get("product_code"),
                "file_name": extra.get("file_name"),
                "changes": changes[:50],
                "changes_total": len(changes),
            }
        )
    return fixes


async def list_archives(
    session: AsyncSession,
    *,
    page: int,
    page_size: int,
    product_code: str = "FA",
) -> tuple[list[tuple[ScheduleExcelArchive, str | None]], int]:
    """分页列表，附带上传人姓名（created_by 外键 join identity.users）。"""
    from app.platform.identity.models import User

    total = await session.scalar(
        select(func.count(ScheduleExcelArchive.id)).where(
            ScheduleExcelArchive.is_deleted.is_(False),
            ScheduleExcelArchive.product_code == product_code,
        )
    )
    query = (
        select(ScheduleExcelArchive, User.name)
        .outerjoin(User, User.id == ScheduleExcelArchive.created_by)
        .where(
            ScheduleExcelArchive.is_deleted.is_(False),
            ScheduleExcelArchive.product_code == product_code,
        )
        .order_by(ScheduleExcelArchive.created_at.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
    )
    result = await session.execute(query)
    items = [(archive, name) for archive, name in result.all()]
    return items, int(total or 0)


async def get_user_name(
    session: AsyncSession, user_id: uuid.UUID | None
) -> str | None:
    """按用户 ID 取姓名（用于详情等单条响应）。"""
    if user_id is None:
        return None
    from app.platform.identity.models import User

    name = await session.scalar(
        select(User.name).where(User.id == user_id)
    )
    return name


async def get_archive(
    session: AsyncSession, archive_id: uuid.UUID
) -> ScheduleExcelArchive | None:
    result = await session.execute(
        select(ScheduleExcelArchive).where(
            ScheduleExcelArchive.id == archive_id,
            ScheduleExcelArchive.is_deleted.is_(False),
        )
    )
    return result.scalar_one_or_none()


async def delete_archive(
    session: AsyncSession,
    archive: ScheduleExcelArchive,
    *,
    deleted_by: uuid.UUID | None = None,
) -> None:
    archive.is_deleted = True
    archive.updated_by = deleted_by
    await session.commit()
